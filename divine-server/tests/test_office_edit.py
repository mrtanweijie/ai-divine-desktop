"""document_outline / office 备份恢复 / 脚本端到端（外部副本、原地备份）。"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402
from openpyxl import Workbook  # noqa: E402

from app.docgen.outline import docx_outline, xlsx_outline  # noqa: E402
from app.tools.office_backup import backup_file, restore_latest_batch  # noqa: E402
from app.tools.office_files import (  # noqa: E402
    OfficePathError,
    prepare_edit_target,
    resolve_office_path,
)


def _make_docx(path: Path) -> None:
    document = Document()
    document.add_heading("测试文档", 1)
    document.add_paragraph("第一段内容")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "名称"
    table.rows[0].cells[1].text = "数量"
    document.save(str(path))


def test_docx_outline_coordinates():
    path = Path("/tmp/outline_test.docx")
    _make_docx(path)
    result = docx_outline(path)
    assert result["kind"] == "docx"
    assert result["paragraphs"][0]["preview"] == "测试文档"
    assert result["paragraphs"][0]["style"] == "Heading 1"
    assert result["tables"][0]["index"] == 1
    assert result["tables"][0]["headers"] == ["名称", "数量"]


def test_xlsx_outline_sheets_and_formulas():
    path = Path("/tmp/outline_test.xlsx")
    workbook = Workbook()
    ws = workbook.active
    ws.title = "9月"
    ws.append(["日期", "金额"])
    ws["B2"] = 100
    ws["B3"] = "=SUM(B2:B2)"
    workbook.save(str(path))
    result = xlsx_outline(path)
    sheet = result["sheets"][0]
    assert sheet["name"] == "9月"
    assert sheet["headers"][:2] == ["日期", "金额"]
    assert sheet["has_formulas"] == 1
    assert sheet["has_chart"] is False


def test_resolve_office_path_rejects_doc_and_missing(tmp_path):
    ws = tmp_path / "ws"
    ws.mkdir()
    try:
        resolve_office_path(ws, "不存在.docx")
        raised = False
    except OfficePathError as exc:
        raised = True
        assert "不存在" in str(exc)
    assert raised

    doc = ws / "old.doc"
    doc.write_bytes(b"x")
    try:
        resolve_office_path(ws, str(doc))
        raised = False
    except OfficePathError as exc:
        raised = True
        assert ".doc" in str(exc)
    assert raised


def test_prepare_edit_target_backs_up_workspace_file(tmp_path, monkeypatch):
    monkeypatch.setenv("DIVINE_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("DB_PATH", str(tmp_path / "eval-office.db"))
    ws = tmp_path / "ws"
    ws.mkdir()
    doc = ws / "report.docx"
    _make_docx(doc)
    prep = prepare_edit_target(ws, "sess-1", doc)
    assert prep["in_workspace"] is True and prep["is_copy"] is False
    assert prep["backup"] is not None and prep["backup"].is_file()


def test_prepare_edit_target_copies_external_file(tmp_path):
    ws = tmp_path / "ws"
    ws.mkdir()
    external = tmp_path / "Desktop" / "合同.docx"
    external.parent.mkdir(parents=True)
    _make_docx(external)
    original_bytes = external.read_bytes()

    prep = prepare_edit_target(ws, "sess-1", external)
    assert prep["is_copy"] is True
    assert prep["target"].parent == ws / "documents" / "edited"
    assert external.read_bytes() == original_bytes  # 原文件未被触碰
    assert prep["backup"] is None


def test_backup_restore_batch_and_prune(tmp_path, monkeypatch):
    monkeypatch.setenv("DIVINE_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("DB_PATH", str(tmp_path / "eval-office.db"))
    doc = tmp_path / "report.docx"

    # 真实时序：每次编辑前先备份（备份内容 = 编辑前状态）
    doc.write_text("v1", encoding="utf-8")
    first = backup_file("s1", doc)  # b1 = v1
    doc.write_text("v2", encoding="utf-8")
    backup_file("s1", doc)  # b2 = v2
    doc.write_text("v3", encoding="utf-8")
    backup_file("s1", doc)  # b3 = v3
    doc.write_text("v4", encoding="utf-8")
    fourth = backup_file("s1", doc)  # b4 = v4 → b1 被剪枝
    doc.write_text("v5", encoding="utf-8")  # 最新一次编辑未备份

    assert first is not None and not first.exists()  # 最旧被剪枝
    assert fourth is not None and fourth.is_file()

    # undo：恢复最近一次备份（b4 = v4），恢复即消费
    result = restore_latest_batch("s1")
    assert len(result["restored"]) == 1
    assert doc.read_text(encoding="utf-8") == "v4"
    # 连续 undo 逐批回溯到 v3
    restore_latest_batch("s1")
    assert doc.read_text(encoding="utf-8") == "v3"


def _run_script(name: str, payload: dict) -> dict:
    script = Path(__file__).resolve().parents[1] / "app" / "tools" / "scripts" / name
    completed = subprocess.run(
        [sys.executable, str(script)],
        input=json.dumps(payload, ensure_ascii=False),
        capture_output=True,
        text=True,
        timeout=60,
        check=False,
    )
    assert completed.returncode == 0, completed.stderr[-2000:]
    return json.loads(completed.stdout)


def test_edit_docx_script_external_file_creates_copy(tmp_path):
    ws = tmp_path / "ws"
    (ws / "documents").mkdir(parents=True)
    external = tmp_path / "Desktop"
    external.mkdir()
    doc = external / "周报.docx"
    _make_docx(doc)
    original_bytes = doc.read_bytes()

    output = _run_script(
        "edit_docx.py",
        {
            "workspace": str(ws),
            "output_dir": str(ws),
            "params": {
                "path": str(doc),
                "session_id": "s-script",
                "ops": [
                    {"op": "replace_text", "find": "第一段内容", "replace": "已更新内容"}
                ],
            },
        },
    )
    assert output["success"] is True
    assert doc.read_bytes() == original_bytes  # 原文件未动
    copy_path = Path(output["filepath"])
    assert copy_path.parent == ws / "documents" / "edited"
    reopened = Document(str(copy_path))
    assert any("已更新内容" in p.text for p in reopened.paragraphs)


def test_edit_docx_script_transaction_failure_leaves_file_untouched(tmp_path):
    ws = tmp_path / "ws"
    ws.mkdir()
    doc = ws / "report.docx"
    _make_docx(doc)
    before = doc.read_bytes()

    output = _run_script(
        "edit_docx.py",
        {
            "workspace": str(ws),
            "output_dir": str(ws),
            "params": {
                "path": str(doc),
                "session_id": "s-script",
                "ops": [
                    {"op": "replace_text", "find": "不存在的文本", "replace": "x"}
                ],
            },
        },
    )
    assert output["success"] is False
    assert doc.read_bytes() == before  # 文件未做任何修改


def test_document_outline_script(tmp_path):
    ws = tmp_path / "ws"
    ws.mkdir()
    doc = ws / "doc.docx"
    _make_docx(doc)
    output = _run_script(
        "document_outline.py",
        {"workspace": str(ws), "output_dir": str(ws), "params": {"path": str(doc)}},
    )
    assert output["success"] is True
    assert output["data"]["kind"] == "docx"
    assert "表格1" in output["content"] or "表格 1" in output["content"]


def test_pptx_outline_pages(tmp_path):
    """document_outline pptx 分支：页码/标题/预览/备注。"""
    from pptx import Presentation
    from pptx.util import Inches

    from app.docgen.outline import pptx_outline

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide1 = prs.slides.add_slide(prs.slide_layouts[5])
    slide1.shapes.title.text_frame.text = "封面标题"
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1.5))
    box.text_frame.text = "第二页正文内容"
    prs.save(str(tmp_path / "deck.pptx"))

    result = pptx_outline(tmp_path / "deck.pptx")
    assert result["kind"] == "pptx"
    assert len(result["slides"]) == 2
    assert result["slides"][0]["title"] == "封面标题"
    assert result["slides"][1]["index"] == 2
    assert any("第二页正文内容" in p for p in result["slides"][1]["previews"])
def test_word_base_style_and_header(tmp_path):
    """样式基底模式：清空正文、保留样式定义/页眉，新内容字体派生自基底。"""
    script = Path(__file__).resolve().parents[1] / "app" / "tools" / "scripts" / "word_document.py"
    ws_dir = tmp_path / "ws"
    ws_dir.mkdir()

    base = Document()
    base.styles["Normal"].font.name = "Courier New"
    base.sections[0].header.paragraphs[0].text = "公司机密"
    base.add_paragraph("旧正文，应当被清空")
    base.save(str(ws_dir / "模板.docx"))

    completed = subprocess.run(
        [sys.executable, str(script)],
        input=json.dumps(
            {
                "workspace": str(ws_dir),
                "output_dir": str(ws_dir),
                "params": {
                    "title": "新报告",
                    "base": str(ws_dir / "模板.docx"),
                    "markdown": "# 新报告标题\n\n生成的新内容。",
                },
            },
            ensure_ascii=False,
        ),
        capture_output=True,
        text=True,
        timeout=60,
        check=False,
    )
    assert completed.returncode == 0, completed.stderr[-2000:]
    output = json.loads(completed.stdout)
    assert output["success"] is True and output["data"]["base"] is True

    produced = Document(output["filepath"])
    texts = [p.text for p in produced.paragraphs]
    assert not any("旧正文" in t for t in texts)  # 基底正文已清空
    assert any("生成的新内容" in t for t in texts)
    assert produced.styles["Normal"].font.name == "Courier New"  # 样式定义保留
    assert produced.sections[0].header.paragraphs[0].text == "公司机密"  # 页眉保留
    # 新内容 run 字体派生自基底 Normal（Courier New），未被 PingFang 覆盖
    body_run = next(p.runs[0] for p in produced.paragraphs if p.runs and "生成的新内容" in p.text)
    assert body_run.font.name == "Courier New"


def test_word_base_missing_file_fails_cleanly(tmp_path):
    script = Path(__file__).resolve().parents[1] / "app" / "tools" / "scripts" / "word_document.py"
    ws_dir = tmp_path / "ws"
    ws_dir.mkdir()
    completed = subprocess.run(
        [sys.executable, str(script)],
        input=json.dumps(
            {
                "workspace": str(ws_dir),
                "output_dir": str(ws_dir),
                "params": {"title": "x", "base": str(ws_dir / "不存在.docx")},
            },
            ensure_ascii=False,
        ),
        capture_output=True,
        text=True,
        timeout=60,
        check=False,
    )
    assert completed.returncode == 0
    output = json.loads(completed.stdout)
    assert output["success"] is False and "不存在" in output["message"]
