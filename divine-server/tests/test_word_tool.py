"""Tests for word_document tool: rich inline text, element types, page structure, markdown."""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from docx import Document

SCRIPT = Path(__file__).resolve().parents[1] / "app" / "tools" / "scripts" / "word_document.py"


def _run(workspace: Path, params: dict) -> dict:
    payload = {
        "workspace": str(workspace.resolve()),
        "output_dir": str(workspace.resolve() / "out"),
        "params": params,
        "locale": "zh-CN",
    }
    completed = subprocess.run(
        [sys.executable, str(SCRIPT)],
        input=json.dumps(payload, ensure_ascii=False),
        capture_output=True,
        text=True,
        timeout=60,
        check=False,
    )
    assert completed.returncode == 0, completed.stderr
    return json.loads(completed.stdout)


def _load(out: dict) -> Document:
    return Document(out["filepath"])


def test_inline_format_and_tolerant_body(tmp_path: Path) -> None:
    ws_dir = tmp_path / "ws"
    ws_dir.mkdir()
    out = _run(
        ws_dir,
        {
            "title": "行内格式测试",
            "sections": [
                {
                    "heading": "进展",
                    "body": (
                        "完成了 **核心模块** 与 `docx 渲染`，详见 [文档](https://example.com/a)。\n"
                        "- 第一项\n"
                        "- 第二项\n"
                        "1. 步骤一\n"
                        "2. 步骤二\n"
                        "> 引用一句话"
                    ),
                }
            ],
        },
    )
    assert out["success"] is True
    assert out["data"]["warnings"] == []
    doc = _load(out)
    styles = [p.style.name for p in doc.paragraphs]
    assert "List Bullet" in styles
    assert "List Number" in styles
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "**" not in all_text  # markdown 标记不再原样漏出
    assert "[" not in all_text or "文档" in all_text
    bold_runs = [
        run
        for p in doc.paragraphs
        for run in p.runs
        if run.font.bold and run.text == "核心模块"
    ]
    assert bold_runs, "加粗 run 未生成"


def test_table_section_with_header_shading(tmp_path: Path) -> None:
    ws_dir = tmp_path / "ws"
    ws_dir.mkdir()
    out = _run(
        ws_dir,
        {
            "title": "会议纪要",
            "sections": [
                {
                    "type": "table",
                    "caption": "决议事项",
                    "headers": ["事项", "负责人", "截止"],
                    "rows": [
                        ["接入文件问答", "张三", "2026-09-10"],
                        ["RAG 检索联调", "李四", "2026-09-20"],
                    ],
                }
            ],
        },
    )
    assert out["success"] is True
    assert out["data"]["tables"] == 1
    doc = _load(out)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3
    assert table.rows[0].cells[0].text == "事项"
    header_xml = table.rows[0].cells[0]._tc.xml
    assert "w:shd" in header_xml and "2563EB" in header_xml  # 表头着色


def test_cover_footer_toc(tmp_path: Path) -> None:
    ws_dir = tmp_path / "ws"
    ws_dir.mkdir()
    out = _run(
        ws_dir,
        {
            "title": "正式报告",
            "cover": True,
            "toc": True,
            "sections": [{"heading": "背景", "body": "项目背景说明。"}],
        },
    )
    assert out["success"] is True
    doc = _load(out)
    body_xml = doc.element.body.xml
    # 真封面：独立分页 + 首页页脚关闭
    assert 'w:type="page"' in body_xml
    assert doc.sections[0].different_first_page_header_footer is True
    # 真页脚：每页 PAGE/NUMPAGES 域
    footer_xml = doc.sections[0].footer.paragraphs[0]._p.xml
    assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml
    # 目录域
    assert "TOC" in body_xml


def test_markdown_param_conversion(tmp_path: Path) -> None:
    ws_dir = tmp_path / "ws"
    ws_dir.mkdir()
    out = _run(
        ws_dir,
        {
            "title": "Markdown 直转",
            "markdown": (
                "# 项目周报\n\n"
                "本周完成 **核心功能**，重点见 [计划](https://example.com/p)。\n\n"
                "- 模块 A\n"
                "- 模块 B\n\n"
                "1. 联调\n"
                "2. 上线\n\n"
                "| 指标 | 数值 |\n"
                "| --- | --- |\n"
                "| 覆盖率 | 85% |\n\n"
                "```python\nprint('hello')\n```\n\n"
                "> 周报结语\n"
            ),
        },
    )
    assert out["success"] is True
    data = out["data"]
    assert data["tables"] == 1
    assert data["code_blocks"] == 1
    assert data["headings"] >= 1
    assert data["lists"] >= 4
    assert data["quotes"] == 1
    doc = _load(out)
    assert len(doc.tables) == 1
    styles = [p.style.name for p in doc.paragraphs]
    assert "Heading 1" in styles
    assert "List Bullet" in styles and "List Number" in styles


def test_quote_code_and_missing_image_warning(tmp_path: Path) -> None:
    ws_dir = tmp_path / "ws"
    ws_dir.mkdir()
    out = _run(
        ws_dir,
        {
            "title": "混合元素",
            "sections": [
                {"type": "quote", "content": "重要提示：请及时更新依赖。"},
                {"type": "code_block", "content": "pip install docx"},
                {"type": "image", "src": "not_exists/missing.png"},
                {"type": "page_break"},
                {"heading": "附录", "body": "补充说明。"},
            ],
        },
    )
    assert out["success"] is True
    data = out["data"]
    assert data["quotes"] == 1
    assert data["code_blocks"] == 1
    assert data["images"] == 0
    assert any("图片不可用" in warning for warning in data["warnings"])
    doc = _load(out)
    code_runs = [
        run
        for p in doc.paragraphs
        for run in p.runs
        if run.text == "pip install docx"
    ]
    assert code_runs and code_runs[0].font.name == "Consolas"
