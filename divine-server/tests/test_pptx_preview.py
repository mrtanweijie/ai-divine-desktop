"""pptx -> PNG 预览验证：核心渲染 / API 端点守卫 / agent 工具脚本端到端。"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest
from fastapi import HTTPException
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

SCRIPT = Path(__file__).resolve().parents[1] / "app" / "tools" / "scripts" / "pptx_to_png.py"


def _make_deck(path: Path, slides: int = 3) -> Path:
    """造一个最小的多页 pptx：文本框 + 圆角矩形（覆盖两种绘制路径）。"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for index in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        box.text_frame.text = f"第 {index + 1} 页"
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(3), Inches(4), Inches(2)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0xFF, 0x5C, 0x93)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


# ---------------------------------------------------------------- 核心渲染


def test_render_pptx_pages_returns_data_urls(tmp_path):
    from app.docgen.ppt.pptx_to_png import render_pptx_pages

    deck = _make_deck(tmp_path / "deck.pptx", slides=3)
    pages = render_pptx_pages(deck, scale=1.0)
    assert len(pages) == 3
    assert all(page.startswith("data:image/png;base64,") for page in pages)
    # data URL 体积合理（不为空，也不至于异常巨大）
    assert all(1_000 < len(page) < 8_000_000 for page in pages)


def test_render_pptx_pages_caps_page_count(tmp_path):
    from app.docgen.ppt.pptx_to_png import MAX_PPTX_PREVIEW_PAGES, render_pptx_pages

    deck = _make_deck(tmp_path / "deck.pptx", slides=2)
    pages = render_pptx_pages(deck, max_pages=1, scale=1.0)
    assert len(pages) == 1
    assert MAX_PPTX_PREVIEW_PAGES >= 1


def test_render_pptx_page_images_page_filter(tmp_path):
    from app.docgen.ppt.pptx_to_png import render_pptx_page_images

    deck = _make_deck(tmp_path / "deck.pptx", slides=3)
    picked = render_pptx_page_images(deck, scale=1.0, page=2)
    assert [number for number, _ in picked] == [2]

    with pytest.raises(RuntimeError, match="页码超出范围"):
        render_pptx_page_images(deck, scale=1.0, page=9)


def test_render_pptx_rejects_legacy_ppt(tmp_path):
    from app.docgen.ppt.pptx_to_png import render_pptx_pages

    legacy = tmp_path / "old.ppt"
    legacy.write_bytes(b"not a real ppt")
    with pytest.raises(RuntimeError, match="只支持"):
        render_pptx_pages(legacy)


def test_render_pptx_rejects_corrupt_file(tmp_path):
    from app.docgen.ppt.pptx_to_png import render_pptx_pages

    broken = tmp_path / "broken.pptx"
    broken.write_bytes(b"PK\x03\x04 garbage")
    with pytest.raises(RuntimeError, match="损坏或格式不支持"):
        render_pptx_pages(broken)


def test_render_pptx_uses_deck_dimensions(tmp_path):
    """4:3 演示文稿应按实际尺寸渲染，而不是写死 16:9。"""
    from PIL import Image

    from app.docgen.ppt.pptx_to_png import render_pptx_page_images

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "四比三"
    path = tmp_path / "ratio.pptx"
    prs.save(path)

    images = render_pptx_page_images(path, scale=1.0)
    assert len(images) == 1
    size = images[0][1].size
    assert size == Image.new("RGB", (int(10 * 96), int(7.5 * 96))).size


# ---------------------------------------------------------------- API 端点


def test_pptx_preview_endpoint_guards(tmp_path):
    from app.api.files import PptxPreviewRequest, pptx_preview

    deck = _make_deck(tmp_path / "deck.pptx", slides=1)

    result = pptx_preview(PptxPreviewRequest(path=str(deck)))
    assert result["page_count"] == 1
    assert result["pages"][0].startswith("data:image/png;base64,")

    with pytest.raises(HTTPException) as not_found:
        pptx_preview(PptxPreviewRequest(path=str(tmp_path / "missing.pptx")))
    assert not_found.value.status_code == 404

    legacy = tmp_path / "old.ppt"
    legacy.write_bytes(b"x")
    with pytest.raises(HTTPException) as bad_suffix:
        pptx_preview(PptxPreviewRequest(path=str(legacy)))
    assert bad_suffix.value.status_code == 400


# ------------------------------------------------------- agent 工具脚本


def _run_script(payload: dict, tmp_path: Path) -> dict:
    (tmp_path / "workspace").mkdir(parents=True, exist_ok=True)
    (tmp_path / "output").mkdir(parents=True, exist_ok=True)
    completed = subprocess.run(
        [sys.executable, str(SCRIPT)],
        input=json.dumps(payload, ensure_ascii=False),
        capture_output=True, text=True, timeout=120, check=False,
    )
    assert completed.returncode == 0, completed.stderr[-2000:]
    return json.loads(completed.stdout)


def test_pptx_to_png_script_exports_all_pages(tmp_path):
    workspace = tmp_path / "workspace"
    deck = _make_deck(workspace / "deck.pptx", slides=3)
    result = _run_script(
        {
            "workspace": str(workspace),
            "output_dir": str(tmp_path / "output"),
            "params": {"file_path": "deck.pptx", "scale": 1.0},
        },
        tmp_path,
    )
    assert result["success"] is True, result
    pages = result["data"]["pages"]
    assert result["data"]["page_count"] == 3
    assert len(pages) == 3
    for page in pages:
        assert Path(page).is_file() and Path(page).stat().st_size > 0
    assert result["data"]["pages"][0].endswith("deck-01.png")


def test_pptx_to_png_script_single_page_and_custom_dir(tmp_path):
    workspace = tmp_path / "workspace"
    deck = _make_deck(workspace / "deck.pptx", slides=3)
    result = _run_script(
        {
            "workspace": str(workspace),
            "output_dir": str(tmp_path / "output"),
            "params": {"file_path": "deck.pptx", "page": 2, "scale": 1.0,
                       "output_path": "exports/slides"},
        },
        tmp_path,
    )
    assert result["success"] is True, result
    assert result["data"]["page_count"] == 1
    assert Path(result["data"]["pages"][0]).parent.name == "slides"


def test_pptx_to_png_script_rejects_outside_path(tmp_path):
    _make_deck(tmp_path / "outside.pptx", slides=1)
    result = _run_script(
        {
            "workspace": str(tmp_path / "workspace"),
            "output_dir": str(tmp_path / "output"),
            "params": {"file_path": str(tmp_path / "outside.pptx")},
        },
        tmp_path,
    )
    # 沙箱外文件默认拒绝（越界），除非走 external_paths 白名单
    assert result["success"] is False
    assert "越界" in result["message"]


def test_pptx_to_png_script_rejects_bad_scale(tmp_path):
    workspace = tmp_path / "workspace"
    _make_deck(workspace / "deck.pptx", slides=1)
    result = _run_script(
        {
            "workspace": str(workspace),
            "output_dir": str(tmp_path / "output"),
            "params": {"file_path": "deck.pptx", "scale": 9},
        },
        tmp_path,
    )
    assert result["success"] is False
    assert "scale" in result["message"]
