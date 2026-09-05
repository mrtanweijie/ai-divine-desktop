"""PPT 空页修复回归（docs/fix/001-ppt-empty-slides-20260906.md）：

核心 fixture 是霸王茶姬事故的真实模型输入（tool_runs 2026-09-01，
tests/fixtures/ppt_chagee_input.json）——归一化前 cards/kpi/quote 内容
全部静默丢失，归一化后必须完整进入渲染窄键。
"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest  # noqa: E402
from pptx import Presentation  # noqa: E402

from app.docgen.ppt.normalize import (  # noqa: E402
    check_empty_pages,
    normalize_slide,
    normalize_slides,
)

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "ppt_chagee_input.json"


def _load_fixture() -> dict:
    return json.loads(FIXTURE.read_text(encoding="utf-8"))


def test_normalize_chagee_kpi_items_reach_points():
    """事故页 slide2：kpi items[{label,value,desc}×6] → points 数值|标签|说明。"""
    slides = _load_fixture()["slides"]
    normalized, warnings = normalize_slides(slides)
    kpi = next(s for s in normalized if s.get("layout") == "kpi")
    points = kpi["points"]
    assert len(points) == 6
    assert any("295亿" in p and "GMV" in p for p in points)
    assert any("124亿" in p and "营业收入" in p for p in points)
    assert any(w for w in warnings if "items" in w)


def test_normalize_chagee_cards_items_reach_points():
    """事故页 slide4/8/11：cards items[{title,desc}] → points 标题|描述。"""
    slides = _load_fixture()["slides"]
    normalized, _ = normalize_slides(slides)
    cards = [s for s in normalized if s.get("layout") == "cards"]
    assert len(cards) == 3
    for slide in cards:
        assert isinstance(slide.get("points"), list) and len(slide["points"]) >= 3
        assert all("|" in p for p in slide["points"])  # 标题|描述 结构
    first = cards[0]["points"]
    assert any(p.startswith("品类减法") for p in first)


def test_normalize_chagee_quote_key_reaches_heading():
    """事故页 slide13：quote 键正文（"健康的增长…"）此前整句丢失。"""
    slides = _load_fixture()["slides"]
    normalized, _ = normalize_slides(slides)
    quote = next(s for s in normalized if s.get("layout") == "quote")
    assert "健康的增长" in str(quote.get("heading"))
    assert quote.get("subtitle") == "霸王姬的商业启示".replace("姬", "茶姬")


def test_normalize_chagee_bullets_compare_untouched():
    """本来就正常的 bullets/compare 页归一化后保持不变。"""
    slides = _load_fixture()["slides"]
    normalized, _ = normalize_slides(slides)
    bullets = [s for s in normalized if s.get("layout") == "bullets"]
    assert len(bullets) == 2
    assert any("云南茶马古道" in p for p in bullets[0]["points"])
    compare = next(s for s in normalized if s.get("layout") == "compare")
    assert compare["left"]["title"] == "霸王茶姬·原叶鲜奶茶"
    assert len(compare["right"]["points"]) == 4


def test_normalize_compare_side_bullets_alias():
    item = {
        "layout": "compare",
        "heading": "对比",
        "left": {"title": "旧", "bullets": ["a", "b"]},
        "right": {"title": "新", "items": [{"title": "x", "desc": "y"}]},
    }
    normalized, warnings = normalize_slide(item)
    assert normalized["left"]["points"] == ["a", "b"]
    assert normalized["right"]["points"] == ["x|y"]
    assert len(warnings) >= 2


def test_normalize_unknown_key_reported():
    item = {"layout": "bullets", "heading": "T", "points": ["a"], "theme_color": "red"}
    _, warnings = normalize_slide(item)
    assert any("theme_color" in w for w in warnings)


def test_empty_page_detection_flags_content_layouts():
    layouts = ["cover", "kpi", "section", "bullets"]
    items = [
        {"layout": "cover", "heading": "封面"},
        {"layout": "kpi", "heading": "指标"},  # kpi 无 points → 空
        {"layout": "section", "heading": "01 章节"},  # section 豁免
        {"layout": "bullets", "heading": "要点", "points": ["内容充足的要点页"]},
    ]
    warnings = check_empty_pages(items, layouts)
    assert len(warnings) == 1
    assert "第 2 页" in warnings[0] and "kpi" in warnings[0]


def test_speaker_note_lands_in_notes_slide():
    from app.docgen.ppt.renderer import get_ppt_theme
    from app.docgen.ppt.renderer.layouts import add_slide

    prs = Presentation()
    prs.slide_width = 12192000
    prs.slide_height = 6858000
    get_ppt_theme("universal")
    layout = add_slide(
        prs,
        {"layout": "bullets", "heading": "标题", "points": ["要点"], "note": "口播提醒"},
        "universal",
        1,
    )
    assert layout == "bullets"
    assert prs.slides[0].notes_slide.notes_text_frame.text == "口播提醒"


def test_end_to_end_chagee_fixture_no_empty_pages(tmp_path):
    """端到端回归：真实事故输入 → 脚本渲染 → 重开逐页断言无空页。"""
    params = dict(_load_fixture())
    payload = {
        "workspace": str(tmp_path),
        "output_dir": str(tmp_path),
        "params": params,
    }
    script = Path(__file__).resolve().parents[1] / "app" / "tools" / "scripts" / "presentation.py"
    completed = subprocess.run(
        [sys.executable, str(script)],
        input=json.dumps(payload, ensure_ascii=False),
        capture_output=True,
        text=True,
        timeout=120,
        check=False,
    )
    assert completed.returncode == 0, completed.stderr[-2000:]
    output = json.loads(completed.stdout)
    assert output["success"] is True

    # 重开产物：内容版式页必须都有正文（此前 kpi/cards/quote 页全空）
    prs = Presentation(output["filepath"])
    content_layouts = {"bullets", "cards", "timeline", "columns", "two_col",
                       "compare", "kpi", "quote", "table", "chart"}
    empty = []
    for index, slide in enumerate(prs.slides, 1):
        if index > len(params["slides"]) + 1:  # 自动补齐的 end 页豁免
            break
        chars = sum(
            len(shape.text_frame.text.strip())
            for shape in slide.shapes
            if shape.has_text_frame
        )
        if index == 1:
            continue  # 封面豁免
        if chars < 12:
            empty.append(index)
    assert empty == [], f"空页: {empty}; warnings={output['data'].get('warnings')}"

    # 消息不再对空页视而不见
    assert output["message"].startswith("PPT 生成完成")
    warnings = output["data"]["warnings"]
    assert any("items" in w for w in warnings)  # 归一化提示可见


def test_script_path_with_pptx_suffix_treated_as_directory(tmp_path):
    """path 带 .pptx 文件名时按目录处理，不再建出名叫 xxx.pptx 的目录。"""
    params = {
        "title": "路径防御测试",
        "theme": "universal",
        "path": "presentations/路径防御测试.pptx",
        "slides": [{"layout": "bullets", "heading": "标题", "points": ["要点"]}],
    }
    script = Path(__file__).resolve().parents[1] / "app" / "tools" / "scripts" / "presentation.py"
    completed = subprocess.run(
        [sys.executable, str(script)],
        input=json.dumps({"workspace": str(tmp_path), "output_dir": str(tmp_path), "params": params}, ensure_ascii=False),
        capture_output=True,
        text=True,
        timeout=120,
        check=False,
    )
    assert completed.returncode == 0, completed.stderr[-2000:]
    output = json.loads(completed.stdout)
    assert output["success"] is True
    produced = Path(output["filepath"])
    assert produced.suffix == ".pptx" and produced.is_file()
    assert produced.parent == tmp_path / "presentations"  # 不再嵌套
    assert not (tmp_path / "presentations" / "路径防御测试.pptx").is_dir()
