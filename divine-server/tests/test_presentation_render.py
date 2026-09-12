"""presentation 脚本渲染验证：14 种版式 / 原生图表 / 表格 / 图片 / 兜底逻辑。"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

SCRIPT = Path(__file__).resolve().parents[1] / "app" / "tools" / "scripts" / "presentation.py"


def _run(payload: dict, tmp_path: Path) -> dict:
    body = {
        "workspace": str(tmp_path / "workspace"),
        "output_dir": str(tmp_path / "output"),
        "params": payload,
    }
    (tmp_path / "workspace").mkdir(parents=True, exist_ok=True)
    (tmp_path / "output").mkdir(parents=True, exist_ok=True)
    completed = subprocess.run(
        [sys.executable, str(SCRIPT)],
        input=json.dumps(body, ensure_ascii=False),
        capture_output=True, text=True, timeout=90, check=False,
    )
    assert completed.returncode == 0, completed.stderr[-2000:]
    return json.loads(completed.stdout)


def _all_layout_slides() -> list[dict]:
    return [
        {"layout": "toc", "points": ["市场洞察", "增长引擎", "产品路线",
                                      "组织保障", "财务规划", "风险与应对"]},
        {"layout": "section", "heading": "产品路线"},
        {"layout": "two_col", "heading": "智能体平台 2.0",
         "points": ["多模态理解升级", "零代码编排引擎", "内置行业模板"]},
        {"layout": "cards", "heading": "三大增长引擎",
         "points": ["企业客户|深耕 KA 标杆", "生态合作|覆盖 200 城", "智能提效|人效 +40%"]},
        {"layout": "timeline", "heading": "2026 关键里程碑",
         "points": ["Q1|联盟成立", "Q2|扩编完成", "Q3|新品发布", "Q4|年度复盘"]},
        {"layout": "columns", "heading": "三维度展开",
         "points": ["产品力", "渠道力", "运营力"]},
        {"layout": "compare", "heading": "增长路径对比",
         "left": {"title": "自建团队", "points": ["可控性强", "周期 6 个月"]},
         "right": {"title": "生态合作", "points": ["起量快", "周期 2 个月"]}},
        {"layout": "kpi", "heading": "上半年核心指标",
         "points": ["1,668|营收（万元）|同比 +38%", "1,240|客户数|+52%",
                    "91%|复购率|+6pt", "68|NPS|+12"]},
        {"layout": "quote", "heading": "增长的本质是构建可复用的价值网络",
         "subtitle": "2026 增长战略纲要"},
        {"layout": "table", "heading": "区域销售明细",
         "table": {"columns": ["区域", "营收", "同比"],
                   "rows": [["华东", "634", "+45%"], ["华北", "428", "+32%"]]}},
        {"layout": "chart", "heading": "季度营收趋势",
         "chart": {"type": "column", "categories": ["Q1", "Q2", "Q3", "Q4"],
                   "series": [{"name": "营收", "values": [312, 444, 386, 528]}]},
         "points": ["Q2 环比 +42%", "Q4 冲刺 528 万"]},
        {"layout": "bullets", "heading": "风险与应对",
         "points": ["市场波动：多元化布局", "供给不足：提前锁产能"]},
        {"layout": "end", "heading": "谢谢观看", "subtitle": "让每一页都有设计感"},
    ]


def test_all_layouts_render(tmp_path):
    result = _run({"title": "2026 增长战略规划", "theme": "universal",
                   "slides": _all_layout_slides()}, tmp_path)
    assert result["success"] is True, result
    output = Path(result["filepath"])
    assert output.is_file() and output.stat().st_size > 0

    prs = Presentation(str(output))
    # 封面 + 12 内容页 + 结尾页
    assert len(prs.slides) == 14

    all_shapes = [shape for slide in prs.slides for shape in slide.shapes]
    assert any(shape.has_chart for shape in all_shapes), "原生图表缺失"
    assert any(shape.has_table for shape in all_shapes), "表格缺失"


def test_chart_slide_has_editable_chart(tmp_path):
    result = _run({"title": "图表验证",
                   "slides": [{"layout": "chart", "heading": "趋势",
                               "chart": {"type": "pie",
                                         "categories": ["华东", "华北", "华南"],
                                         "series": [{"name": "占比",
                                                     "values": [38, 26, 21]}]}}]},
                  tmp_path)
    assert result["success"] is True, result
    prs = Presentation(result["filepath"])
    charts = [shape.chart for slide in prs.slides for shape in slide.shapes
              if shape.has_chart]
    assert len(charts) == 1
    plot = charts[0].plots[0]
    assert len(list(plot.categories)) == 3
    assert len(list(charts[0].series)) == 1


def test_chart_unsupported_type_fails_gracefully(tmp_path):
    result = _run({"title": "非法图表",
                   "slides": [{"layout": "chart", "heading": "桑基图",
                               "chart": {"type": "sankey", "categories": ["A"],
                                         "series": [{"name": "s", "values": [1]}]}}]},
                  tmp_path)
    assert result["success"] is False
    assert "generate_chart" in result["message"]


def test_chart_over_limit_fails_gracefully(tmp_path):
    values = list(range(20))
    result = _run({"title": "超限图表",
                   "slides": [{"layout": "chart", "heading": "过多数据点",
                               "chart": {"type": "bar",
                                         "categories": [f"c{i}" for i in values],
                                         "series": [{"name": "s",
                                                     "values": values}]}}]},
                  tmp_path)
    assert result["success"] is False
    assert "聚合" in result["message"]


def test_two_col_with_local_image(tmp_path):
    workspace = tmp_path / "workspace"
    workspace.mkdir(parents=True, exist_ok=True)
    from PIL import Image

    image_path = workspace / "product.png"
    Image.new("RGB", (800, 600), color=(37, 99, 235)).save(image_path)
    result = _run({"title": "图文页",
                   "slides": [{"layout": "two_col", "heading": "产品发布",
                               "points": ["要点一"],
                               "image": {"path": "product.png",
                                         "caption": "产品主视觉"}}]},
                  tmp_path)
    assert result["success"] is True, result
    prs = Presentation(result["filepath"])
    pictures = [shape for slide in prs.slides for shape in slide.shapes
                if shape.shape_type is not None and str(shape.shape_type).startswith("PICTURE")]
    assert pictures, "本地图片未插入"


def test_svg_image_converted_to_picture(tmp_path):
    workspace = tmp_path / "workspace"
    workspace.mkdir(parents=True, exist_ok=True)
    (workspace / "logo.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="200" height="120">'
        '<rect width="200" height="120" fill="#7C3AED"/></svg>',
        encoding="utf-8",
    )
    result = _run({"title": "SVG 配图",
                   "slides": [{"layout": "two_col", "heading": "品牌页",
                               "points": ["品牌资产"],
                               "image": {"path": "logo.svg"}}]},
                  tmp_path)
    assert result["success"] is True, result
    prs = Presentation(result["filepath"])
    pictures = [shape for slide in prs.slides for shape in slide.shapes
                if shape.shape_type is not None and str(shape.shape_type).startswith("PICTURE")]
    assert pictures, "SVG 未转 PNG 插入"


def test_image_outside_workspace_falls_back_to_placeholder(tmp_path):
    outside = tmp_path / "outside.png"
    from PIL import Image

    Image.new("RGB", (100, 100), color=(0, 0, 0)).save(outside)
    result = _run({"title": "越界图片",
                   "slides": [{"layout": "two_col", "heading": "安全校验",
                               "points": ["要点"],
                               "image": {"path": str(outside)}}]},
                  tmp_path)
    # 越界路径不应导致失败，降级为占位块
    assert result["success"] is True, result
    prs = Presentation(result["filepath"])
    pictures = [shape for slide in prs.slides for shape in slide.shapes
                if shape.shape_type is not None and str(shape.shape_type).startswith("PICTURE")]
    assert not pictures, "越界图片不应被插入"


def test_auto_end_slide(tmp_path):
    result = _run({"title": "自动结尾",
                   "slides": [{"layout": "bullets", "heading": "要点",
                               "points": ["一", "二"]}]}, tmp_path)
    assert result["success"] is True
    assert result["data"]["slides"] == 3  # 封面 + 内容 + 自动结尾

    result2 = _run({"title": "显式结尾",
                    "slides": [{"layout": "bullets", "heading": "要点",
                                "points": ["一"]},
                               {"layout": "end", "heading": "再见"}]}, tmp_path)
    assert result2["success"] is True
    assert result2["data"]["slides"] == 3  # 封面 + 内容 + 显式结尾，不重复


def test_all_design_presets_render(tmp_path):
    """8 套现代设计预设全版式冒烟：每种风格生成完整 deck 不报错。"""
    import pytest

    sample = [
        {"layout": "toc", "points": ["市场洞察", "增长引擎", "产品路线"]},
        {"layout": "section", "heading": "产品路线"},
        {"layout": "two_col", "heading": "智能体平台 2.0",
         "points": ["多模态理解升级", "零代码编排"]},
        {"layout": "cards", "heading": "增长引擎",
         "points": ["企业客户|KA 标杆", "生态合作|200 城", "智能提效|+40%"]},
        {"layout": "timeline", "heading": "里程碑",
         "points": ["Q1|联盟成立", "Q2|扩编完成", "Q3|新品发布", "Q4|年度复盘"]},
        {"layout": "kpi", "heading": "核心指标",
         "points": ["1,668|营收（万元）|+38%", "91%|复购率|+6pt"]},
        {"layout": "chart", "heading": "季度营收",
         "chart": {"type": "column", "categories": ["Q1", "Q2"],
                   "series": [{"name": "营收", "values": [312, 444]}]}},
        {"layout": "table", "heading": "区域明细",
         "table": {"columns": ["区域", "营收"], "rows": [["华东", "634"]]}},
        {"layout": "compare", "heading": "路径对比",
         "left": {"title": "自建", "points": ["可控"]},
         "right": {"title": "合作", "points": ["起量快"]}},
        {"layout": "quote", "heading": "构建可复用的价值网络", "subtitle": "战略纲要"},
        {"layout": "bullets", "heading": "风险与应对", "points": ["多元化布局"]},
    ]
    for theme in ("aurora", "glass", "brutalist", "swiss", "mesh",
                  "editorial", "darktech", "pop"):
        result = _run({"title": "预设验证", "theme": theme, "slides": sample},
                      tmp_path)
        assert result["success"] is True, f"{theme}: {result}"
        assert result["data"]["slides"] == 13, theme  # 封面 + 11 + 自动结尾


def test_unknown_theme_falls_back(tmp_path):
    result = _run({"title": "回退验证", "theme": "not-exist",
                   "slides": [{"layout": "bullets", "heading": "要点",
                               "points": ["一"]}]}, tmp_path)
    assert result["success"] is True
    assert result["data"]["theme"] == "not-exist"  # 原样返回，渲染回退默认


def test_textfit_estimation():
    from app.docgen.ppt.renderer.textfit import pick_size, text_em_width, wrapped_lines

    assert text_em_width("中文") == 2.0
    assert text_em_width("ab") == 1.1
    assert wrapped_lines("一" * 20, 5.0, 20) == 2  # 每行 18 字，20 字折 2 行
    assert wrapped_lines("hello world", 5.0, 20) == 1
    long_text = "要点" * 30
    short_text = "短要点"
    assert pick_size(long_text, 6.0, 0.8, [20, 18, 16, 14]) < \
        pick_size(short_text, 6.0, 0.8, [20, 18, 16, 14])


def test_guess_layout():
    from app.docgen.ppt.renderer.layouts import guess_layout

    assert guess_layout({"heading": "目录"}) == "toc"
    assert guess_layout({"heading": "2026 路线图"}) == "timeline"
    assert guess_layout({"heading": "引擎", "points": list("abcde")}) == "cards"
    assert guess_layout({"heading": "引擎", "points": list("abc")}) == "columns"
    assert guess_layout({"heading": "总结", "body": "长文本"}) == "bullets"


def test_overloaded_bullets_upgrade_to_cards():
    """超密 bullets（≥7 条）自动升级 cards 渲染。"""
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    from pptx import Presentation as NewPres

    from app.docgen.ppt.renderer.layouts import add_slide

    prs = NewPres()
    item = {"layout": "bullets", "heading": "融资历程",
            "points": [f"轮次{i}|金额{i}万" for i in range(1, 9)]}
    used = add_slide(prs, item, "aurora", 2)
    assert used == "cards"
    # 卡片面板存在（圆角矩形）
    rounded = [sh for sh in prs.slides[0].shapes
               if "ROUND" in str(sh.shape_type).upper()
               or (sh.shape_type is not None and "AUTO_SHAPE" in str(sh.shape_type))]
    assert rounded, "cards 面板缺失"

    prs2 = NewPres()
    used2 = add_slide(prs2, {"layout": "bullets", "heading": "正常页",
                             "points": ["一", "二", "三"]}, "aurora", 2)
    assert used2 == "bullets"


def test_dict_points_normalized():
    """LLM 违反 schema 传 dict 数组时按字段语义渲染，不泄漏 dict 原文。"""
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    from pptx import Presentation as NewPres

    from app.docgen.ppt.renderer.layouts import add_slide

    # timeline：{time, desc} → "2015|Seed 种子轮"
    prs = NewPres()
    used = add_slide(prs, {
        "layout": "timeline", "heading": "融资历程",
        "points": [
            {"time": "2015", "desc": "Seed 种子轮 · 4:33 Creative Lab"},
            {"time": "2016", "desc": "两轮 Early VC"},
            {"time": "2018", "desc": "Early VC · SBVA"},
            {"time": "2019", "desc": "Late VC · 约 1,000 万美元"},
        ],
    }, "aurora", 3)
    assert used == "timeline"
    all_text = "\n".join(sh.text_frame.text for sh in prs.slides[0].shapes
                         if sh.has_text_frame)
    assert "{'" not in all_text, "dict 原文泄漏"
    assert "2015" in all_text and "Seed 种子轮" in all_text

    # kpi：{value, label, delta}
    prs2 = NewPres()
    add_slide(prs2, {
        "layout": "kpi", "heading": "核心指标",
        "points": [{"value": "4,000 万", "label": "月活用户", "delta": "+18%"}],
    }, "aurora", 2)
    text2 = "\n".join(sh.text_frame.text for sh in prs2.slides[0].shapes
                      if sh.has_text_frame)
    assert "{'" not in text2 and "4,000 万" in text2 and "+18%" in text2

    # cards：{title, desc}
    prs3 = NewPres()
    add_slide(prs3, {
        "layout": "cards", "heading": "优势",
        "points": [{"title": "数据壁垒", "desc": "17 款游戏全覆盖"}],
    }, "swiss", 2)
    text3 = "\n".join(sh.text_frame.text for sh in prs3.slides[0].shapes
                      if sh.has_text_frame)
    assert "{'" not in text3 and "数据壁垒" in text3 and "17 款游戏" in text3


def test_auto_toc_from_sections(tmp_path):
    """未传目录页时按章节页自动生成（条目带页码）。"""
    slides = [
        {"layout": "section", "heading": "市场洞察"},
        {"layout": "bullets", "heading": "格局", "points": ["一", "二"]},
        {"layout": "section", "heading": "增长引擎"},
        {"layout": "bullets", "heading": "引擎详情", "points": ["一"]},
        {"layout": "section", "heading": "风险与应对"},
        {"layout": "bullets", "heading": "风险", "points": ["一"]},
    ]
    result = _run({"title": "自动目录", "theme": "aurora", "slides": slides},
                  tmp_path)
    assert result["success"] is True, result
    assert result["data"]["toc"] is True
    assert result["data"]["slides"] == 9  # 封面 + 目录 + 6 内容 + 结尾

    prs = Presentation(result["filepath"])
    toc_slide = prs.slides[1]  # 第 2 页应为目录
    texts = "\n".join(sh.text_frame.text for sh in toc_slide.shapes
                      if sh.has_text_frame)
    assert "目录" in texts
    assert "市场洞察" in texts and "增长引擎" in texts and "风险与应对" in texts
    assert "P03" in texts  # 第一个 section 在第 3 页


def test_toc_disabled_and_explicit(tmp_path):
    slides = [
        {"layout": "section", "heading": "章节一"},
        {"layout": "bullets", "heading": "内容", "points": ["一"]},
        {"layout": "section", "heading": "章节二"},
        {"layout": "bullets", "heading": "内容二", "points": ["一"]},
        {"layout": "section", "heading": "章节三"},
        {"layout": "bullets", "heading": "内容三", "points": ["一"]},
    ]
    # 显式关闭
    r1 = _run({"title": "无目录", "theme": "aurora", "toc": False,
               "slides": slides}, tmp_path)
    assert r1["success"] is True and r1["data"]["toc"] is False
    assert r1["data"]["slides"] == 8  # 封面 + 6 + 结尾

    # LLM 已自带目录 → 不重复插入
    slides_with_toc = [{"layout": "toc", "heading": "目录",
                        "points": ["章节一|P02", "章节二|P04"]}] + slides
    r2 = _run({"title": "自带目录", "theme": "aurora",
               "slides": slides_with_toc}, tmp_path)
    assert r2["success"] is True and r2["data"]["toc"] is False
    prs = Presentation(r2["filepath"])
    toc_pages = [i for i, slide in enumerate(prs.slides, 1)
                 if any(sh.has_text_frame and sh.text_frame.text.strip() == "目录"
                        for sh in slide.shapes)]
    assert len(toc_pages) == 1 and toc_pages[0] == 2  # 仅封面后一页


def test_explicit_cover_layout_not_duplicated(tmp_path):
    """回归：显式 layout=cover 只算一张封面，不再额外生成自动封面。

    修复前 _is_cover_slide 只认「无 layout 但带 subtitle」的隐式封面，
    显式声明会被当内容页保留 → 自动封面 + 显式封面并存。
    """
    result = _run({
        "title": "演示标题",
        "slides": [{"layout": "cover", "heading": "唯一封面", "subtitle": "副标题"},
                   {"layout": "bullets", "heading": "要点", "points": ["一"]}],
    }, tmp_path)
    assert result["success"] is True, result
    assert result["data"]["slides"] == 3  # 封面 + 内容 + 自动结尾

    prs = Presentation(result["filepath"])
    covers = [i for i, slide in enumerate(prs.slides, 1)
              if any(sh.has_text_frame and "唯一封面" in sh.text_frame.text
                     for sh in slide.shapes)]
    assert covers == [1], covers


def test_pop_style_primitives(tmp_path):
    """pop 预设（波普潮流风）：三色渐变下划条 / 页脚页码 / 页内进度点 / 封面礼盒。"""
    from lxml import etree

    result = _run({
        "title": "潮流玩具", "theme": "pop",
        "slides": [
            {"layout": "cover", "heading": "泡泡玛特", "subtitle": "潮玩帝国",
             "points": ["盲盒经济", "IP 飞轮", "全球化"]},
            {"layout": "cards", "heading": "公司速览",
             "points": ["2010|成立", "潮玩|品类"]},
        ],
    }, tmp_path)
    assert result["success"] is True, result
    assert result["data"]["slides"] == 3
    assert result["data"]["presets"] >= 17
    assert "pop" in result["data"]["styles"]

    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    prs = Presentation(result["filepath"])
    content = prs.slides[1]
    root = content._element

    # 1) 标题下的三色渐变下划条：单条 gradFill 含 ≥3 个色标
    multi = [gs for gs in root.findall(".//a:gradFill/a:gsLst", ns)
             if len(gs) >= 3]
    assert multi, "pop 风格内容页应有三色渐变下划条"

    # 2) 页脚页码：内容页（第 2 页 / 共 3 页）
    texts = "\n".join(sh.text_frame.text for sh in content.shapes
                      if sh.has_text_frame)
    assert "2 / 3" in texts

    # 3) 页内进度点：3 页 deck 应有 3 个进度点（含当前页的渐变胶囊）
    #    注：add_shape 产生的自选图形恒带 text_frame，故按几何特征识别
    dots = [sh for sh in content.shapes
            if abs(sh.top / 914400 - 7.09) < 0.03
            and abs(sh.height / 914400 - 0.075) < 0.01]
    assert len(dots) == 3, f"进度点数量异常: {len(dots)}"

    # 4) 封面装饰：盲盒礼盒的白色加号圆 + 四角星点缀
    cover = prs.slides[0]
    cover_xml = cover._element.xml
    assert "star4" in cover_xml or "star" in cover_xml.lower()
    whites = [sh for sh in cover.shapes
              if abs(sh.width / 914400 - 0.40) < 0.01
              and abs(sh.height / 914400 - 0.40) < 0.01
              and not sh.text_frame.text.strip()]
    assert whites, "封面应含礼盒的白色加号圆"
