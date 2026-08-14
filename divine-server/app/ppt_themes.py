from __future__ import annotations

from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


PPT_THEMES: dict[str, dict[str, str]] = {
    "universal": {
        "cover_bg": "10131A",
        "content_bg": "FFFFFF",
        "accent": "2563EB",
        "accent_light": "EFF6FF",
        "text": "1A1D24",
        "muted": "6B7280",
        "cover_text": "FFFFFF",
        "meta": "9AA3B2",
        "font": "PingFang SC",
    },
    "notion": {
        "cover_bg": "FFFFFF",
        "content_bg": "FFFFFF",
        "accent": "37352F",
        "accent_light": "F7F6F3",
        "text": "37352F",
        "muted": "9B9B97",
        "cover_text": "37352F",
        "meta": "B9B9B5",
        "font": "PingFang SC",
    },
    "linear": {
        "cover_bg": "202129",
        "content_bg": "FFFFFF",
        "accent": "5E6AD2",
        "accent_light": "F0F1FF",
        "text": "202129",
        "muted": "5C5F6E",
        "cover_text": "FFFFFF",
        "meta": "8F92A3",
        "font": "PingFang SC",
    },
    "apple": {
        "cover_bg": "F5F5F7",
        "content_bg": "FFFFFF",
        "accent": "0071E3",
        "accent_light": "E8F2FF",
        "text": "1D1D1F",
        "muted": "6E6E73",
        "cover_text": "1D1D1F",
        "meta": "AEAEB2",
        "font": "PingFang SC",
    },
    "modern": {
        "cover_bg": "10131A",
        "content_bg": "FFFFFF",
        "accent": "2563EB",
        "accent_light": "EFF6FF",
        "text": "1A1D24",
        "muted": "6B7280",
        "cover_text": "FFFFFF",
        "meta": "9AA3B2",
        "font": "PingFang SC",
    },
    "academic": {
        "cover_bg": "F7FAFC",
        "content_bg": "FFFFFF",
        "accent": "742A2A",
        "accent_light": "FEF5F5",
        "text": "1A202C",
        "muted": "718096",
        "cover_text": "1A202C",
        "meta": "A0AEC0",
        "font": "Noto Serif SC",
    },
    "travel": {
        "cover_bg": "FFF7ED",
        "content_bg": "FFFDF7",
        "accent": "EA580C",
        "accent_light": "FFF7ED",
        "text": "1C1917",
        "muted": "A8A29E",
        "cover_text": "9A3412",
        "meta": "FB923C",
        "font": "PingFang SC",
    },
    "roadmap": {
        "cover_bg": "1E1B4B",
        "content_bg": "FFFFFF",
        "accent": "7C3AED",
        "accent_light": "F3E8FF",
        "text": "1E1B4B",
        "muted": "A78BFA",
        "cover_text": "FFFFFF",
        "meta": "C4B5FD",
        "font": "PingFang SC",
    },
    "minimal": {
        "cover_bg": "FFFFFF",
        "content_bg": "FFFFFF",
        "accent": "222222",
        "accent_light": "F5F5F5",
        "text": "222222",
        "muted": "999999",
        "cover_text": "111111",
        "meta": "BBBBBB",
        "font": "PingFang SC",
    },
}

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def get_ppt_theme(theme_id: str) -> dict[str, str]:
    return PPT_THEMES.get(theme_id, PPT_THEMES["universal"])


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def solid(shape: Any, color: str) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    shape.line.fill.background()
    shape.shadow.inherit = False


def set_run(run: Any, size: int, color: str, bold: bool = False, font: str = "PingFang SC") -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    run.font.name = font
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.find(qn("a:ea"))
    if ea is None:
        ea = r_pr.makeelement(qn("a:ea"), {})
        r_pr.append(ea)
    ea.set("typeface", font)


def add_title_slide(
    prs: Presentation,
    title: str,
    count: int,
    theme_id: str,
    subtitle: str = "",
) -> None:
    theme = get_ppt_theme(theme_id)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    solid(bg, theme["cover_bg"])

    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.95),
        Inches(2.62),
        Inches(0.14),
        Inches(1.25),
    )
    solid(accent, theme["accent"])

    title_box = slide.shapes.add_textbox(
        Inches(1.35), Inches(2.55), Inches(10.6), Inches(1.6)
    )
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    set_run(p.runs[0], 46, theme["cover_text"], bold=True, font=theme["font"])

    meta_y = Inches(4.25)
    if subtitle:
        subtitle_box = slide.shapes.add_textbox(
            Inches(1.35), Inches(4.05), Inches(10.6), Inches(0.8)
        )
        tf = subtitle_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = subtitle
        set_run(p.runs[0], 16, theme["meta"], font=theme["font"])
        meta_y = Inches(5.0)

    meta = slide.shapes.add_textbox(
        Inches(1.35), meta_y, Inches(10), Inches(0.6)
    )
    tf = meta.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"AI Divine  ·  {count} 页"
    set_run(p.runs[0], 18, theme["meta"], font=theme["font"])

    footer = slide.shapes.add_textbox(
        Inches(1.35), Inches(6.75), Inches(10), Inches(0.4)
    )
    tf = footer.text_frame
    p = tf.paragraphs[0]
    p.text = "AI Divine"
    set_run(p.runs[0], 13, theme["muted"], font=theme["font"])


def add_content_slide(
    prs: Presentation,
    index: int,
    heading: str,
    body: str,
    theme_id: str,
) -> None:
    theme = get_ppt_theme(theme_id)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    solid(bg, theme["content_bg"])

    top = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.09)
    )
    solid(top, theme["accent"])

    label = slide.shapes.add_textbox(
        Inches(0.95), Inches(0.55), Inches(8), Inches(0.4)
    )
    tf = label.text_frame
    p = tf.paragraphs[0]
    p.text = f"{index:02d}"
    set_run(p.runs[0], 15, theme["accent"], bold=True, font=theme["font"])

    title_box = slide.shapes.add_textbox(
        Inches(0.95), Inches(1.0), Inches(11.4), Inches(1.0)
    )
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = heading
    set_run(p.runs[0], 34, theme["text"], bold=True, font=theme["font"])

    body_box = slide.shapes.add_textbox(
        Inches(0.95), Inches(2.25), Inches(11.4), Inches(4.4)
    )
    tf = body_box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, raw in enumerate(body.splitlines()):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        level = 1 if raw.startswith(("  ", "\t", "- ", "• ")) else 0
        text = raw.lstrip("-•· \t")
        p.text = f"{'• ' if level == 0 else '– '}{text}"
        set_run(
            p.runs[0],
            22 if level == 0 else 18,
            theme["text"],
            bold=(level == 0),
            font=theme["font"],
        )
        p.space_after = Pt(12)
        p.line_spacing = 1.25

    footer = slide.shapes.add_textbox(
        Inches(0.95), Inches(6.85), Inches(11.4), Inches(0.35)
    )
    tf = footer.text_frame
    p = tf.paragraphs[0]
    p.text = heading
    p.alignment = PP_ALIGN.RIGHT
    set_run(p.runs[0], 11, theme["muted"], font=theme["font"])


def _add_shape(
    slide: Any,
    shape_type: Any,
    x: int,
    y: int,
    width: int,
    height: int,
    fill: str,
    line: str | None = None,
) -> Any:
    shape = slide.shapes.add_shape(shape_type, x, y, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    if line:
        shape.line.color.rgb = _rgb(line)
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_cards_slide(
    prs: Presentation,
    index: int,
    heading: str,
    items: list[str],
    theme_id: str,
) -> None:
    theme = get_ppt_theme(theme_id)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H, theme["content_bg"])
    _add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.09), theme["accent"])

    label = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(8), Inches(0.4))
    p = label.text_frame.paragraphs[0]
    p.text = f"{index:02d}"
    set_run(p.runs[0], 15, theme["accent"], bold=True, font=theme["font"])

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.95), Inches(11.8), Inches(0.8))
    p = title_box.text_frame.paragraphs[0]
    p.text = heading
    set_run(p.runs[0], 30, theme["text"], bold=True, font=theme["font"])

    cards = items or ["暂无内容"]
    columns = 2 if len(cards) <= 4 else 3
    rows = (len(cards) + columns - 1) // columns
    card_w = Inches(3.8) if columns == 2 else Inches(2.9)
    card_h = Inches(1.8) if rows > 1 else Inches(2.1)
    gap_x = Inches(0.3)
    gap_y = Inches(0.25)
    start_x = Inches(0.7)
    start_y = Inches(2.0)

    for card_index, text in enumerate(cards):
        row = card_index // columns
        col = card_index % columns
        x = int(start_x + col * (card_w + gap_x))
        y = int(start_y + row * (card_h + gap_y))
        card = _add_shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            x,
            y,
            card_w,
            card_h,
            theme["accent_light"],
            theme.get("border"),
        )
        dot = _add_shape(
            slide,
            MSO_SHAPE.OVAL,
            int(x + Inches(0.18)),
            int(y + Inches(0.18)),
            Inches(0.14),
            Inches(0.14),
            theme["accent"],
        )
        box = slide.shapes.add_textbox(
            int(x + Inches(0.18)),
            int(y + Inches(0.42)),
            int(card_w - Inches(0.36)),
            int(card_h - Inches(0.6)),
        )
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        set_run(p.runs[0], 12, theme["text"], font=theme["font"])

    footer = slide.shapes.add_textbox(Inches(0.7), Inches(6.85), Inches(11.8), Inches(0.35))
    p = footer.text_frame.paragraphs[0]
    p.text = heading
    p.alignment = PP_ALIGN.RIGHT
    set_run(p.runs[0], 11, theme["muted"], font=theme["font"])


def add_timeline_slide(
    prs: Presentation,
    index: int,
    heading: str,
    items: list[str],
    theme_id: str,
) -> None:
    theme = get_ppt_theme(theme_id)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H, theme["content_bg"])
    _add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.09), theme["accent"])

    label = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(8), Inches(0.4))
    p = label.text_frame.paragraphs[0]
    p.text = f"{index:02d}"
    set_run(p.runs[0], 15, theme["accent"], bold=True, font=theme["font"])

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.95), Inches(11.8), Inches(0.8))
    p = title_box.text_frame.paragraphs[0]
    p.text = heading
    set_run(p.runs[0], 30, theme["text"], bold=True, font=theme["font"])

    steps = items or ["阶段一", "阶段二", "阶段三"]
    line_y = Inches(3.35)
    start_x = Inches(1.0)
    end_x = Inches(12.3)
    _add_shape(slide, MSO_SHAPE.RECTANGLE, int(start_x), int(line_y), int(end_x - start_x), Inches(0.06), theme["accent"])

    step_count = max(len(steps), 1)
    spacing = int((end_x - start_x) / step_count)
    for step_index, text in enumerate(steps):
        cx = int(start_x + spacing * step_index + spacing // 2)
        _add_shape(
            slide,
            MSO_SHAPE.OVAL,
            int(cx - Inches(0.16)),
            int(line_y - Inches(0.13)),
            Inches(0.32),
            Inches(0.32),
            theme["accent"],
        )
        box = slide.shapes.add_textbox(
            int(cx - Inches(1.4)),
            int(line_y + Inches(0.35)),
            Inches(2.8),
            Inches(1.3),
        )
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = PP_ALIGN.CENTER
        set_run(p.runs[0], 12, theme["text"], font=theme["font"])

    footer = slide.shapes.add_textbox(Inches(0.7), Inches(6.85), Inches(11.8), Inches(0.35))
    p = footer.text_frame.paragraphs[0]
    p.text = heading
    p.alignment = PP_ALIGN.RIGHT
    set_run(p.runs[0], 11, theme["muted"], font=theme["font"])


def add_columns_slide(
    prs: Presentation,
    index: int,
    heading: str,
    items: list[str],
    theme_id: str,
) -> None:
    theme = get_ppt_theme(theme_id)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H, theme["content_bg"])
    _add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.09), theme["accent"])

    label = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(8), Inches(0.4))
    p = label.text_frame.paragraphs[0]
    p.text = f"{index:02d}"
    set_run(p.runs[0], 15, theme["accent"], bold=True, font=theme["font"])

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.95), Inches(11.8), Inches(0.8))
    p = title_box.text_frame.paragraphs[0]
    p.text = heading
    set_run(p.runs[0], 30, theme["text"], bold=True, font=theme["font"])

    columns = items or ["第一栏", "第二栏", "第三栏"]
    count = min(max(len(columns), 2), 3)
    col_w = Inches(3.8)
    gap = Inches(0.35)
    start_x = Inches(0.7)
    col_h = Inches(4.2)
    for col_index, text in enumerate(columns[:count]):
        x = int(start_x + col_index * (col_w + gap))
        panel = _add_shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            x,
            Inches(2.0),
            col_w,
            col_h,
            theme["accent_light"],
            theme.get("border"),
        )
        box = slide.shapes.add_textbox(
            int(x + Inches(0.25)),
            int(Inches(2.25)),
            int(col_w - Inches(0.5)),
            int(col_h - Inches(0.5)),
        )
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        set_run(p.runs[0], 13, theme["text"], font=theme["font"])

    footer = slide.shapes.add_textbox(Inches(0.7), Inches(6.85), Inches(11.8), Inches(0.35))
    p = footer.text_frame.paragraphs[0]
    p.text = heading
    p.alignment = PP_ALIGN.RIGHT
    set_run(p.runs[0], 11, theme["muted"], font=theme["font"])
