"""Office 编辑工具（fill_template / edit_docx / edit_xlsx / document_outline）
脚本侧与服务侧共享的工具函数：路径解析、编辑目标准备、产物重开自检。

路径策略（docs/design/017-office-edit-closed-loop-20260906.md §5.3）：
- workspace 内文件：原地编辑，编辑前字节级备份到 data_dir/backups/<session>/；
- 外部文件：**不原地改**——复制进 workspace/documents/edited/ 后编辑副本，
  原文件零风险；
- .doc 老格式直接拒绝并引导转换。
"""

from __future__ import annotations

import re
import time
from pathlib import Path
from typing import Any

_DOC_SUFFIXES = (".docx",)
_SHEET_SUFFIXES = (".xlsx", ".xlsm")
_SLIDE_SUFFIXES = (".pptx",)
_KIND_SUFFIXES = {"doc": _DOC_SUFFIXES, "sheet": _SHEET_SUFFIXES, "slide": _SLIDE_SUFFIXES}
_KIND_LABEL = {"doc": "Word(.docx)", "sheet": "Excel(.xlsx)", "slide": "PowerPoint(.pptx)"}


class OfficePathError(ValueError):
    """路径无效/类型不支持/文件缺失，文案直接面向用户。"""


def resolve_office_path(workspace: Path, raw: str, kind: str = "doc") -> Path:
    """解析输入路径（workspace 相对或绝对外部路径），校验类型与存在性。

    kind: doc / sheet / slide。
    """
    text = (raw or "").strip().strip('"').strip("'")
    if not text:
        raise OfficePathError("缺少文件路径")
    path = Path(text)
    if not path.is_absolute():
        path = workspace / path
    path = path.resolve()
    if not path.exists():
        raise OfficePathError(f"文件不存在: {path}")
    if path.is_dir():
        raise OfficePathError(f"路径是目录而不是文件: {path}")
    allowed = _KIND_SUFFIXES.get(kind, _DOC_SUFFIXES)
    if path.suffix.lower() == ".doc":
        raise OfficePathError(
            "暂不支持 .doc 老格式：请先另存为 .docx 再试，"
            "或让我读取内容后重新生成一份 .docx"
        )
    if path.suffix.lower() not in allowed:
        raise OfficePathError(
            f"文件类型不支持: {path.suffix}（需要 {_KIND_LABEL.get(kind, 'Word(.docx)')}）"
        )
    return path


def _unique_edited_copy(directory: Path, stem: str, suffix: str) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    ts = time.strftime("%Y%m%d_%H%M%S")
    candidate = directory / f"{stem}_{ts}{suffix}"
    counter = 1
    while candidate.exists():
        candidate = directory / f"{stem}_{ts}_{counter}{suffix}"
        counter += 1
    return candidate


def prepare_edit_target(
    workspace: Path,
    session_id: str,
    path: Path,
) -> dict[str, Any]:
    """决定编辑目标：workspace 内原地改（先备份），外部文件改副本。

    返回 {target, original, is_copy, backup, in_workspace}。
    """
    from .office_backup import backup_file

    workspace = workspace.resolve()
    try:
        path.relative_to(workspace)
        in_workspace = True
    except ValueError:
        in_workspace = False

    if in_workspace:
        backup = (
            backup_file(session_id, path) if session_id else None
        )
        return {
            "target": path,
            "original": path,
            "is_copy": False,
            "backup": backup,
            "in_workspace": True,
        }

    edited_dir = workspace / "documents" / "edited"
    stem = re.sub(r'[\\/:*?"<>|]+', "_", path.stem).strip("._") or "document"
    copy_path = _unique_edited_copy(edited_dir, stem, path.suffix.lower())
    copy_path.write_bytes(path.read_bytes())
    return {
        "target": copy_path,
        "original": path,
        "is_copy": True,
        "backup": None,
        "in_workspace": False,
    }


def verify_output(path: Path) -> list[str]:
    """改后自检：重新打开产物文件，任何异常都视为不可交付。"""
    suffix = path.suffix.lower()
    try:
        if suffix == ".docx":
            from docx import Document

            document = Document(str(path))
            if not document.paragraphs and not document.tables:
                return ["重开后文档为空"]
        elif suffix in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook

            workbook = load_workbook(str(path), data_only=False)
            if not workbook.worksheets:
                return ["重开后工作簿没有工作表"]
        elif suffix == ".pptx":
            from pptx import Presentation

            presentation = Presentation(str(path))
            if not len(presentation.slides):
                return ["重开后演示文稿没有任何页面"]
        else:
            return [f"未知文件类型: {suffix}"]
    except Exception as exc:  # noqa: BLE001 — 重开失败即产物不可交付
        return [f"产物重开失败（可能已损坏）: {exc}"]
    return []
