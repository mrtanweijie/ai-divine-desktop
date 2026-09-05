"""edit_pptx 脚本：对已有 pptx 做文本替换（replace_text/set_title）。

边界：只动文本，不动母版/版式/图表。事务性同 edit_docx——任一步失败不写盘。
外部文件复制进 workspace/documents/edited/ 编辑副本；workspace 内原地编辑前备份。
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

from pptx import Presentation  # noqa: E402

from app.docgen.ppt.editor import PptxOpError, apply_ops  # noqa: E402
from app.tools.office_files import (  # noqa: E402
    OfficePathError,
    prepare_edit_target,
    resolve_office_path,
)


def main() -> None:
    data = json.load(sys.stdin)
    workspace = Path(data["workspace"])
    params = data.get("params") or {}
    session_id = str(params.get("session_id") or "")
    ops = params.get("ops") or []
    if not isinstance(ops, list) or not ops:
        print(
            json.dumps({"success": False, "message": "缺少 ops 操作序列"}, ensure_ascii=False)
        )
        return

    try:
        source = resolve_office_path(workspace, str(params.get("path") or ""), kind="slide")
    except OfficePathError as exc:
        print(json.dumps({"success": False, "message": str(exc)}, ensure_ascii=False))
        return

    prep = prepare_edit_target(workspace, session_id, source)
    target: Path = prep["target"]
    try:
        prs = Presentation(str(target))
    except Exception as exc:  # noqa: BLE001
        _cleanup(prep)
        print(
            json.dumps(
                {
                    "success": False,
                    "message": f"文件打开失败（可能已加密或损坏）: {exc}",
                },
                ensure_ascii=False,
            )
        )
        return

    try:
        report = apply_ops(prs, ops)
        prs.save(str(target))
    except PptxOpError as exc:
        candidates = getattr(exc, "candidates", None) or []
        message = str(exc)
        if isinstance(getattr(exc, "op_index", None), int):
            message = f"第 {exc.op_index} 个操作失败：{message}"  # type: ignore[attr-defined]
        print(
            json.dumps(
                {
                    "success": False,
                    "message": message,
                    "content": message + "。文件未做任何修改。",
                    "data": {"candidates": candidates},
                },
                ensure_ascii=False,
            )
        )
        return
    except Exception as exc:  # noqa: BLE001
        _cleanup(prep)
        print(
            json.dumps(
                {"success": False, "message": f"编辑失败: {exc}"}, ensure_ascii=False
            )
        )
        return

    # 重开自检
    try:
        reopened = Presentation(str(target))
        if not len(reopened.slides):
            raise ValueError("重开后没有任何页面")
    except Exception as exc:  # noqa: BLE001
        _cleanup(prep)
        print(
            json.dumps(
                {"success": False, "message": f"产物重开失败（可能已损坏）: {exc}"},
                ensure_ascii=False,
            )
        )
        return

    prefix = "已创建编辑副本" if prep["is_copy"] else "已编辑"
    message = f"{prefix}：{len(report['ops_applied'])} 个操作全部生效"
    print(
        json.dumps(
            {
                "success": True,
                "message": message,
                "file_name": target.name,
                "filepath": str(target),
                "content": message,
                "data": {
                    **report,
                    "source_path": str(prep["original"]) if prep["is_copy"] else None,
                    "backup_path": str(prep["backup"]) if prep["backup"] else None,
                    "edited_copy": prep["is_copy"],
                },
            },
            ensure_ascii=False,
        )
    )


def _cleanup(prep: dict) -> None:
    """失败现场清理：副本删除，原地编辑用编辑前备份覆盖回去。"""
    target: Path = prep["target"]
    backup = prep.get("backup")
    if prep["is_copy"]:
        if target.exists():
            target.unlink(missing_ok=True)
    elif backup and Path(backup).is_file():
        target.write_bytes(Path(backup).read_bytes())


if __name__ == "__main__":
    main()
