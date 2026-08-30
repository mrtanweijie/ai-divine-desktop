"""生成 Excel 表格（专业样式版）。

样式能力：主题配色（自动/指定）、标题栏与信息行、表头着色、斑马纹、
细边框、CJK 感知列宽、冻结表头、自动筛选、数字/百分比/日期智能识别、
货币格式（按列头关键词）、可选合计行。
"""

import json
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

from app.tools.scripts.file_utils import unique_output_path  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402


def safe_name(name: str, fallback: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|\s]+', "_", name).strip("._")
    if not cleaned:
        cleaned = fallback
    if not cleaned.endswith(".xlsx"):
        cleaned += ".xlsx"
    return cleaned


# ---------------------------------------------------------------------------
# 主题
# ---------------------------------------------------------------------------

THEMES: dict[str, dict[str, str]] = {
    "ocean": {
        "label": "青碧",
        "primary": "0F766E",
        "dark": "134E4A",
        "banded": "EEF7F6",
        "accent": "CCFBF1",
        "border": "CBD5D3",
    },
    "blue": {
        "label": "商务蓝",
        "primary": "2563EB",
        "dark": "1E3A8A",
        "banded": "EFF4FE",
        "accent": "DBEAFE",
        "border": "C7D4EC",
    },
    "emerald": {
        "label": "翡翠绿",
        "primary": "059669",
        "dark": "064E3B",
        "banded": "EAF8F1",
        "accent": "D1FAE5",
        "border": "C4E3D5",
    },
    "orange": {
        "label": "活力橙",
        "primary": "EA580C",
        "dark": "7C2D12",
        "banded": "FDF2EA",
        "accent": "FFEDD5",
        "border": "EAD3C0",
    },
    "violet": {
        "label": "优雅紫",
        "primary": "7C3AED",
        "dark": "4C1D95",
        "banded": "F5F1FE",
        "accent": "EDE9FE",
        "border": "D8CEEE",
    },
    "graphite": {
        "label": "极简灰",
        "primary": "4B5563",
        "dark": "1F2937",
        "banded": "F4F5F7",
        "accent": "E5E7EB",
        "border": "D4D7DC",
    },
}

_THEME_ALIASES: dict[str, str] = {
    "ocean": "ocean",
    "teal": "ocean",
    "青碧": "ocean",
    "默认": "ocean",
    "blue": "blue",
    "商务蓝": "blue",
    "emerald": "emerald",
    "green": "emerald",
    "绿色": "emerald",
    "翡翠绿": "emerald",
    "orange": "orange",
    "amber": "orange",
    "活力橙": "orange",
    "橙色": "orange",
    "violet": "violet",
    "purple": "violet",
    "优雅紫": "violet",
    "紫色": "violet",
    "graphite": "graphite",
    "gray": "graphite",
    "grey": "graphite",
    "极简灰": "graphite",
    "灰色": "graphite",
}

# 按标题/文件名关键词自动选主题；顺序即优先级（越靠前越具体）
_AUTO_KEYWORDS: list[tuple[str, tuple[str, ...]]] = [
    ("emerald", ("财务", "预算", "报销", "工资", "成本", "费用", "账", "发票", "finance", "budget")),
    ("orange", ("销售", "业绩", "营收", "增长", "市场", "营销", "sales", "marketing")),
    ("violet", ("名单", "通讯录", "联系人", "花名册", "人员", "成员", "库存", "清单")),
    ("graphite", ("技术", "开发", "缺陷", "bug", "日志", "服务器", "系统", "监控", "tech", "dev")),
    ("blue", ("进度", "计划", "排期", "项目", "日程", "任务", "汇总", "roadmap", "plan", "schedule", "report")),
]

_CURRENCY_HEADER_KEYWORDS = (
    "金额",
    "价格",
    "单价",
    "成本",
    "费用",
    "工资",
    "营收",
    "收入",
    "支出",
    "总额",
    "总价",
    "预算",
    "余额",
)

_INT_RE = re.compile(r"^-?\d+$")
_FLOAT_RE = re.compile(r"^-?\d+\.\d+$")
_THOUSANDS_RE = re.compile(r"^-?\d{1,3}(,\d{3})+(\.\d+)?$")
_PERCENT_RE = re.compile(r"^(-?\d+(?:\.\d+)?)\s*%$")
_CNY_RE = re.compile(r"^[¥￥]\s*(-?\d+(?:,\d{3})*(?:\.\d+)?)$")
_DATE_RE = re.compile(r"^(\d{4})[-/年.](\d{1,2})[-/月.](\d{1,2})日?$")
_DATETIME_RE = re.compile(
    r"^(\d{4})[-/年.](\d{1,2})[-/月.](\d{1,2})日?\s+(\d{1,2}):(\d{2})(?::(\d{2}))?$"
)


def _display_width(value: object) -> int:
    """按显示宽度估列宽：东亚全角字符计 2，其余计 1。"""
    text = str(value) if value is not None else ""
    return sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1 for ch in text)


def _pick_theme(style: str, title: str, file_name: str) -> str:
    key = (style or "").strip()
    if key:
        theme = _THEME_ALIASES.get(key.lower() if key.isascii() else key)
        if theme:
            return theme
    haystack = f"{title} {file_name}"
    for theme, keywords in _AUTO_KEYWORDS:
        if any(k in haystack for k in keywords):
            return theme
    return "ocean"


def _sanitize_sheet_name(raw: str) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "", (raw or "").strip()).strip()
    if not cleaned:
        cleaned = "数据"
    return cleaned[:31]


def _is_currency_header(header: object) -> bool:
    text = str(header or "")
    return any(k in text for k in _CURRENCY_HEADER_KEYWORDS)


def _coerce_cell(value: object, currency: bool) -> tuple[object, str | None, str]:
    """把单元格值归一为 (写入值, 数字格式, 对齐类型)。

    对齐类型：text（左对齐）/ number（右对齐）/ center（居中）。
    """
    if value is None:
        return None, None, "text"
    if isinstance(value, bool):
        return ("是" if value else "否"), None, "center"
    if isinstance(value, (int, float)):
        if isinstance(value, int) or float(value).is_integer():
            return int(value), ("¥#,##0" if currency else "#,##0"), "number"
        return float(value), ("¥#,##0.00" if currency else "#,##0.00"), "number"
    if not isinstance(value, str):
        return str(value), None, "text"

    text = value.strip()
    if not text:
        return None, None, "text"

    m = _PERCENT_RE.match(text)
    if m:
        num = float(m.group(1)) / 100
        decimals = len(m.group(1).partition(".")[2])
        fmt = "0%" if decimals == 0 else "0." + "0" * decimals + "%"
        return num, fmt, "number"

    m = _CNY_RE.match(text)
    if m:
        return float(m.group(1).replace(",", "")), "¥#,##0.00", "number"

    m = _DATETIME_RE.match(text)
    if m:
        try:
            parts = [int(p) for p in m.groups() if p is not None]
            dt = datetime(
                parts[0], parts[1], parts[2],
                parts[3], parts[4], parts[5] if len(parts) > 5 else 0,
            )
            return dt, "yyyy-mm-dd hh:mm", "center"
        except ValueError:
            pass

    m = _DATE_RE.match(text)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3))), "yyyy-mm-dd", "center"
        except ValueError:
            pass

    if _INT_RE.match(text):
        return int(text), ("¥#,##0" if currency else "#,##0"), "number"
    if _THOUSANDS_RE.match(text):
        cleaned = text.replace(",", "")
        if "." in cleaned:
            return float(cleaned), ("¥#,##0.00" if currency else "#,##0.00"), "number"
        return int(cleaned), ("¥#,##0" if currency else "#,##0"), "number"
    if _FLOAT_RE.match(text):
        num = float(text)
        if num.is_integer():
            return int(num), ("¥#,##0" if currency else "#,##0"), "number"
        return num, ("¥#,##0.00" if currency else "#,##0.00"), "number"

    return value, None, "text"


def _style_header_cell(cell, theme: dict[str, str]) -> None:
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill("solid", fgColor=theme["primary"])
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(
        left=Side(style="thin", color=theme["primary"]),
        right=Side(style="thin", color=theme["primary"]),
        top=Side(style="thin", color=theme["primary"]),
        bottom=Side(style="medium", color=theme["dark"]),
    )


def _style_data_cell(cell, theme: dict[str, str], banded: bool, kind: str) -> None:
    if banded:
        cell.fill = PatternFill("solid", fgColor=theme["banded"])
    align = "right" if kind == "number" else "center" if kind == "center" else "left"
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border = Border(
        left=Side(style="thin", color=theme["border"]),
        right=Side(style="thin", color=theme["border"]),
        top=Side(style="thin", color=theme["border"]),
        bottom=Side(style="thin", color=theme["border"]),
    )


def _style_total_cell(cell, theme: dict[str, str]) -> None:
    cell.font = Font(bold=True, color=theme["dark"], size=11)
    cell.fill = PatternFill("solid", fgColor=theme["accent"])
    cell.alignment = Alignment(horizontal="right", vertical="center")
    cell.border = Border(
        left=Side(style="thin", color=theme["border"]),
        right=Side(style="thin", color=theme["border"]),
        top=Side(style="medium", color=theme["dark"]),
        bottom=Side(style="double", color=theme["dark"]),
    )


def _numeric_display(value: object) -> str:
    if isinstance(value, float) and not value.is_integer():
        return f"{value:,.2f}"
    return f"{int(value):,}"


def _column_widths(
    headers: list[str],
    coerced_rows: list[list[tuple[object, str | None, str]]],
    totals: list[object] | None = None,
) -> list[int]:
    widths: list[int] = []
    for col, header in enumerate(headers):
        candidates = [_display_width(header) + 4]
        for row in coerced_rows:
            if col >= len(row):
                continue
            value, _fmt, _kind = row[col]
            if value is None:
                continue
            if isinstance(value, bool):
                candidates.append(4)
            elif isinstance(value, (int, float)):
                candidates.append(len(_numeric_display(value)) + 2)
            elif isinstance(value, (date, datetime)):
                candidates.append(12)
            else:
                candidates.append(_display_width(value) + 2)
        if totals and col < len(totals) and isinstance(totals[col], (int, float)):
            candidates.append(len(_numeric_display(totals[col])) + 4)
        widths.append(min(max(max(candidates) + 2, 9), 46))
    return widths


def build_workbook(
    headers: list[str],
    rows: list[list[object]],
    *,
    title: str,
    theme_name: str,
    total_row: bool,
    sheet_name: str,
) -> tuple[Workbook, dict]:
    theme = THEMES.get(theme_name) or THEMES["ocean"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _sanitize_sheet_name(sheet_name)
    sheet.sheet_view.showGridLines = False

    if not headers:
        width = max((len(r) for r in rows), default=0)
        headers = [f"列{i + 1}" for i in range(width)]
    else:
        max_row_len = max((len(r) for r in rows), default=0)
        if max_row_len > len(headers):
            headers = headers + [f"列{i + 1}" for i in range(len(headers), max_row_len)]
    n_cols = len(headers)

    coerced_rows: list[list[tuple[object, str | None, str]]] = []
    for row in rows:
        cells = []
        for col in range(n_cols):
            raw = row[col] if col < len(row) else None
            cells.append(_coerce_cell(raw, _is_currency_header(headers[col])))
        coerced_rows.append(cells)

    # 列级格式统一：多数单元格是百分比时，该列全部百分比单元格（含裸小数 0 < v ≤ 1）
    # 统一为同一小数位（取列内最大小数位，上限 2 位）
    col_percent_fmt: dict[int, str] = {}
    for col in range(n_cols):
        numeric = [
            (value, fmt)
            for row in coerced_rows
            if col < len(row)
            for value, fmt, _kind in [row[col]]
            if isinstance(value, (int, float)) and not isinstance(value, bool)
        ]
        if not numeric:
            continue
        percent_count = sum(1 for _v, fmt in numeric if fmt and "%" in fmt)
        if percent_count / len(numeric) < 0.6:
            continue

        def _fmt_decimals(fmt: str) -> int:
            if "." not in fmt:
                return 0
            return fmt.index("%") - fmt.index(".") - 1

        max_decimals = min(
            max((_fmt_decimals(fmt) for _v, fmt in numeric if fmt and "%" in fmt), default=1),
            2,
        )
        unified = "0%" if max_decimals == 0 else "0." + "0" * max_decimals + "%"
        col_percent_fmt[col] = unified
        for row in coerced_rows:
            if col < len(row):
                value, fmt, kind = row[col]
                if (
                    isinstance(value, (int, float))
                    and not isinstance(value, bool)
                    and 0 < value <= 1
                    and (fmt and "%" in fmt or fmt in (None, "#,##0", "#,##0.00"))
                ):
                    row[col] = (value, unified, "number")

    # ---- 标题栏与信息行 ----
    row_cursor = 1
    if title:
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(n_cols, 2))
        cell = sheet.cell(row=1, column=1, value=title)
        cell.font = Font(bold=True, size=16, color=theme["dark"])
        cell.alignment = Alignment(horizontal="left", vertical="center")
        sheet.row_dimensions[1].height = 32

        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(n_cols, 2))
        meta_cell = sheet.cell(
            row=2,
            column=1,
            value=f"共 {len(coerced_rows)} 行 · 生成于 {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        )
        meta_cell.font = Font(size=9, color="6B7280")
        meta_cell.alignment = Alignment(horizontal="left", vertical="center")
        sheet.row_dimensions[2].height = 16
        row_cursor = 3

    # ---- 表头 ----
    header_row = row_cursor
    for col, header in enumerate(headers, start=1):
        _style_header_cell(sheet.cell(row=header_row, column=col, value=str(header)), theme)
    sheet.row_dimensions[header_row].height = 24

    # ---- 数据行（斑马纹）----
    for i, cells in enumerate(coerced_rows):
        r = header_row + 1 + i
        banded = i % 2 == 1
        for col, (value, fmt, kind) in enumerate(cells, start=1):
            cell = sheet.cell(row=r, column=col, value=value)
            if fmt:
                cell.number_format = fmt
            _style_data_cell(cell, theme, banded, kind)

    # ---- 合计行 ----
    totals: list[object] | None = None
    if total_row and coerced_rows:
        totals = []
        for col in range(n_cols):
            if col == 0:
                totals.append("合计")
                continue
            numeric = [
                c[0]
                for row in coerced_rows
                if col < len(row)
                for c in [row[col]]
                if isinstance(c[0], (int, float)) and not isinstance(c[0], bool)
            ]
            if numeric and len(numeric) >= max(1, int(len(coerced_rows) * 0.6)):
                totals.append(sum(numeric))
            else:
                totals.append(None)
        has_total = any(isinstance(v, (int, float)) for v in totals)
    else:
        has_total = False

    if has_total:
        r = header_row + 1 + len(coerced_rows)
        for col, value in enumerate(totals, start=1):
            cell = sheet.cell(row=r, column=col, value=value)
            _style_total_cell(cell, theme)
            if isinstance(value, (int, float)):
                if _is_currency_header(headers[col - 1]):
                    cell.number_format = (
                        "¥#,##0.00"
                        if isinstance(value, float) and not value.is_integer()
                        else "¥#,##0"
                    )
                elif (col - 1) in col_percent_fmt:
                    cell.number_format = col_percent_fmt[col - 1]
                else:
                    cell.number_format = (
                        "#,##0.00"
                        if isinstance(value, float) and not value.is_integer()
                        else "#,##0"
                    )
        sheet.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="center")

    last_row = header_row + len(coerced_rows) + (1 if has_total else 0)

    # ---- 列宽（CJK 感知）----
    for col, width in enumerate(_column_widths(headers, coerced_rows, totals if has_total else None), start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width

    # ---- 冻结表头 + 自动筛选 ----
    sheet.freeze_panes = f"A{header_row + 1}"
    if coerced_rows:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(n_cols)}{last_row}"

    meta = {
        "rows": len(coerced_rows),
        "columns": n_cols,
        "theme": theme_name,
        "theme_label": theme["label"],
        "total_row": has_total,
    }
    return workbook, meta


def main() -> None:
    data = json.load(sys.stdin)
    workspace = Path(data["workspace"])
    output_root = Path(data.get("output_dir") or workspace)
    params = data.get("params") or {}
    output_dir = output_root / "spreadsheets"
    raw_path = str(params.get("path") or "").strip()
    if raw_path:
        candidate = Path(raw_path)
        if not candidate.is_absolute():
            candidate = workspace / candidate
        output_dir = candidate.resolve()
        try:
            output_dir.relative_to(workspace.resolve())
        except ValueError:
            print(
                json.dumps(
                    {"success": False, "message": f"保存路径越界: {output_dir}"},
                    ensure_ascii=False,
                )
            )
            return
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = unique_output_path(
        output_dir,
        params.get("file_name") or "表格",
        "表格",
        ".xlsx",
    )

    headers = params.get("headers") or []
    rows = params.get("rows") or []
    if not headers and not rows:
        print(
            json.dumps(
                {"success": False, "message": "headers 和 rows 不能同时为空"},
                ensure_ascii=False,
            )
        )
        return

    file_name = str(params.get("file_name") or "表格")
    title = str(params.get("title") or "").strip()
    if not title and file_name and file_name not in ("表格", "数据"):
        title = file_name
    theme_name = _pick_theme(str(params.get("style") or ""), title, file_name)

    workbook, meta = build_workbook(
        headers,
        rows,
        title=title,
        theme_name=theme_name,
        total_row=bool(params.get("total_row")),
        sheet_name=str(params.get("sheet_name") or "数据"),
    )
    workbook.save(output_path)

    print(
        json.dumps(
            {
                "success": True,
                "message": (
                    f"Excel 表格生成完成（{meta['theme_label']}主题，"
                    f"{meta['rows']} 行 × {meta['columns']} 列）"
                ),
                "file_name": output_path.name,
                "filepath": str(output_path),
                "data": meta,
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
