"""fill_template 脚本：按后缀路由填充引擎。

- .docx：python-docx 填充正文/表格/页眉页脚 {{占位符}}，list_values 支持表格行展开；
- .xlsx：openpyxl 填充所有 sheet 的字符串单元格（公式不触碰，行展开不支持）。

外部文件先复制进 workspace/documents/edited/ 再编辑（原文件零风险）；
workspace 内文件原地编辑，编辑前字节级备份到 data_dir/backups/<session>/。
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

from docx import Document  # noqa: E402

from app.docgen.word.editor import fill_template as fill_docx  # noqa: E402
from app.docgen.xlsx_editor import fill_template as fill_xlsx  # noqa: E402
from app.tools.office_files import (  # noqa: E402
    OfficePathError,
    prepare_edit_target,
    resolve_office_path,
    verify_output,
)


def main() -> None:
    data = json.load(sys.stdin)
    workspace = Path(data["workspace"])
    params = data.get("params") or {}
    session_id = str(params.get("session_id") or "")
    values = params.get("values") or {}
    list_values = params.get("list_values") or {}
    if not isinstance(values, dict) or not values:
        print(
            json.dumps(
                {"success": False, "message": "缺少 values（占位符→值 映射）"},
                ensure_ascii=False,
            )
        )
        return

    try:
        source = resolve_office_path(workspace, str(params.get("path") or ""), kind="office")
    except OfficePathError as exc:
        print(json.dumps({"success": False, "message": str(exc)}, ensure_ascii=False))
        return

    prep = prepare_edit_target(workspace, session_id, source)
    target: Path = prep["target"]
    try:
        if source.suffix.lower() in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook  # noqa: PLC0415

            workbook = load_workbook(str(target))
            report = fill_xlsx(workbook, values)
            if list_values:
                for key in list_values:
                    report["missing_keys"].append(key)
                report["warnings"].append(
                    "list_values 表格行展开仅支持 Word 模板，Excel 已忽略"
                )
            report = _normalize_report(report)
            if not report["filled"] and not report["warnings"]:
                print(
                    json.dumps(
                        {
                            "success": False,
                            "message": "未找到任何可填充的 {{占位符}}；"
                            "请确认文件内容或改用 edit_xlsx 直接编辑",
                            "data": report,
                        },
                        ensure_ascii=False,
                    )
                )
                return
            workbook.save(str(target))
        else:
            document = Document(str(target))
            report = fill_docx(document, values, list_values)
            report = _normalize_report(report)
            if not report["filled"] and not report["warnings"]:
                print(
                    json.dumps(
                        {
                            "success": False,
                            "message": "未找到任何可填充的 {{占位符}}；"
                            "请确认文件内容或改用 edit_docx 直接编辑",
                            "data": report,
                        },
                        ensure_ascii=False,
                    )
                )
                return
            document.save(str(target))
    except Exception as exc:  # noqa: BLE001 — 保存失败需清理现场
        _cleanup(prep)
        print(
            json.dumps(
                {"success": False, "message": f"填充失败: {exc}"}, ensure_ascii=False
            )
        )
        return

    problems = verify_output(target)
    if problems:
        _cleanup(prep)
        print(
            json.dumps(
                {"success": False, "message": "；".join(problems)}, ensure_ascii=False
            )
        )
        return

    missing = report.get("missing_keys") or []
    total = sum(item["count"] for item in report["filled"])
    message = f"模板填充完成：命中 {total} 处占位符"
    if missing:
        message += f"，未命中 {len(missing)} 个：{', '.join(missing)}"
    print(
        json.dumps(
            {
                "success": True,
                "message": message,
                "file_name": target.name,
                "filepath": str(target),
                "content": message
                + "。"
                + "；".join(
                    f"{item['key']}×{item['count']}" for item in report["filled"]
                ),
                "data": {
                    **report,
                    "source_path": str(prep["original"]) if prep["is_copy"] else None,
                    "backup_path": str(prep["backup"]) if prep["backup"] else None,
                    "edited_copy": prep["is_copy"],
                },
            },
            ensure_ascii=False,
        )
    )


def _normalize_report(report: dict) -> dict:
    """两种引擎的报告字段对齐（filled/missing_keys/warnings 必为列表）。"""
    report.setdefault("filled", [])
    report.setdefault("missing_keys", [])
    report.setdefault("warnings", [])
    return report


def _cleanup(prep: dict) -> None:
    """失败现场清理：副本删除，原地编辑用编辑前备份覆盖回去。"""
    target: Path = prep["target"]
    backup = prep.get("backup")
    if prep["is_copy"]:
        if target.exists():
            target.unlink(missing_ok=True)
    elif backup and Path(backup).is_file():
        target.write_bytes(Path(backup).read_bytes())


if __name__ == "__main__":
    main()
