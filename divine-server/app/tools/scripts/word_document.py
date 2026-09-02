import json
import sys
import time
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

from app.tools.scripts.file_utils import unique_output_path  # noqa: E402
from app.docgen.word.elements import (  # noqa: E402
    add_body_text,
    add_code_block,
    add_image,
    add_quote,
)
from app.docgen.word.markdown import markdown_to_document  # noqa: E402
from app.docgen.word.themes import (  # noqa: E402
    add_footer,
    add_styled_table,
    add_title_page,
    add_toc,
    apply_word_theme,
)

HEADING_TYPES = {
    "heading": 2,
    "h1": 1,
    "heading1": 1,
    "title": 1,
    "chapter": 1,
    "h2": 2,
    "heading2": 2,
    "h3": 3,
    "heading3": 3,
    "subheading": 3,
}
BULLET_TYPES = ("bullet_list", "list", "bullets", "bullet", "unordered_list")
NUMBERED_TYPES = ("numbered_list", "numbered", "ordered_list", "steps", "ol")
TABLE_TYPES = ("table", "grid")
QUOTE_TYPES = ("quote", "blockquote", "callout")
CODE_TYPES = ("code", "code_block", "codeblock")
IMAGE_TYPES = ("image", "picture", "img")
PARAGRAPH_TYPES = ("paragraph", "text", "body", "note", "")


def _section_content(section: dict) -> str:
    return str(
        section.get("content")
        or section.get("body")
        or section.get("text")
        or ""
    )


def _section_items(section: dict) -> list[str]:
    raw = section.get("items") or section.get("points") or section.get("bullets")
    if isinstance(raw, list):
        return [str(item).strip() for item in raw if str(item).strip()]
    return []


def _heading_level(section: dict, section_type: str) -> int:
    if section_type in HEADING_TYPES:
        return HEADING_TYPES[section_type]
    level = section.get("level")
    if isinstance(level, int) and 1 <= level <= 3:
        return level
    return 2


def add_sections(
    document: Document,
    sections: list,
    workspace: Path,
    theme: str,
) -> tuple[dict, list[str]]:
    counts = {
        "headings": 0,
        "paragraphs": 0,
        "lists": 0,
        "tables": 0,
        "images": 0,
        "code_blocks": 0,
        "quotes": 0,
    }
    warnings: list[str] = []

    for item in sections or []:
        section = item if isinstance(item, dict) else {"body": str(item)}
        section_type = str(section.get("type") or "").lower().strip()
        content = _section_content(section)
        heading = str(section.get("heading") or section.get("title") or "").strip()

        if not content and not heading and section_type not in TABLE_TYPES + IMAGE_TYPES + ("page_break", "toc"):
            continue

        if section_type in HEADING_TYPES:
            text = content or heading
            if text:
                document.add_heading(text, level=_heading_level(section, section_type))
                counts["headings"] += 1
            continue

        if section_type in BULLET_TYPES or section_type in NUMBERED_TYPES:
            lines = _section_items(section) or [
                line.strip() for line in content.splitlines() if line.strip()
            ]
            style = "List Number" if section_type in NUMBERED_TYPES else "List Bullet"
            # List Number 的编号在多个列表间会连续（python-docx 无重启编号 API）
            for line in lines:
                document.add_paragraph(line, style=style)
                counts["lists"] += 1
            continue

        if section_type in TABLE_TYPES:
            headers = section.get("headers") or section.get("columns") or []
            rows = section.get("rows") or section.get("data") or []
            table = add_styled_table(
                document, headers, rows, theme, section.get("caption")
            )
            if table is not None:
                counts["tables"] += 1
            else:
                warnings.append("表格元素缺少 headers/rows，已跳过")
            continue

        if section_type in QUOTE_TYPES:
            add_quote(document, content or heading, theme)
            counts["quotes"] += 1
            continue

        if section_type in CODE_TYPES:
            add_code_block(document, content, theme)
            counts["code_blocks"] += 1
            continue

        if section_type in IMAGE_TYPES:
            src = str(
                section.get("src") or section.get("url") or section.get("path") or ""
            ).strip()
            if not src:
                warnings.append("图片元素缺少 src，已跳过")
                continue
            if add_image(document, src, workspace, theme, section.get("caption")):
                counts["images"] += 1
            else:
                warnings.append(f"图片不可用，已跳过：{src}")
            continue

        if section_type == "page_break":
            document.add_page_break()
            continue

        if section_type == "toc":
            add_toc(document, theme)
            continue

        # 普通段落：可选 heading + 正文宽容解析
        if heading:
            document.add_heading(heading, level=_heading_level(section, section_type))
            counts["headings"] += 1
        body_counts = add_body_text(document, content, workspace, theme)
        for key, value in body_counts.items():
            counts[key] = counts.get(key, 0) + value

    return counts, warnings


def verify_document(path: Path, expected: dict) -> list[str]:
    """生成后读回自检：元素数量、空标题、残留 markdown 标记。"""
    warnings: list[str] = []
    document = Document(str(path))
    if "tables" in expected and len(document.tables) != expected["tables"]:
        warnings.append(
            f"表格数量不符：预期 {expected['tables']}，实际 {len(document.tables)}"
        )
    if "images" in expected and len(document.inline_shapes) != expected["images"]:
        warnings.append(
            f"图片数量不符：预期 {expected['images']}，实际 {len(document.inline_shapes)}"
        )
    empty_headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading") and not paragraph.text.strip()
    ]
    if empty_headings:
        warnings.append(f"存在 {len(empty_headings)} 个空标题")
    leftover = sum(
        1
        for paragraph in document.paragraphs
        if "**" in paragraph.text or "](" in paragraph.text
    )
    if leftover:
        warnings.append(f"有 {leftover} 个段落疑似残留 markdown 标记")
    return warnings


def main() -> None:
    data = json.load(sys.stdin)
    workspace = Path(data["workspace"])
    output_root = Path(data.get("output_dir") or workspace)
    params = data.get("params") or {}
    title = params.get("title") or "未命名文档"
    sections = params.get("sections") or []
    theme = str(params.get("theme") or params.get("style") or "universal").strip()
    markdown = str(params.get("markdown") or "").strip()
    cover = bool(params.get("cover"))
    toc = bool(params.get("toc"))
    output_dir = output_root / "documents"
    raw_path = str(params.get("path") or "").strip()
    if raw_path:
        candidate = Path(raw_path)
        if not candidate.is_absolute():
            candidate = workspace / candidate
        output_dir = candidate.resolve()
        try:
            output_dir.relative_to(workspace.resolve())
        except ValueError:
            print(
                json.dumps(
                    {"success": False, "message": f"保存路径越界: {output_dir}"},
                    ensure_ascii=False,
                )
            )
            return
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = unique_output_path(output_dir, title, "文档", ".docx")

    for attempt in range(2):
        try:
            document = Document()
            apply_word_theme(document, theme)
            add_title_page(document, title, theme, cover=cover)
            if toc:
                add_toc(document, theme)

            warnings: list[str] = []
            if markdown:
                counts = markdown_to_document(document, markdown, workspace, theme)
            elif sections:
                counts, warnings = add_sections(
                    document, sections, workspace, theme
                )
            else:
                counts = {}
                document.add_paragraph("此文档由 AI Divine 自动生成。")
            add_footer(document, theme)
            document.save(output_path)
            break
        except Exception:
            if output_path.exists():
                output_path.unlink()
            if attempt == 1:
                raise
            time.sleep(1)

    expected = {
        key: counts.get(key, 0)
        for key in ("tables", "images")
        if counts
    }
    warnings.extend(verify_document(output_path, expected))
    print(
        json.dumps(
            {
                "success": True,
                "message": "Word 文档生成完成",
                "file_name": output_path.name,
                "filepath": str(output_path),
                "data": {
                    "theme": theme,
                    "cover": cover,
                    "toc": toc,
                    **(counts or {}),
                    "warnings": warnings,
                },
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
