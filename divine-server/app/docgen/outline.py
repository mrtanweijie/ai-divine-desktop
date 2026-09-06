"""Office 文档结构大纲（document_outline 工具的读取层）。

与代码侧 list_definitions 对偶：输出可编辑坐标（段落 index、表格序号、
sheet 清单），供 edit_docx / edit_xlsx 定位使用。
坐标约定与编辑器一致：docx 段落 index 是 document.paragraphs 的 1 基下标；
表格 index 是 document.tables 的 1 基下标。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

_PREVIEW_CHARS = 50


def _preview(text: str) -> str:
    cleaned = " ".join(text.split())
    return cleaned[:_PREVIEW_CHARS]


def docx_outline(path: Path) -> dict[str, Any]:
    from docx import Document

    document = Document(str(path))
    paragraphs: list[dict[str, Any]] = []
    for index, paragraph in enumerate(document.paragraphs, 1):
        text = paragraph.text
        if not text.strip():
            continue
        paragraphs.append(
            {
                "index": index,
                "style": paragraph.style.name if paragraph.style is not None else "",
                "preview": _preview(text),
            }
        )
    tables: list[dict[str, Any]] = []
    for index, table in enumerate(document.tables, 1):
        headers: list[str] = []
        if table.rows:
            headers = [cell.text.strip() for cell in table.rows[0].cells]
        tables.append(
            {
                "index": index,
                "rows": len(table.rows),
                "cols": len(table.columns),
                "headers": headers,
            }
        )
    return {
        "kind": "docx",
        "paragraph_count": len(document.paragraphs),
        "paragraphs": paragraphs,
        "tables": tables,
    }


def xlsx_outline(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=False, data_only=False)
    sheets: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        headers: list[str] = []
        first_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if first_row:
            headers = [str(v).strip() if v is not None else "" for v in first_row]
        formula_count = 0
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
        sheets.append(
            {
                "name": worksheet.title,
                "dims": worksheet.dimensions or "",
                "headers": headers,
                "has_formulas": formula_count,
                "has_chart": bool(getattr(worksheet, "_charts", None)),
            }
        )
    return {"kind": "xlsx", "sheets": sheets}


def pptx_outline(path: Path) -> dict[str, Any]:
    """PPT 页级结构大纲：页序号 + 标题 + 文本预览 + 是否有备注。"""
    from pptx import Presentation

    prs = Presentation(str(path))
    slides: list[dict[str, Any]] = []
    for index, slide in enumerate(prs.slides, 1):
        title = ""
        try:
            if slide.shapes.title is not None:
                title = slide.shapes.title.text_frame.text.strip()
        except Exception:  # noqa: BLE001 — 无标题占位符的版式
            title = ""
        previews: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame is not title:
                text = " ".join(shape.text_frame.text.split())
                if text:
                    previews.append(text[:_PREVIEW_CHARS])
        has_note = False
        try:
            has_note = slide.has_notes_slide and bool(
                slide.notes_slide.notes_text_frame.text.strip()
            )
        except Exception:  # noqa: BLE001 — 备注读取失败不影响大纲
            has_note = False
        slides.append(
            {
                "index": index,
                "title": title,
                "previews": previews[:4],
                "has_note": has_note,
            }
        )
    return {"kind": "pptx", "slides": slides}


def outline(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return docx_outline(path)
    if suffix in (".xlsx", ".xlsm"):
        return xlsx_outline(path)
    if suffix == ".pptx":
        return pptx_outline(path)
    raise ValueError(f"不支持的文件类型: {suffix}（仅支持 .docx/.xlsx/.pptx）")
