from __future__ import annotations

from datetime import datetime
from typing import Any

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

WORD_THEMES: dict[str, dict[str, Any]] = {
    "universal": {
        "name": "通用现代",
        "font": "PingFang SC",
        "font_western": "Helvetica",
        "size": 10.5,
        "line": 1.6,
        "margin": 2.2,
        "accent": "2563EB",
        "title": "111827",
        "heading": "1E40AF",
        "heading_style": "underline",
    },
    "notion": {
        "name": "Notion 极简",
        "font": "PingFang SC",
        "font_western": "Helvetica",
        "size": 10.5,
        "line": 1.8,
        "margin": 2.5,
        "accent": "37352F",
        "title": "37352F",
        "heading": "37352F",
        "heading_style": "plain",
    },
    "linear": {
        "name": "Linear 紧凑",
        "font": "PingFang SC",
        "font_western": "Helvetica",
        "size": 9.5,
        "line": 1.4,
        "margin": 1.8,
        "accent": "5E6AD2",
        "title": "202129",
        "heading": "202129",
        "heading_style": "bar",
    },
    "apple": {
        "name": "Apple 舒适",
        "font": "PingFang SC",
        "font_western": "Helvetica",
        "size": 10.5,
        "line": 1.9,
        "margin": 2.4,
        "accent": "0071E3",
        "title": "1D1D1F",
        "heading": "1D1D1F",
        "heading_style": "underline",
    },
    "academic": {
        "name": "学术典雅",
        "font": "Noto Serif SC",
        "font_western": "Georgia",
        "size": 11,
        "line": 2.0,
        "margin": 2.5,
        "accent": "742A2A",
        "title": "1A202C",
        "heading": "742A2A",
        "heading_style": "plain",
    },
    "travel": {
        "name": "旅行暖橙",
        "font": "PingFang SC",
        "font_western": "Helvetica",
        "size": 10.5,
        "line": 1.7,
        "margin": 2.0,
        "accent": "EA580C",
        "title": "9A3412",
        "heading": "EA580C",
        "heading_style": "bar",
    },
    "roadmap": {
        "name": "知识图谱",
        "font": "PingFang SC",
        "font_western": "Helvetica",
        "size": 10.5,
        "line": 1.7,
        "margin": 2.0,
        "accent": "7C3AED",
        "title": "4C1D95",
        "heading": "7C3AED",
        "heading_style": "bar",
    },
    "minimal": {
        "name": "极简黑白",
        "font": "PingFang SC",
        "font_western": "Helvetica",
        "size": 10.5,
        "line": 2.0,
        "margin": 2.6,
        "accent": "222222",
        "title": "111111",
        "heading": "222222",
        "heading_style": "plain",
    },
}

# modern 与 universal 七个 token 完全相同，收敛为别名而非独立主题
WORD_THEME_ALIASES = {"modern": "universal"}


def resolve_theme(theme_id: str) -> dict[str, Any]:
    normalized = WORD_THEME_ALIASES.get(
        str(theme_id or "").strip().lower(),
        str(theme_id or "").strip().lower(),
    )
    return WORD_THEMES.get(normalized, WORD_THEMES["universal"])


def _tint(color_hex: str, ratio: float = 0.92) -> str:
    """把十六进制颜色向白色混合，ratio 为白色占比。"""
    channels = [
        int(color_hex[index : index + 2], 16) for index in (0, 2, 4)
    ]
    mixed = [round(c + (255 - c) * ratio) for c in channels]
    return "".join(f"{c:02X}" for c in mixed)


def _set_style_font(
    style: Any,
    font_western: str,
    font_east: str,
    size: float,
    color: str,
) -> None:
    style.font.name = font_western
    style.font.size = Pt(size)
    if isinstance(color, str):
        style.font.color.rgb = RGBColor.from_string(color)
    else:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    # 西文走 ascii/hAnsi、中文走 eastAsia，避免 Windows 上西文也落入中文字体
    rfonts.set(qn("w:ascii"), font_western)
    rfonts.set(qn("w:hAnsi"), font_western)
    rfonts.set(qn("w:eastAsia"), font_east)


def _style_border(style: Any, side: str, color: str, sz: int, space: int = 4) -> None:
    ppr = style.element.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    border = OxmlElement(f"w:{side}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(sz))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color)
    pbdr.append(border)


def _style_shading(style: Any, fill: str) -> None:
    ppr = style.element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def apply_word_theme(document: Document, theme_id: str) -> None:
    theme = resolve_theme(theme_id)
    font_east = theme["font"]
    font_western = theme["font_western"]
    accent = theme["accent"]

    normal = document.styles["Normal"]
    _set_style_font(normal, font_western, font_east, theme["size"], "111827")
    normal.paragraph_format.line_spacing = theme["line"]
    normal.paragraph_format.space_after = Pt(6)

    heading_map = {
        "Heading 1": (18, theme["title"], 10),
        "Heading 2": (14, theme["heading"], 8),
        "Heading 3": (12, theme["heading"], 6),
    }
    for style_name, (size, color, space_before) in heading_map.items():
        style = document.styles[style_name]
        _set_style_font(style, font_western, font_east, size, color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(4)

    # 版式级主题差异：标题装饰（下划线/侧条/底色块），不再只是换色
    decoration = theme.get("heading_style", "plain")
    if decoration == "underline":
        _style_border(document.styles["Heading 1"], "bottom", accent, 12, 6)
    elif decoration == "bar":
        for style_name, sz in (
            ("Heading 1", 28),
            ("Heading 2", 20),
            ("Heading 3", 14),
        ):
            _style_border(document.styles[style_name], "left", accent, sz, 6)
    elif decoration == "block":
        _style_shading(document.styles["Heading 2"], _tint(accent, 0.90))

    for section in document.sections:
        section.top_margin = Cm(theme["margin"])
        section.bottom_margin = Cm(theme["margin"] + 0.4)
        section.left_margin = Cm(theme["margin"] - 0.2)
        section.right_margin = Cm(theme["margin"] - 0.2)


def add_title_page(
    document: Document,
    title: str,
    theme_id: str,
    cover: bool = False,
    subtitle: str | None = None,
) -> None:
    theme = resolve_theme(theme_id)
    font_east = theme["font"]
    font_western = theme["font_western"]

    def _run_font(run: Any, size: float) -> None:
        run.font.size = Pt(size)
        rpr = run._r.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:ascii"), font_western)
        rfonts.set(qn("w:hAnsi"), font_western)
        rfonts.set(qn("w:eastAsia"), font_east)

    if not cover:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(title)
        run.font.size = Pt(26)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(theme["title"])
        _run_font(run, 26)
        paragraph.paragraph_format.space_before = Pt(120)

        sub = document.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub.add_run(subtitle or "AI Divine 生成文档")
        sub_run.font.size = Pt(12)
        sub_run.font.color.rgb = RGBColor.from_string(theme["accent"])
        _run_font(sub_run, 12)
        sub.paragraph_format.space_after = Pt(36)
        return

    # 正式封面：独占一页，首页页脚关闭
    section = document.sections[0]
    section.different_first_page_header_footer = True

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(150)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(theme["title"])
    _run_font(run, 30)
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "12")
    bottom.set(qn("w:color"), theme["accent"])
    pbdr.append(bottom)
    ppr.append(pbdr)

    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_before = Pt(18)
    sub_run = sub.add_run(subtitle or "AI Divine 生成文档")
    sub_run.font.size = Pt(13)
    sub_run.font.color.rgb = RGBColor.from_string(theme["accent"])
    _run_font(sub_run, 13)

    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_before = Pt(10)
    date_run = date.add_run(f"{datetime.now().strftime('%Y年%m月%d日')} 生成")
    date_run.font.size = Pt(11)
    date_run.font.color.rgb = RGBColor.from_string("6B7280")
    _run_font(date_run, 11)

    brand = document.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    brand.paragraph_format.space_before = Pt(200)
    brand_run = brand.add_run("AI Divine")
    brand_run.font.size = Pt(9)
    brand_run.font.color.rgb = RGBColor.from_string("9CA3AF")
    _run_font(brand_run, 9)

    document.add_page_break()


def _add_field(paragraph: Any, instruction: str, size: float, color: str) -> None:
    run = paragraph.add_run()
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for element in (fld_begin, instr, fld_sep, text, fld_end):
        run._r.append(element)


def add_footer(
    document: Document,
    theme_id: str,
    brand: str = "AI Divine · Word Document",
) -> None:
    theme = resolve_theme(theme_id)
    muted = "9CA3AF"

    def _footer_run(paragraph: Any, text: str) -> None:
        run = paragraph.add_run(text)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(muted)
        rpr = run._r.get_or_add_rPr()
        rfonts = OxmlElement("w:rFonts")
        rfonts.set(qn("w:ascii"), theme["font_western"])
        rfonts.set(qn("w:hAnsi"), theme["font_western"])
        rfonts.set(qn("w:eastAsia"), theme["font"])
        rpr.append(rfonts)

    for section in document.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        content_width = section.page_width - section.left_margin - section.right_margin
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Emu(int(content_width)), WD_TAB_ALIGNMENT.RIGHT
        )
        _footer_run(paragraph, brand)
        _footer_run(paragraph, "\t第 ")
        _add_field(paragraph, " PAGE ", 8, muted)
        _footer_run(paragraph, " 页 / 共 ")
        _add_field(paragraph, " NUMPAGES ", 8, muted)
        _footer_run(paragraph, " 页")


def add_toc(document: Document, theme_id: str) -> None:
    theme = resolve_theme(theme_id)
    label = document.add_paragraph()
    label_run = label.add_run("目录")
    label_run.font.size = Pt(16)
    label_run.font.bold = True
    label_run.font.color.rgb = RGBColor.from_string(theme["heading"])

    paragraph = document.add_paragraph()
    placeholder = paragraph.add_run(
        "（目录尚未生成：请在 Word 中全选后按 F9，或右键目录 → 更新域）"
    )
    placeholder.font.size = Pt(10)
    placeholder.font.color.rgb = RGBColor.from_string("9CA3AF")
    field_run = paragraph.add_run()
    field_run.font.size = Pt(10)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for element in (fld_begin, instr, fld_sep, fld_end):
        field_run._r.append(element)

    document.add_page_break()


def _display_width(text: str) -> int:
    return sum(2 if ord(char) > 0x2E80 else 1 for char in text)


def shade_table_header(document: Document, table: Any, theme_id: str) -> None:
    if not table.rows:
        return
    theme = resolve_theme(theme_id)
    accent = theme["accent"]
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = tc_pr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            tc_pr.append(shd)
        shd.set(qn("w:fill"), accent)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor.from_string("FFFFFF")
                run.font.bold = True


def _cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_styled_table(
    document: Document,
    headers: list,
    rows: list,
    theme_id: str,
    caption: str | None = None,
) -> Any:
    theme = resolve_theme(theme_id)
    header_cells = [str(header or "") for header in (headers or [])]
    data_rows = [
        ["" if cell is None else str(cell) for cell in row] for row in (rows or [])
    ]
    if not header_cells and not data_rows:
        return None
    ncols = len(header_cells) or max(len(row) for row in data_rows)
    header_cells += [""] * (ncols - len(header_cells))
    data_rows = [row + [""] * (ncols - len(row)) for row in data_rows]

    if caption:
        cap = document.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_run = cap.add_run(str(caption))
        cap_run.font.size = Pt(9)
        cap_run.font.color.rgb = RGBColor.from_string("6B7280")
        cap.paragraph_format.space_after = Pt(4)

    table = document.add_table(rows=len(data_rows) + 1, cols=ncols)
    table.style = "Table Grid"
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    for index, text in enumerate(header_cells):
        cell = table.rows[0].cells[index]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.paragraphs[0].add_run(text)
    shade_table_header(document, table, theme_id)

    zebra = _tint(theme["accent"], 0.94)
    for row_index, row in enumerate(data_rows):
        for col_index, text in enumerate(row):
            cell = table.rows[row_index + 1].cells[col_index]
            run = cell.paragraphs[0].add_run(text)
            run.font.size = Pt(theme["size"])
            rpr = run._r.get_or_add_rPr()
            rfonts = OxmlElement("w:rFonts")
            rfonts.set(qn("w:ascii"), theme["font_western"])
            rfonts.set(qn("w:hAnsi"), theme["font_western"])
            rfonts.set(qn("w:eastAsia"), theme["font"])
            rpr.append(rfonts)
            if row_index % 2 == 1:
                _cell_shading(cell, zebra)

    section = document.sections[0]
    content_cm = Emu(
        int(section.page_width - section.left_margin - section.right_margin)
    ).cm
    weights = []
    for col in range(ncols):
        cells_text = [header_cells[col]] + [
            row[col] for row in data_rows if col < len(row)
        ]
        weight = max((_display_width(text) for text in cells_text), default=4)
        weights.append(min(max(weight, 4), 42))
    total = sum(weights)
    widths = [max(content_cm * w / total, 1.6) for w in weights]
    scale = content_cm / sum(widths) if sum(widths) > content_cm else 1.0
    for col, width in enumerate(widths):
        table.columns[col].width = Cm(width * scale)
        for row in table.rows:
            row.cells[col].width = Cm(width * scale)

    separator = document.add_paragraph()
    separator.paragraph_format.space_before = Pt(2)
    separator.paragraph_format.space_after = Pt(2)
    return table
