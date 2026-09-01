"""Word 元素构建：行内格式、引用、代码块、图片与宽容 markdown 正文解析。"""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .themes import add_styled_table, resolve_theme

CODE_FONT = "Consolas"
CODE_BG = "F3F4F6"
MUTED_TEXT = "4B5563"

INLINE_RE = re.compile(
    r"\*\*(?P<bold>.+?)\*\*"
    r"|__(?P<bold2>.+?)__"
    r"|\*(?P<italic>[^*\n]+?)\*"
    r"|`(?P<code>[^`\n]+?)`"
    r"|\[(?P<link_text>[^\]\n]+)\]\((?P<link_url>https?://[^)\s]+)\)"
)

_BULLET_RE = re.compile(r"^(\s*)[-*+•·]\s+(.+)$")
_NUMBER_RE = re.compile(r"^\s*(\d+)[.、)]\s+(.+)$")
_HEADING_RE = re.compile(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$")
_HR_RE = re.compile(r"^\s*([-*_])\s*\1\s*\1[\s\-*_]*$")
_IMAGE_RE = re.compile(r"^\s*!\[([^\]]*)\]\(([^)\s]+)\)\s*$")
_QUOTE_PREFIX_RE = re.compile(r"^\s*>\s?")


def set_run_fonts(run, western: str, east: str) -> None:
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), western)
    rfonts.set(qn("w:hAnsi"), western)
    rfonts.set(qn("w:eastAsia"), east)


def shade_run(run, fill: str) -> None:
    rpr = run._r.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    rpr.append(shd)


def add_hyperlink(paragraph, text: str, url: str, theme_id: str) -> None:
    theme = resolve_theme(theme_id)
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), theme["font_western"])
    rfonts.set(qn("w:hAnsi"), theme["font_western"])
    rfonts.set(qn("w:eastAsia"), theme["font"])
    rpr.append(rfonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), theme["accent"])
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    text_el = OxmlElement("w:t")
    text_el.set(qn("xml:space"), "preserve")
    text_el.text = text
    run.append(rpr)
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_rich_text(paragraph, text: str, theme_id: str) -> None:
    """把 **加粗** / *斜体* / `代码` / [文字](链接) 解析为多个 run。"""
    theme = resolve_theme(theme_id)
    text = str(text or "")
    pos = 0

    def _plain(value: str) -> None:
        if not value:
            return
        run = paragraph.add_run(value)
        set_run_fonts(run, theme["font_western"], theme["font"])

    for match in INLINE_RE.finditer(text):
        _plain(text[pos : match.start()])
        if match.group("bold") is not None or match.group("bold2") is not None:
            value = match.group("bold") or match.group("bold2")
            run = paragraph.add_run(value)
            run.font.bold = True
            set_run_fonts(run, theme["font_western"], theme["font"])
        elif match.group("italic") is not None:
            run = paragraph.add_run(match.group("italic"))
            run.font.italic = True
            set_run_fonts(run, theme["font_western"], theme["font"])
        elif match.group("code") is not None:
            run = paragraph.add_run(match.group("code"))
            run.font.size = Pt(max(theme["size"] - 1, 8))
            run.font.name = CODE_FONT
            set_run_fonts(run, CODE_FONT, theme["font"])
            shade_run(run, CODE_BG)
        elif match.group("link_text") is not None:
            add_hyperlink(
                paragraph, match.group("link_text"), match.group("link_url"), theme_id
            )
        pos = match.end()
    _plain(text[pos:])


def add_quote(document, text: str, theme_id: str):
    """引用段：左侧强调色竖线 + 缩进 + 弱化文字色。"""
    theme = resolve_theme(theme_id)
    paragraph = document.add_paragraph()
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), theme["accent"])
    pbdr.append(left)
    ppr.append(pbdr)
    paragraph.paragraph_format.left_indent = Cm(0.4)
    run = paragraph.add_run(text)
    run.font.color.rgb = RGBColor.from_string(MUTED_TEXT)
    run.font.size = Pt(theme["size"])
    set_run_fonts(run, theme["font_western"], theme["font"])
    return paragraph


def add_code_block(document, code: str, theme_id: str) -> None:
    theme = resolve_theme(theme_id)
    lines = str(code or "").rstrip("\n").splitlines() or [""]
    for index, line in enumerate(lines):
        paragraph = document.add_paragraph()
        ppr = paragraph._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), CODE_BG)
        ppr.append(shd)
        paragraph.paragraph_format.space_before = Pt(6 if index == 0 else 0)
        paragraph.paragraph_format.space_after = Pt(6 if index == len(lines) - 1 else 0)
        run = paragraph.add_run(line or " ")
        run.font.name = CODE_FONT
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("1F2937")
        set_run_fonts(run, CODE_FONT, theme["font"])


def add_image(document, src: str, workspace: Path, theme_id: str, caption: str | None = None) -> bool:
    """插入图片：本地路径（workspace 相对/绝对）或 http(s) URL；超宽自动缩到版心。"""
    src = str(src or "").strip()
    if not src:
        return False
    data: bytes | None = None
    if src.startswith(("http://", "https://")):
        from app.pdf_theme_manager import _fetch_image_data  # 复用 PDF 侧下载器

        fetched = _fetch_image_data(src)
        if not fetched:
            return False
        data = fetched[0]
    else:
        path = Path(src)
        if not path.is_absolute():
            path = workspace / path
        if not path.is_file() or path.stat().st_size > 10 * 1024 * 1024:
            return False
        data = path.read_bytes()

    document.add_picture(BytesIO(data))
    shape = document.inline_shapes[-1]
    section = document.sections[0]
    max_width = section.page_width - section.left_margin - section.right_margin
    if shape.width > max_width:
        shape.width = max_width

    if caption:
        cap = document.add_paragraph()
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(str(caption))
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("6B7280")
    return True


def _is_table_line(line: str) -> bool:
    stripped = line.strip()
    return (
        stripped.startswith("|")
        and stripped.endswith("|")
        and stripped.count("|") >= 2
    )


def _is_table_separator(line: str) -> bool:
    stripped = line.strip().strip("|")
    return bool(stripped) and all(
        char in "-: " for part in stripped.split("|") for char in part
    ) and "-" in stripped


def _split_table_row(line: str) -> list[str]:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return cells


_COUNT_KEYS = ("headings", "paragraphs", "lists", "tables", "images", "code_blocks", "quotes")


def _new_counts() -> dict[str, int]:
    return dict.fromkeys(_COUNT_KEYS, 0)


def add_body_text(document, text: str, workspace: Path, theme_id: str) -> dict[str, int]:
    """宽容解析正文：逐行识别 markdown 常见块级语法，未命中时按富文本段落输出。

    返回实际创建的元素计数，供上层与生成后自检对账。
    """
    counts = _new_counts()
    lines = str(text or "").splitlines()
    index = 0
    while index < len(lines):
        raw = lines[index]
        line = raw.rstrip()
        stripped = line.strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith("```"):
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            index += 1  # 跳过收尾 fence
            add_code_block(document, "\n".join(code_lines), theme_id)
            counts["code_blocks"] += 1
            continue

        heading_match = _HEADING_RE.match(line)
        if heading_match:
            level = min(len(heading_match.group(1)), 3)
            document.add_heading(heading_match.group(2), level=level)
            counts["headings"] += 1
            index += 1
            continue

        if _HR_RE.match(stripped):
            index += 1
            continue

        image_match = _IMAGE_RE.match(line)
        if image_match:
            if add_image(document, image_match.group(2), workspace, theme_id):
                counts["images"] += 1
            index += 1
            continue

        if _is_table_line(line) and index + 1 < len(lines) and _is_table_separator(
            lines[index + 1]
        ):
            table_lines = [line]
            index += 2
            while index < len(lines) and _is_table_line(lines[index]):
                table_lines.append(lines[index])
                index += 1
            headers = _split_table_row(table_lines[0])
            rows = [_split_table_row(row) for row in table_lines[1:]]
            if add_styled_table(document, headers, rows, theme_id) is not None:
                counts["tables"] += 1
            continue

        bullet_match = _BULLET_RE.match(line)
        if bullet_match:
            level2 = 2 if len(bullet_match.group(1)) >= 2 else 1
            style = "List Bullet" if level2 == 1 else "List Bullet 2"
            document.add_paragraph(bullet_match.group(2).strip(), style=style)
            counts["lists"] += 1
            index += 1
            continue

        number_match = _NUMBER_RE.match(line)
        if number_match:
            document.add_paragraph(number_match.group(2).strip(), style="List Number")
            counts["lists"] += 1
            index += 1
            continue

        if stripped.startswith(">"):
            quote_text = _QUOTE_PREFIX_RE.sub("", line)
            add_quote(document, quote_text, theme_id)
            counts["quotes"] += 1
            index += 1
            continue

        paragraph = document.add_paragraph()
        add_rich_text(paragraph, line.strip(), theme_id)
        counts["paragraphs"] += 1
        index += 1

    return counts
