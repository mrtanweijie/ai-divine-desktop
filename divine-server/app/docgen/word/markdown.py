"""Markdown → Word 转换：markdown 库转 HTML，再用 HTMLParser 遍历映射 docx 元素。

行内格式（加粗/斜体/代码/链接）在 handle_data 时快照为 run 元组，
块结束时统一落入段落/标题/列表/引用/表格等元素。
"""

from __future__ import annotations

from html.parser import HTMLParser
from pathlib import Path
from typing import Any

import markdown as _markdown
from docx.shared import Pt

from .elements import (
    CODE_BG,
    CODE_FONT,
    add_code_block,
    add_hyperlink,
    add_image,
    add_quote,
    set_run_fonts,
    shade_run,
)
from .themes import add_styled_table, resolve_theme

_HEADING_LEVELS = {"h1": 1, "h2": 2, "h3": 3, "h4": 3, "h5": 3, "h6": 3}


class _MarkdownBuilder(HTMLParser):
    def __init__(self, document: Any, workspace: Path, theme_id: str) -> None:
        super().__init__(convert_charrefs=True)
        self.document = document
        self.workspace = workspace
        self.theme_id = theme_id
        self.theme = resolve_theme(theme_id)
        self.list_stack: list[bool] = []  # 每层 True=有序列表
        self.quote_depth = 0
        self.in_pre = False
        self.pre_buf: list[str] = []
        self.table_rows: list[list[tuple[bool, list[tuple[str, dict]]]]] = []
        self.cur_row: list[tuple[bool, list[tuple[str, dict]]]] | None = None
        self.cur_cell: list[tuple[str, dict]] | None = None
        self.bold = 0
        self.italic = 0
        self.code = 0
        self.link_url: str | None = None
        self.block: Any = None  # ("p", 0) | ("h", level) | "li"
        self.runs: list[tuple[str, dict]] = []
        self.counts = {
            "paragraphs": 0,
            "headings": 0,
            "tables": 0,
            "images": 0,
            "code_blocks": 0,
            "lists": 0,
            "quotes": 0,
        }

    # ---- 行内状态 ----

    def _fmt(self) -> dict:
        return {
            "bold": self.bold > 0,
            "italic": self.italic > 0,
            "code": self.code > 0,
            "href": self.link_url,
        }

    # ---- 块构建 ----

    def _emit_runs(self, paragraph: Any, runs: list[tuple[str, dict]]) -> None:
        theme = self.theme
        for text, fmt in runs:
            if fmt.get("break"):
                paragraph.add_run().add_break()
                continue
            if not text:
                continue
            href = fmt.get("href")
            if href:
                add_hyperlink(paragraph, text, href, self.theme_id)
                continue
            run = paragraph.add_run(text)
            if fmt.get("bold"):
                run.font.bold = True
            if fmt.get("italic"):
                run.font.italic = True
            if fmt.get("code"):
                run.font.name = CODE_FONT
                run.font.size = Pt(9)
                shade_run(run, CODE_BG)
            set_run_fonts(run, theme["font_western"], theme["font"])

    def _flush(self) -> None:
        block, runs = self.block, self.runs
        self.block, self.runs = None, []
        if block is None:
            return
        has_text = any(text.strip() for text, _ in runs)
        if not has_text and block != "li":
            return
        if not has_text and block == "li":
            return

        if block == "li":
            ordered = self.list_stack[-1] if self.list_stack else False
            level = min(max(len(self.list_stack), 1), 3)
            style = ("List Number" if ordered else "List Bullet") + (
                "" if level == 1 else f" {level}"
            )
            paragraph = self.document.add_paragraph(style=style)
            self.counts["lists"] += 1
        elif isinstance(block, tuple) and block[0] == "h":
            paragraph = self.document.add_paragraph(style=f"Heading {block[1]}")
            self.counts["headings"] += 1
        elif self.quote_depth > 0:
            paragraph = add_quote(self.document, "", self.theme_id)
            self.counts["quotes"] += 1
        else:
            paragraph = self.document.add_paragraph()
            self.counts["paragraphs"] += 1
        self._emit_runs(paragraph, runs)

    def _build_table(self) -> None:
        rows = self.table_rows
        self.table_rows = []
        if not rows:
            return
        headers: list[str] = []
        data: list[list[str]] = []
        for row_index, row in enumerate(rows):
            texts = [
                "".join(text for text, _ in cell_runs) for _, cell_runs in row
            ]
            has_header_cell = any(is_header for is_header, _ in row)
            if row_index == 0 and has_header_cell:
                headers = texts
            else:
                data.append(texts)
        table = add_styled_table(self.document, headers, data, self.theme_id)
        if table is not None:
            self.counts["tables"] += 1

    # ---- HTMLParser 回调 ----

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attributes = dict(attrs)
        if tag == "p":
            if self.block != "li":
                self._flush()
                self.block = ("p", 0)
            return
        if tag in _HEADING_LEVELS:
            self._flush()
            self.block = ("h", _HEADING_LEVELS[tag])
            return
        if tag in ("ul", "ol"):
            self._flush()
            self.list_stack.append(tag == "ol")
            return
        if tag == "li":
            self._flush()
            self.block = "li"
            return
        if tag == "blockquote":
            self._flush()
            self.quote_depth += 1
            return
        if tag == "pre":
            self._flush()
            self.in_pre = True
            self.pre_buf = []
            return
        if tag == "table":
            self._flush()
            self.table_rows = []
            return
        if tag == "tr":
            self.cur_row = []
            return
        if tag in ("td", "th"):
            self.cur_cell = []
            return
        if tag == "img":
            self._flush()
            src = str(attributes.get("src") or "")
            if add_image(self.document, src, self.workspace, self.theme_id):
                self.counts["images"] += 1
            return
        if tag == "hr":
            self._flush()
            return
        if tag == "br":
            if self.cur_cell is not None:
                self.cur_cell.append(("", {"break": True}))
            elif self.block is not None:
                self.runs.append(("", {"break": True}))
            return
        if tag in ("strong", "b"):
            self.bold += 1
            return
        if tag in ("em", "i"):
            self.italic += 1
            return
        if tag == "code" and not self.in_pre:
            self.code += 1
            return
        if tag == "a":
            href = attributes.get("href")
            if href and str(href).startswith(("http://", "https://")):
                self.link_url = str(href)
            return

    def handle_endtag(self, tag: str) -> None:
        if tag == "p":
            if self.block != "li":
                self._flush()
            return
        if tag in _HEADING_LEVELS or tag == "li":
            self._flush()
            return
        if tag in ("ul", "ol"):
            self._flush()
            if self.list_stack:
                self.list_stack.pop()
            return
        if tag == "blockquote":
            self._flush()
            self.quote_depth = max(self.quote_depth - 1, 0)
            return
        if tag == "pre":
            code = "".join(self.pre_buf)
            self.in_pre = False
            self.pre_buf = []
            if code.strip():
                add_code_block(self.document, code, self.theme_id)
                self.counts["code_blocks"] += 1
            return
        if tag == "tr":
            if self.cur_row is not None:
                self.table_rows.append(self.cur_row)
                self.cur_row = None
            return
        if tag in ("td", "th"):
            if self.cur_row is not None and self.cur_cell is not None:
                self.cur_row.append((tag == "th", self.cur_cell))
                self.cur_cell = None
            return
        if tag == "table":
            self._build_table()
            return
        if tag in ("strong", "b"):
            self.bold = max(self.bold - 1, 0)
            return
        if tag in ("em", "i"):
            self.italic = max(self.italic - 1, 0)
            return
        if tag == "code" and not self.in_pre:
            self.code = max(self.code - 1, 0)
            return
        if tag == "a":
            self.link_url = None
            return

    def handle_data(self, data: str) -> None:
        if self.in_pre:
            self.pre_buf.append(data)
            return
        if self.cur_cell is not None:
            self.cur_cell.append((data, self._fmt()))
            return
        if self.block is not None:
            self.runs.append((data, self._fmt()))

    def close(self) -> None:  # noqa: D102 - HTMLParser.close
        super().close()
        self._flush()


def markdown_to_document(
    document: Any,
    text: str,
    workspace: Path,
    theme_id: str,
) -> dict[str, int]:
    html = _markdown.markdown(
        str(text or ""),
        extensions=["tables", "fenced_code", "sane_lists"],
    )
    if not html.strip():
        return {}
    builder = _MarkdownBuilder(document, workspace, theme_id)
    builder.feed(html)
    builder.close()
    return builder.counts
