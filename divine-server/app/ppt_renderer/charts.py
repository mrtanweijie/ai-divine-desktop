"""原生图表渲染：python-pptx add_chart 直绘，PowerPoint 内可二次编辑。

覆盖简单图表（柱状/条形/折线/饼/环/面积）；复杂图表（sankey/组合图等）
由上层路由到 generate_chart SVG 管线，不经此模块。
"""

from __future__ import annotations

from typing import Any

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

from .themes import rgb

SUPPORTED_TYPES: dict[str, int] = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "area": XL_CHART_TYPE.AREA,
}

# 简单图表数据量上限：超出应走 generate_chart SVG 管线
MAX_SERIES = 3
MAX_POINTS = 12


def validate_chart_spec(spec: dict[str, Any]) -> None:
    chart_type = str(spec.get("type") or "column").lower()
    if chart_type not in SUPPORTED_TYPES:
        raise ValueError(
            f"原生图表不支持类型 {chart_type}，"
            f"仅支持 {'/'.join(SUPPORTED_TYPES)}；"
            "复杂图表请用 generate_chart 生成 SVG 后以 image.path 引用"
        )
    categories = spec.get("categories") or []
    series = spec.get("series") or []
    if not categories or not series:
        raise ValueError("chart 需要非空 categories 与 series")
    if len(series) > MAX_SERIES:
        raise ValueError(
            f"原生图表最多 {MAX_SERIES} 个系列，请改用 generate_chart SVG 管线"
        )
    if len(categories) > MAX_POINTS or any(
        len(s.get("values") or []) > MAX_POINTS for s in series
    ):
        raise ValueError(
            f"原生图表最多 {MAX_POINTS} 个数据点，请先聚合或改用 generate_chart"
        )


def add_native_chart(slide: Any, theme: dict[str, Any], spec: dict[str, Any],
                     x: float, y: float, w: float, h: float) -> Any:
    """在指定框内渲染原生图表，返回 graphic frame。"""
    validate_chart_spec(spec)
    chart_type = SUPPORTED_TYPES[str(spec.get("type") or "column").lower()]
    palette = theme.get("chart_palette") or ["2563EB"]
    font_name = theme.get("font") or "PingFang SC"
    muted = theme.get("muted") or "6B7280"
    text_color = theme.get("text") or "1A1D24"

    chart_data = CategoryChartData()
    chart_data.categories = [str(c) for c in spec["categories"]]
    for entry in spec["series"]:
        if isinstance(entry, dict):
            chart_data.add_series(
                str(entry.get("name") or "系列"), entry.get("values") or []
            )
        else:
            chart_data.add_series("系列", entry or [])

    frame = slide.shapes.add_chart(
        chart_type, Inches(x), Inches(y), Inches(w), Inches(h), chart_data
    )
    chart = frame.chart
    try:
        chart.font.size = Pt(11)
        chart.font.name = font_name
        chart.font.color.rgb = rgb(muted)
    except Exception:  # noqa: BLE001 — 图表级字体设置失败不阻塞
        pass

    title = str(spec.get("title") or "").strip()
    if title:
        chart.has_title = True
        chart.chart_title.text_frame.text = title
        run = chart.chart_title.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = rgb(text_color)
        run.font.name = font_name
    else:
        chart.has_title = False

    is_pie = chart_type in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT)
    try:
        if is_pie:
            points = chart.series[0].points
            for idx, point in enumerate(points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = rgb(
                    palette[idx % len(palette)]
                )
        else:
            for idx, series in enumerate(chart.series):
                color = palette[idx % len(palette)]
                if chart_type == XL_CHART_TYPE.LINE_MARKERS:
                    series.format.line.color.rgb = rgb(color)
                    series.format.line.width = Pt(2.25)
                    series.smooth = False
                else:
                    series.format.fill.solid()
                    series.format.fill.fore_color.rgb = rgb(color)
    except Exception:  # noqa: BLE001 — 配色失败保留默认样式
        pass

    try:
        plot = chart.plots[0]
        if chart_type in (XL_CHART_TYPE.BAR_CLUSTERED,
                          XL_CHART_TYPE.COLUMN_CLUSTERED):
            plot.gap_width = 80
        if is_pie:
            plot.has_data_labels = True
            labels = plot.data_labels
            labels.show_percentage = True
            labels.show_value = False
            labels.font.size = Pt(10.5)
            labels.font.color.rgb = rgb(text_color)
            labels.font.name = font_name
            try:
                labels.position = XL_LABEL_POSITION.OUTSIDE_END
            except Exception:  # noqa: BLE001 — 环形图不支持外部标签
                pass
    except Exception:  # noqa: BLE001
        pass

    multi_series = len(spec["series"]) > 1
    try:
        chart.has_legend = is_pie or multi_series
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(10.5)
            chart.legend.font.name = font_name
    except Exception:  # noqa: BLE001
        pass

    for axis_name in ("category_axis", "value_axis"):
        try:
            axis = getattr(chart, axis_name)
        except Exception:  # noqa: BLE001 — 饼图/环形图没有坐标轴
            continue
        try:
            axis.tick_labels.font.size = Pt(10.5)
            axis.tick_labels.font.name = font_name
            axis.tick_labels.font.color.rgb = rgb(muted)
        except Exception:  # noqa: BLE001
            continue

    return frame
