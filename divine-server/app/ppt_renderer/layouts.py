"""14 种版式渲染与分发，按主题的装饰语言（style）出画。

版式：cover / toc / section / bullets / two_col / cards / timeline /
columns / compare / kpi / quote / table / chart / end。
所有版式从主题 token 取色、从 STYLES 取画法，正文按 textfit 自适应字号。
"""

from __future__ import annotations

import math
import re
from typing import Any, Callable

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from . import charts, tables
from .textfit import pick_size, split_title_desc
from .themes import (add_oval, add_rect, add_textbox, blend, get_ppt_theme,
                     get_style, gradient, rgb, set_alpha, set_run, solid)

MARGIN = 0.7
CONTENT_W = 13.333 - MARGIN * 2
CONTENT_TOP = 2.0
CONTENT_H = 4.55
MONO_FONT = "Menlo"

LAYOUTS = {
    "cover", "toc", "section", "bullets", "two_col", "cards", "timeline",
    "columns", "compare", "kpi", "quote", "table", "chart", "end", "default",
}


# ---------------------------------------------------------------- 小工具

def _heading_font(theme: dict[str, Any], st: dict[str, Any]) -> str:
    if st.get("serif_heading"):
        return theme.get("heading_font") or theme["font"]
    return theme["font"]


def _mono_font() -> str:
    return MONO_FONT


def _rule_color(theme: dict[str, Any], st: dict[str, Any]) -> str:
    kind = st.get("toc_rule") or "muted_mix"
    if kind == "text":
        return theme["text"]
    if kind == "hairline":
        return blend(theme["content_bg"], theme["text"], 0.55)
    if kind == "accent_light":
        return blend(theme["accent_light"], theme["accent"], 0.35)
    return blend(theme["content_bg"], theme["muted"], 0.82)


# ---------------------------------------------------------------- 页面公共块

def _bg_decor(slide: Any, theme: dict[str, Any], st: dict[str, Any]) -> None:
    """内容页背景装饰：暗黑网格 / 玻璃光斑 / 瑞士红角标。"""
    decor = st.get("bg_decor")
    if decor == "grid":
        line_color = blend(theme["content_bg"], "FFFFFF", 0.05)
        for i in range(1, 7):
            add_rect(slide, i * 13.333 / 7, 0, 0.012, 7.5, line_color)
        for i in range(1, 4):
            add_rect(slide, 0, i * 7.5 / 4, 13.333, 0.012, line_color)
    elif decor == "blob":
        halo = add_oval(slide, 10.6, -1.4, 3.6, 3.6, theme["accent"])
        set_alpha(halo, 6)
    if st.get("corner_mark"):
        add_rect(slide, 12.35, 0.55, 0.22, 0.22, theme["accent"])


def _topbar(slide: Any, theme: dict[str, Any], st: dict[str, Any]) -> None:
    kind = st.get("topbar") or "solid"
    if kind == "none":
        return
    h = float(st.get("topbar_h") or 0.09)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                 Inches(13.333), Inches(h))
    if kind == "gradient":
        gradient(bar, theme["accent"], theme.get("accent2") or theme["accent"], 0)
    elif kind == "solid_dark":
        solid(bar, theme["text"])
    else:
        solid(bar, theme["accent"])


def _chrome(slide: Any, theme: dict[str, Any], st: dict[str, Any],
            index: int, heading: str) -> None:
    """内容页公共骨架：背景 + 顶条 + 序号 + 标题（含风格装饰）+ 页脚。"""
    add_rect(slide, 0, 0, 13.333, 7.5, theme["content_bg"])
    _bg_decor(slide, theme, st)
    _topbar(slide, theme, st)

    label = add_textbox(slide, MARGIN, 0.5, 8, 0.4)
    p = label.text_frame.paragraphs[0]
    set_run(p.add_run(), 14, theme["accent"], bold=True,
            font=_heading_font(theme, st) if st.get("index_serif") else theme["font"])
    p.runs[0].text = f"{index:02d}"

    if heading:
        size = pick_size(heading, CONTENT_W, 0.85, [30, 26, 22])
        box = add_textbox(slide, MARGIN, 0.92, CONTENT_W, 0.9)
        p = box.text_frame.paragraphs[0]
        set_run(p.add_run(), size, theme["text"], bold=True,
                font=_heading_font(theme, st))
        p.runs[0].text = heading
        rule_kind = st.get("heading_rule")
        if rule_kind:
            rule_h = float(st.get("heading_rule_h") or 0.02)
            rule_color = theme["text"] if rule_kind == "text" else \
                blend(theme["content_bg"], theme["text"], 0.55)
            add_rect(slide, MARGIN, 1.9, CONTENT_W, rule_h, rule_color)

    _footer(slide, theme, heading)


def _footer(slide: Any, theme: dict[str, Any], text: str) -> None:
    if not text:
        return
    box = add_textbox(slide, MARGIN, 6.98, CONTENT_W, 0.32)
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    display = text if len(text) <= 32 else text[:31] + "…"
    set_run(p.add_run(), 10.5, blend(theme["content_bg"], theme["muted"], 0.75),
            font=theme["font"])
    p.runs[0].text = display


_NUM_RE = re.compile(r"(\d[\d,，.]*\s*(?:万|亿|%|美元|款|城|人|倍|pt)?|\d+%)")


def _text_runs(p: Any, text: str, size: float, theme: dict[str, Any],
               font: str, bold: bool = False, color: str | None = None) -> None:
    """写一行文本，其中数字/百分比片段用强调色加粗。"""
    pos = 0
    for m in _NUM_RE.finditer(text):
        if m.start() > pos:
            run = p.add_run()
            set_run(run, size, color or theme["text"], bold=bold, font=font)
            run.text = text[pos:m.start()]
        run = p.add_run()
        set_run(run, size, theme["accent"], bold=True, font=font)
        run.text = m.group()
        pos = m.end()
    if pos < len(text):
        run = p.add_run()
        set_run(run, size, color or theme["text"], bold=bold, font=font)
        run.text = text[pos:]


def _bullets(box: Any, theme: dict[str, Any], points: list[str],
             sizes: list[float], height_in: float) -> None:
    """带强调色圆点的要点列表，按总量自适应字号；数字自动强调。"""
    tf = box.text_frame
    total = sum(len(p) for p in points) or 1
    per_size = pick_size("x" * max(24, total // max(len(points), 1)),
                         box.width / 914400, height_in / max(len(points), 1) * 1.6,
                         sizes)
    for i, raw in enumerate(points):
        title, desc = split_title_desc(raw)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        p.line_spacing = 1.22
        dot = p.add_run()
        set_run(dot, per_size, theme["accent"], bold=True, font=theme["font"])
        dot.text = "•  "
        if desc:
            main = p.add_run()
            set_run(main, per_size, theme["text"], bold=True, font=theme["font"])
            main.text = title
            sub = p.add_run()
            set_run(sub, max(per_size - 2.5, 10.5), theme["muted"],
                    font=theme["font"])
            sub.text = f"　{desc}"
        else:
            _text_runs(p, raw, per_size, theme, theme["font"])


def _card_panel(slide: Any, theme: dict[str, Any], st: dict[str, Any],
                x: float, y: float, w: float, h: float) -> Any:
    """风格化卡片面板：圆角浅底 / 黑边硬阴影 / 顶部粗线 / 暗色霓虹左边。"""
    if st.get("card_shadow"):
        add_rect(slide, x + 0.09, y + 0.09, w, h, theme["text"])
        border = st.get("card_border")
        border_color = theme["text"] if border == "text" else border
        card = add_rect(slide, x, y, w, h, theme["accent_light"], rounded=False)
        card.line.color.rgb = rgb(border_color or theme["text"])
        card.line.width = Pt(float(st.get("card_border_w") or 2.5))
        return card

    top_rule = st.get("card_toprule")
    if top_rule:
        rule_h = 0.04 if top_rule == "text" else 0.014
        rule_color = theme["text"]
        add_rect(slide, x, y, w, rule_h, rule_color)
        return None

    card = add_rect(slide, x, y, w, h, theme["accent_light"],
                    rounded=True, radius=float(st.get("card_radius") or 0.07))
    if st.get("card_neon_left"):
        add_rect(slide, x, y + 0.06, 0.05, h - 0.12, theme["accent"])
        return card
    if st.get("card_border_auto"):
        card.line.color.rgb = rgb(blend(theme["content_bg"],
                                        theme["accent_light"], 0.35))
        card.line.width = Pt(1)
    return card


def _dot(slide: Any, theme: dict[str, Any], st: dict[str, Any],
         x: float, y: float, size: float, color: str) -> Any:
    if st.get("card_dot_square"):
        return add_rect(slide, x, y, size, size, color)
    return add_oval(slide, x, y, size, size, color)


# ---------------------------------------------------------------- 封面 / 结尾背景

def _grid_lines(slide: Any, bg_hex: str, light_ratio: float = 0.06) -> None:
    line_color = blend(bg_hex, "FFFFFF", light_ratio)
    for i in range(1, 7):
        add_rect(slide, i * 13.333 / 7, 0, 0.012, 7.5, line_color)
    for i in range(1, 4):
        add_rect(slide, 0, i * 7.5 / 4, 13.333, 0.012, line_color)


def _corner_brackets(slide: Any, color: str) -> None:
    arm, inset, thickness = 0.3, 0.35, 0.028
    for (bx, by, dx, dy) in ((inset, inset, 1, 1),
                             (13.333 - inset, inset, -1, 1),
                             (inset, 7.5 - inset, 1, -1),
                             (13.333 - inset, 7.5 - inset, -1, -1)):
        add_rect(slide, min(bx, bx + dx * arm),
                 by - (thickness if dy < 0 else 0), arm, thickness, color)
        add_rect(slide, bx - (thickness if dx < 0 else 0),
                 min(by, by + dy * arm), thickness, arm, color)


def _cover_bg(slide: Any, theme: dict[str, Any], st: dict[str, Any],
              index: int) -> None:
    kind = st.get("cover") or "clean"
    if kind == "gradient_blobs":
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                    Inches(13.333), Inches(7.5))
        gradient(bg, theme["cover_bg"],
                 blend(theme["cover_bg"], theme["accent"], 0.55), angle=115)
        halo = add_oval(slide, 9.8, 3.6, 4.4, 4.4, theme["accent"])
        set_alpha(halo, 16)
        halo2 = add_oval(slide, 11.2, 0.3, 2.4, 2.4,
                         theme.get("accent2") or theme["accent"])
        set_alpha(halo2, 12)
    elif kind == "glass_panel":
        add_rect(slide, 0, 0, 13.333, 7.5, theme["cover_bg"])
        for (bx, by, size, color, alpha) in (
            (0.7, 0.4, 3.0, theme["accent"], 12),
            (9.8, 4.4, 3.2, theme.get("accent2") or theme["accent"], 11),
            (8.6, 0.2, 2.2, theme["accent"], 8),
        ):
            halo = add_oval(slide, bx, by, size, size, color)
            set_alpha(halo, alpha)
        panel = add_rect(slide, 0.95, 2.05, 8.8, 3.15, "FFFFFF", rounded=True,
                         radius=0.09)
        set_alpha(panel, 62)
        panel.line.color.rgb = rgb(blend("FFFFFF", theme["accent"], 0.18))
        panel.line.width = Pt(1)
    elif kind == "blocks":
        add_rect(slide, 0, 0, 13.333, 7.5, theme["cover_bg"])
        dark = theme["text"]
        add_rect(slide, 9.55, 0.72, 2.65, 1.95, dark)
        yellow = add_rect(slide, 9.45, 0.62, 2.65, 1.95, "FFD400")
        yellow.line.color.rgb = rgb(dark)
        yellow.line.width = Pt(3)
        add_rect(slide, 8.05, 4.72, 1.5, 1.5, dark)
        blue = add_rect(slide, 7.95, 4.62, 1.5, 1.5, theme["accent2"])
        blue.line.color.rgb = rgb(dark)
        blue.line.width = Pt(3)
        orange = add_rect(slide, 0.95, 0.62, 0.42, 1.05, theme["accent"])
        orange.line.color.rgb = rgb(dark)
        orange.line.width = Pt(3)
    elif kind == "swiss_index":
        add_rect(slide, 0, 0, 13.333, 7.5, theme["cover_bg"])
        num = add_textbox(slide, 6.8, 0.0, 6.1, 3.4)
        p = num.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        set_run(p.add_run(), 190, blend(theme["cover_bg"], theme["text"], 0.1),
                bold=True, font=theme["font"])
        p.runs[0].text = f"{index:02d}"
        add_rect(slide, 0.95, 3.62, 1.7, 0.05, theme["text"])
    elif kind == "editorial":
        add_rect(slide, 0, 0, 13.333, 7.5, theme["cover_bg"])
        for ry in (0.35, 6.95):
            add_rect(slide, 0.7, ry, 11.93, 0.045, theme["accent2"])
            add_rect(slide, 0.7, ry + 0.075, 11.93, 0.014, theme["accent2"])
    elif kind == "darktech":
        add_rect(slide, 0, 0, 13.333, 7.5, theme["cover_bg"])
        _grid_lines(slide, theme["cover_bg"], 0.055)
        _corner_brackets(slide, blend(theme["cover_bg"], theme["accent"], 0.5))
    else:  # clean
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                    Inches(13.333), Inches(7.5))
        from .themes import is_light
        if is_light(theme["cover_bg"]):
            solid(bg, theme["cover_bg"])
            halo = add_oval(slide, 9.6, 3.2, 4.6, 4.6, theme["accent"])
            set_alpha(halo, 10)
            halo2 = add_oval(slide, 11.2, 0.4, 2.4, 2.4, theme["accent"])
            set_alpha(halo2, 8)
        else:
            gradient(bg, theme["cover_bg"],
                     blend(theme["cover_bg"], theme["accent"], 0.55), angle=115)
            halo = add_oval(slide, 9.8, 3.6, 4.4, 4.4, theme["accent"])
            set_alpha(halo, 16)
            halo2 = add_oval(slide, 11.4, 0.3, 2.2, 2.2, theme["accent"])
            set_alpha(halo2, 12)


def _layout_cover(slide: Any, theme: dict[str, Any], st: dict[str, Any],
                  item: dict[str, Any], index: int) -> None:
    _cover_bg(slide, theme, st, index)
    title = str(item.get("heading") or item.get("title") or "未命名演示文稿")
    subtitle = str(item.get("subtitle") or "").strip()
    pages = int(item.get("_pages") or max(index - 1, 0))

    if st.get("cover") == "glass_panel":
        size = pick_size(title, 8.0, 1.5, [40, 36, 32])
        box = add_textbox(slide, 1.35, 2.5, 8.0, 1.6)
        p = box.text_frame.paragraphs[0]
        p.line_spacing = 1.18
        set_run(p.add_run(), size, theme["cover_text"], bold=True,
                font=theme["font"])
        p.runs[0].text = title
        if subtitle:
            sub = add_textbox(slide, 1.35, 4.25, 8.0, 0.6)
            set_run(sub.text_frame.paragraphs[0].add_run(), 15,
                    theme["cover_text"], font=theme["font"])
            sub.text_frame.paragraphs[0].runs[0].text = subtitle
        meta = add_textbox(slide, 1.35, 5.55, 8.0, 0.45)
        set_run(meta.text_frame.paragraphs[0].add_run(), 12.5, theme["meta"],
                font=theme["font"])
        meta.text_frame.paragraphs[0].runs[0].text = f"AI Divine · {pages} 页"
        return

    if st.get("cover") == "swiss_index":
        size = pick_size(title, 11.0, 1.8, [46, 42, 36, 32])
        box = add_textbox(slide, 0.95, 3.85, 11.0, 1.9)
        p = box.text_frame.paragraphs[0]
        p.line_spacing = 1.12
        set_run(p.add_run(), size, theme["cover_text"], bold=True,
                font=theme["font"])
        p.runs[0].text = title
        if subtitle:
            sub = add_textbox(slide, 0.95, 5.85, 10.4, 0.55)
            set_run(sub.text_frame.paragraphs[0].add_run(), 15, theme["muted"],
                    font=theme["font"])
            sub.text_frame.paragraphs[0].runs[0].text = subtitle
        meta = add_textbox(slide, 0.95, 6.55, 10.4, 0.45)
        set_run(meta.text_frame.paragraphs[0].add_run(), 11.5, theme["muted"],
                font=theme["font"], letter_spacing_pt=2.2)
        meta.text_frame.paragraphs[0].runs[0].text = "AI DIVINE — STRATEGY"
        return

    if st.get("cover") == "editorial":
        kicker = add_textbox(slide, 1.35, 1.15, 10.6, 0.45)
        p = kicker.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p.add_run(), 12.5, theme["accent"], bold=True, font=theme["font"],
                letter_spacing_pt=3.2)
        p.runs[0].text = "AI DIVINE"
        size = pick_size(title, 10.4, 1.7, [42, 38, 34])
        box = add_textbox(slide, 1.35, 1.95, 10.6, 1.8)
        p = box.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.line_spacing = 1.25
        set_run(p.add_run(), size, theme["cover_text"], bold=True,
                font=_heading_font(theme, st))
        p.runs[0].text = title
        add_rect(slide, 6.27, 4.15, 0.8, 0.014, theme["accent2"])
        if subtitle:
            sub = add_textbox(slide, 1.35, 4.5, 10.6, 0.6)
            p = sub.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            set_run(p.add_run(), 15.5, blend(theme["cover_bg"], theme["text"], 0.62),
                    font=_heading_font(theme, st), letter_spacing_pt=1.2)
            p.runs[0].text = subtitle
        meta = add_textbox(slide, 1.35, 5.55, 10.6, 0.45)
        p = meta.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p.add_run(), 11.5, theme["meta"], font=theme["font"],
                letter_spacing_pt=2.0)
        p.runs[0].text = f"AI DIVINE · {pages} 页"
        return

    if st.get("cover_centered"):  # mesh：深底居中 + 弥散光晕
        size = pick_size(title, 10.6, 1.6, [44, 40, 36, 32])
        box = add_textbox(slide, 1.35, 2.75, 10.6, 1.7)
        p = box.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.line_spacing = 1.15
        set_run(p.add_run(), size, theme["cover_text"], bold=True,
                font=theme["font"])
        p.runs[0].text = title
        if subtitle:
            sub = add_textbox(slide, 1.35, 4.55, 10.6, 0.6)
            p = sub.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            set_run(p.add_run(), 15, theme["meta"], font=theme["font"])
            p.runs[0].text = subtitle
        meta = add_textbox(slide, 1.35, 5.3, 10.6, 0.45)
        p = meta.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p.add_run(), 12, theme["meta"], font=_mono_font(),
                letter_spacing_pt=1.5)
        p.runs[0].text = f"AI DIVINE — {pages} 页"
        return

    if st.get("cover") == "darktech":
        neon = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.95),
                                      Inches(2.52), Inches(0.09), Inches(1.5))
        gradient(neon, theme["accent"], theme["accent2"], 90)
        kicker = add_textbox(slide, 1.28, 2.05, 10, 0.4)
        set_run(kicker.text_frame.paragraphs[0].add_run(), 12, theme["accent"],
                bold=True, font=_mono_font(), letter_spacing_pt=2.0)
        kicker.text_frame.paragraphs[0].runs[0].text = "// AI DIVINE"
        size = pick_size(title, 10.2, 1.7, [42, 38, 34])
        box = add_textbox(slide, 1.28, 2.55, 10.4, 1.8)
        p = box.text_frame.paragraphs[0]
        p.line_spacing = 1.2
        set_run(p.add_run(), size, theme["cover_text"], bold=True,
                font=theme["font"])
        p.runs[0].text = title
        if subtitle:
            sub = add_textbox(slide, 1.28, 4.45, 10.2, 0.55)
            set_run(sub.text_frame.paragraphs[0].add_run(), 14.5, theme["meta"],
                    font=theme["font"])
            sub.text_frame.paragraphs[0].runs[0].text = subtitle
        meta = add_textbox(slide, 1.28, 5.3, 10.2, 0.45)
        set_run(meta.text_frame.paragraphs[0].add_run(), 11.5,
                blend(theme["cover_bg"], theme["cover_text"], 0.4),
                font=_mono_font(), letter_spacing_pt=1.2)
        meta.text_frame.paragraphs[0].runs[0].text = f"AI DIVINE · {pages}P"
        return

    if st.get("cover") == "blocks":
        size = pick_size(title, 6.9, 2.0, [46, 42, 38, 34])
        box = add_textbox(slide, 0.95, 2.6, 6.9, 2.1)
        p = box.text_frame.paragraphs[0]
        p.line_spacing = 1.14
        set_run(p.add_run(), size, theme["cover_text"], bold=True,
                font=theme["font"])
        p.runs[0].text = title
        if subtitle:
            sub = add_textbox(slide, 0.95, 4.85, 6.9, 0.55)
            set_run(sub.text_frame.paragraphs[0].add_run(), 15, theme["muted"],
                    bold=True, font=theme["font"])
            sub.text_frame.paragraphs[0].runs[0].text = subtitle
        chip = add_rect(slide, 0.95, 5.6, 2.5, 0.52, "FFFFFF")
        chip.line.color.rgb = rgb(theme["text"])
        chip.line.width = Pt(2.5)
        p = chip.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p.add_run(), 11.5, theme["text"], bold=True, font=theme["font"])
        p.runs[0].text = f"AI DIVINE · {pages} 页"
        return

    add_rect(slide, 0.95, 2.52, 0.14, 1.35, theme["accent"])
    size = pick_size(title, 10.4, 1.9, [44, 40, 36, 32])
    box = add_textbox(slide, 1.35, 2.35, 10.6, 2.0)
    p = box.text_frame.paragraphs[0]
    p.line_spacing = 1.15
    set_run(p.add_run(), size, theme["cover_text"], bold=True, font=theme["font"])
    p.runs[0].text = title
    y = 4.45
    if subtitle:
        sub = add_textbox(slide, 1.35, 4.42, 10.6, 0.7)
        set_run(sub.text_frame.paragraphs[0].add_run(), 16, theme["meta"],
                font=theme["font"])
        sub.text_frame.paragraphs[0].runs[0].text = subtitle
        y = 5.25
    meta = add_textbox(slide, 1.35, y, 10.6, 0.5)
    set_run(meta.text_frame.paragraphs[0].add_run(), 13, theme["meta"],
            font=theme["font"])
    meta.text_frame.paragraphs[0].runs[0].text = f"AI Divine · {pages} 页"

    brand = add_textbox(slide, 1.35, 6.8, 6, 0.4)
    set_run(brand.text_frame.paragraphs[0].add_run(), 12,
            blend(theme["cover_bg"], theme["cover_text"], 0.45),
            font=theme["font"])
    brand.text_frame.paragraphs[0].runs[0].text = "AI Divine"


# ---------------------------------------------------------------- 媒体块

def _media_block(slide: Any, theme: dict[str, Any], st: dict[str, Any],
                 image: dict[str, Any] | None,
                 x: float, y: float, w: float, h: float,
                 resolve_media: Callable[[str], Any] | None) -> None:
    """图文两栏的媒体区：有图贴图（cover-crop），无图风格化占位。"""
    picture = None
    path_str = str((image or {}).get("path") or "").strip()
    if path_str and resolve_media:
        try:
            picture = resolve_media(path_str)
        except Exception:  # noqa: BLE001 — 图片不可用时降级为占位
            picture = None

    if picture is not None:
        try:
            pic = slide.shapes.add_picture(str(picture), Inches(x), Inches(y),
                                           Inches(w), Inches(h))
            _cover_crop(pic, w, h)
        except Exception:  # noqa: BLE001
            picture = None

    if picture is None:
        panel = add_rect(slide, x, y, w, h, theme["accent_light"],
                         rounded=float(st.get("card_radius") or 0) > 0,
                         radius=float(st.get("card_radius") or 0.07))
        if st.get("card_border_auto"):
            panel.line.color.rgb = rgb(blend(theme["content_bg"],
                                             theme["accent_light"], 0.35))
            panel.line.width = Pt(1)
        elif st.get("card_shadow"):
            panel.line.color.rgb = rgb(theme["text"])
            panel.line.width = Pt(float(st.get("card_border_w") or 2.5))
        glyph = add_oval(slide, x + w / 2 - 0.42, y + h / 2 - 0.86, 0.84, 0.84,
                         theme["accent"])
        set_alpha(glyph, 22)
        note = add_textbox(slide, x + 0.3, y + h / 2 + 0.14, w - 0.6, 0.4)
        p = note.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p.add_run(), 12, theme["muted"], font=theme["font"])
        p.runs[0].text = "配图占位 · 可由 image_generate 生成"

    caption = str((image or {}).get("caption") or "").strip()
    if caption:
        cap = add_textbox(slide, x, y + h + 0.08, w, 0.32)
        p = cap.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p.add_run(), 10.5, theme["muted"], font=theme["font"])
        p.runs[0].text = caption


def _cover_crop(pic: Any, w: float, h: float) -> None:
    """按目标框比例裁剪图片（cover 模式，不拉伸变形）。"""
    try:
        native_w, native_h = pic.image.size
        if not native_w or not native_h:
            return
        native_ar = native_w / native_h
        box_ar = w / h
        if native_ar > box_ar:
            crop = (1 - box_ar / native_ar) / 2
            pic.crop_left = crop
            pic.crop_right = crop
        else:
            crop = (1 - native_ar / box_ar) / 2
            pic.crop_top = crop
            pic.crop_bottom = crop
    except Exception:  # noqa: BLE001 — 裁剪失败保持拉伸
        pass


# ---------------------------------------------------------------- 各版式

def _layout_toc(slide: Any, theme: dict[str, Any], st: dict[str, Any],
                item: dict[str, Any], index: int) -> None:
    _chrome(slide, theme, st, index, str(item.get("heading") or "目录"))
    points = [str(p) for p in (item.get("points") or []) if str(p).strip()][:8]
    rule_color = _rule_color(theme, st)
    rule_h = 0.03 if st.get("toc_rule") == "text" else 0.016
    for i, raw in enumerate(points):
        title, page_ref = split_title_desc(raw)
        col, row = i % 2, i // 2
        x = MARGIN + col * (CONTENT_W / 2 + 0.25)
        y = 2.15 + row * 1.42
        w = CONTENT_W / 2 - 0.25
        num = add_textbox(slide, x, y, 0.75, 0.5)
        set_run(num.text_frame.paragraphs[0].add_run(), 15, theme["accent"],
                bold=True,
                font=_heading_font(theme, st) if st.get("index_serif") else theme["font"])
        num.text_frame.paragraphs[0].runs[0].text = f"{i + 1:02d}"
        entry = add_textbox(slide, x + 0.75, y, w - 0.75, 0.62)
        size = pick_size(title, w - 0.75, 0.6, [17, 15, 13])
        set_run(entry.text_frame.paragraphs[0].add_run(), size, theme["text"],
                bold=True, font=theme["font"])
        entry.text_frame.paragraphs[0].runs[0].text = title
        if page_ref:
            ref = add_textbox(slide, x + w - 0.9, y + 0.06, 0.9, 0.4)
            p = ref.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT
            set_run(p.add_run(), 11.5, theme["muted"], font=theme["font"])
            p.runs[0].text = page_ref
        add_rect(slide, x, y + 0.78, w, rule_h, rule_color)


def _layout_section(slide: Any, theme: dict[str, Any], st: dict[str, Any],
                    item: dict[str, Any], index: int) -> None:
    dark_section = st.get("section_bg") == "dark"
    if dark_section:
        _cover_bg(slide, theme, st, index)
        bg_hex = theme["cover_bg"]
        num_color = blend(bg_hex, theme["accent"], 0.5)
        title_color = theme["cover_text"]
        sub_color = theme["meta"]
    else:
        add_rect(slide, 0, 0, 13.333, 7.5, theme["content_bg"])
        _bg_decor(slide, theme, st)
        if st.get("cover") == "editorial":
            add_rect(slide, 0.7, 0.35, 11.93, 0.045, theme["accent2"])
            add_rect(slide, 0.7, 0.425, 11.93, 0.014, theme["accent2"])
        if st.get("cover") == "blocks":
            bar = add_rect(slide, 0.95, 0.62, 0.42, 1.05, theme["accent"])
            bar.line.color.rgb = rgb(theme["text"])
            bar.line.width = Pt(3)
        bg_hex = theme["content_bg"]
        num_color = blend(bg_hex, theme["text"],
                          0.1 if st.get("cover") == "swiss" else 0.16)
        title_color = theme["text"]
        sub_color = theme["muted"]

    heading = str(item.get("heading") or item.get("title") or "")
    sub = str(item.get("subtitle") or item.get("body") or "").strip()

    if st.get("cover_centered"):
        kicker = add_textbox(slide, 1.35, 1.9, 10.6, 0.45)
        p = kicker.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p.add_run(), 12.5, theme["accent"], bold=True,
                font=theme["font"], letter_spacing_pt=3.0)
        p.runs[0].text = f"PART {index:02d}"
        size = pick_size(heading, 10.6, 1.2, [40, 34, 28])
        title = add_textbox(slide, 1.35, 2.5, 10.6, 1.3)
        p = title.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p.add_run(), size, title_color, bold=True,
                font=_heading_font(theme, st))
        p.runs[0].text = heading
        if sub:
            sub_box = add_textbox(slide, 1.35, 4.0, 10.6, 0.6)
            p = sub_box.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            set_run(p.add_run(), 14.5, sub_color, font=theme["font"])
            p.runs[0].text = sub
        return

    number = add_textbox(slide, 0.9, 1.0, 6, 2.0)
    set_run(number.text_frame.paragraphs[0].add_run(), 96, num_color,
            bold=True, font=theme["font"])
    number.text_frame.paragraphs[0].runs[0].text = f"{index:02d}"

    kicker_font = _mono_font() if st.get("kicker_mono") else theme["font"]
    kicker_text = (f"// PART {index:02d}" if st.get("kicker_mono")
                   else f"PART {index:02d}")
    kicker = add_textbox(slide, 0.95, 3.35, 6, 0.45)
    set_run(kicker.text_frame.paragraphs[0].add_run(), 13, theme["accent"],
            bold=True, font=kicker_font,
            letter_spacing_pt=None if st.get("kicker_mono") else 1.2)
    kicker.text_frame.paragraphs[0].runs[0].text = kicker_text

    size = pick_size(heading, 11.4, 1.1, [40, 34, 28])
    title = add_textbox(slide, 0.95, 3.82, 11.4, 1.15)
    set_run(title.text_frame.paragraphs[0].add_run(), size, title_color,
            bold=True, font=_heading_font(theme, st))
    title.text_frame.paragraphs[0].runs[0].text = heading

    if sub:
        sub_box = add_textbox(slide, 0.95, 5.1, 10.8, 0.7)
        s_size = pick_size(sub, 10.8, 0.7, [15, 13.5, 12])
        set_run(sub_box.text_frame.paragraphs[0].add_run(), s_size, sub_color,
                font=theme["font"])
        sub_box.text_frame.paragraphs[0].runs[0].text = sub


def _layout_bullets(slide: Any, theme: dict[str, Any], st: dict[str, Any],
                    item: dict[str, Any], index: int) -> None:
    heading = str(item.get("heading") or item.get("title") or "")
    _chrome(slide, theme, st, index, heading)
    body = item.get("body") or item.get("content") or item.get("text") or ""
    points = [str(p) for p in (item.get("points") or []) if str(p).strip()]
    if not points and body:
        points = [line.strip() for line in body.splitlines() if line.strip()]
    if not points:
        points = ["（本页暂无内容）"]
    points = points[:6]

    n = len(points)
    area_top = CONTENT_TOP + 0.15
    area_h = CONTENT_H - 0.35
    row_h = area_h / n
    h_font = _heading_font(theme, st)
    num_font = _mono_font() if st.get("kicker_mono") else theme["font"]
    rule_color = _rule_color(theme, st)

    for i, raw in enumerate(points):
        title, desc = split_title_desc(raw)
        y = area_top + i * row_h
        content_h = min(row_h - 0.06, 1.15)

        num_label = add_textbox(slide, MARGIN, y + content_h / 2 - 0.28,
                                0.72, 0.56)
        p = num_label.text_frame.paragraphs[0]
        set_run(p.add_run(), 19, theme["accent"], bold=True, font=num_font)
        p.runs[0].text = f"{i + 1:02d}"

        text_x = MARGIN + 0.85
        text_w = CONTENT_W - 0.85
        box = add_textbox(slide, text_x, y + content_h / 2 - 0.3,
                          text_w, content_h + 0.3)
        tf = box.text_frame
        p = tf.paragraphs[0]
        p.line_spacing = 1.2
        if desc:
            t_size = pick_size(title, text_w, content_h * 0.5, [18, 16, 15, 14])
            _text_runs(p, title, t_size, theme, h_font, bold=True)
            sub = p.add_run()
            set_run(sub, max(t_size - 2.5, 11.5), theme["muted"],
                    font=theme["font"])
            sub.text = f"　{desc}"
        else:
            t_size = pick_size(raw, text_w, content_h, [17, 16, 15, 14, 13])
            _text_runs(p, raw, t_size, theme, h_font)

        if i < n - 1:
            add_rect(slide, text_x, y + row_h - 0.02, text_w, 0.014, rule_color)


def _layout_two_col(slide: Any, theme: dict[str, Any], st: dict[str, Any],
                    item: dict[str, Any], index: int,
                    resolve_media: Callable[[str], Any] | None) -> None:
    heading = str(item.get("heading") or item.get("title") or "")
    _chrome(slide, theme, st, index, heading)
    image = item.get("image") if isinstance(item.get("image"), dict) else {}
    media_w = 5.4
    text_w = CONTENT_W - media_w - 0.45
    media_left = str((image or {}).get("position") or "left") != "right"
    media_x = MARGIN if media_left else MARGIN + text_w + 0.45
    text_x = MARGIN + media_w + 0.45 if media_left else MARGIN

    _media_block(slide, theme, st, image, media_x, CONTENT_TOP, media_w,
                 CONTENT_H - 0.4, resolve_media)
    points = [str(p) for p in (item.get("points") or []) if str(p).strip()]
    if not points:
        points = [str(item.get("body") or item.get("content") or "")]
    box = add_textbox(slide, text_x, CONTENT_TOP + 0.05, text_w, CONTENT_H - 0.3)
    _bullets(box, theme, points, [16, 14.5, 13, 12], CONTENT_H - 0.3)


def _layout_cards(slide: Any, theme: dict[str, Any], st: dict[str, Any],
                  item: dict[str, Any], index: int) -> None:
    heading = str(item.get("heading") or item.get("title") or "")
    _chrome(slide, theme, st, index, heading)
    points = [str(p) for p in (item.get("points") or []) if str(p).strip()][:8]
    if not points:
        points = ["暂无内容"]
    palette = theme.get("chart_palette") or [theme["accent"]]
    n = len(points)
    cols = 2 if n <= 4 else (3 if n <= 6 else 4)
    rows = math.ceil(n / cols)
    gap = 0.3
    card_w = (CONTENT_W - gap * (cols - 1)) / cols
    card_h = min((CONTENT_H - gap * (rows - 1)) / rows, 2.3)
    desc_sizes = [12.5, 11.5, 10.5] if cols <= 3 else [11.5, 10.5, 10]
    for i, raw in enumerate(points):
        title, desc = split_title_desc(raw)
        col, row = i % cols, i // cols
        x = MARGIN + col * (card_w + gap)
        y = CONTENT_TOP + row * (card_h + gap)
        _card_panel(slide, theme, st, x, y, card_w, card_h)
        if st.get("cover") == "blocks":
            dot_color = ["FFD400", theme["accent"], theme["accent2"]][i % 3]
        elif st.get("card_dot_square") and st.get("card_toprule"):
            dot_color = theme["accent"]
        else:
            dot_color = palette[i % len(palette)]
        _dot(slide, theme, st, x + 0.22, y + 0.2, 0.14, dot_color)
        if desc:
            t_box = add_textbox(slide, x + 0.22, y + 0.48, card_w - 0.44, 0.5)
            t_size = pick_size(title, card_w - 0.44, 0.5,
                               [15, 14, 13] if cols <= 3 else [13.5, 12.5, 11.5])
            set_run(t_box.text_frame.paragraphs[0].add_run(), t_size,
                    theme["text"], bold=True, font=theme["font"])
            t_box.text_frame.paragraphs[0].runs[0].text = title
            d_box = add_textbox(slide, x + 0.22, y + 1.02, card_w - 0.44,
                                card_h - 1.2)
            d_size = pick_size(desc, card_w - 0.44, card_h - 1.2, desc_sizes)
            tf = d_box.text_frame
            p = tf.paragraphs[0]
            p.line_spacing = 1.25
            set_run(p.add_run(), d_size, theme["muted"], font=theme["font"])
            p.runs[0].text = desc
        else:
            box = add_textbox(slide, x + 0.22, y + 0.55, card_w - 0.44,
                              card_h - 0.75)
            size = pick_size(raw, card_w - 0.44, card_h - 0.75,
                             [14, 13, 12, 11])
            tf = box.text_frame
            p = tf.paragraphs[0]
            p.line_spacing = 1.25
            set_run(p.add_run(), size, theme["text"], font=theme["font"])
            p.runs[0].text = raw


def _layout_timeline(slide: Any, theme: dict[str, Any], st: dict[str, Any],
                     item: dict[str, Any], index: int) -> None:
    heading = str(item.get("heading") or item.get("title") or "")
    _chrome(slide, theme, st, index, heading)
    steps = [str(p) for p in (item.get("points") or []) if str(p).strip()][:10]
    if not steps:
        steps = ["阶段一", "阶段二", "阶段三"]

    track_kind = st.get("timeline_track") or "accent_light_mix"
    if track_kind == "text":
        track_color, track_h = theme["text"], 0.05
    elif track_kind == "hairline":
        track_color = blend(theme["content_bg"], theme["text"], 0.6)
        track_h = 0.018
    elif track_kind == "muted_mix":
        track_color = blend(theme["content_bg"], theme["muted"], 0.4)
        track_h = 0.035
    else:
        track_color = blend(theme["content_bg"], theme["accent"], 0.35)
        track_h = 0.045

    if len(steps) <= 5:
        rows = [steps]
        ys = [3.85]
    else:
        mid = math.ceil(len(steps) / 2)
        rows = [steps[:mid], steps[mid:]]
        ys = [2.75, 5.15]

    for row_points, track_y in zip(rows, ys):
        n = len(row_points)
        start_x = MARGIN + 0.35
        end_x = 13.333 - MARGIN - 0.35
        add_rect(slide, start_x, track_y, end_x - start_x, track_h, track_color)
        span = (end_x - start_x) / n
        for i, raw in enumerate(row_points):
            title, desc = split_title_desc(raw)
            cx = start_x + span * (i + 0.5)
            if st.get("timeline_dot_square"):
                if st.get("cover") == "blocks":
                    dot_color = "FFD400"
                else:
                    dot_color = theme["accent"]
                dot = add_rect(slide, cx - 0.13, track_y - 0.105, 0.26, 0.26,
                               dot_color)
            else:
                dot = add_oval(slide, cx - 0.14, track_y - 0.115, 0.3, 0.3,
                               theme["accent"])
            dot.line.color.rgb = rgb(theme["content_bg"])
            dot.line.width = Pt(2)
            box_w = min(span, 2.8)
            t_box = add_textbox(slide, cx - box_w / 2, track_y + 0.32,
                                box_w, 0.55)
            t_size = pick_size(title, box_w, 0.55, [14, 13, 12, 11])
            p = t_box.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            set_run(p.add_run(), t_size, theme["text"], bold=True,
                    font=theme["font"])
            p.runs[0].text = title
            if desc:
                d_box = add_textbox(slide, cx - box_w / 2, track_y + 0.88,
                                    box_w, 0.85)
                d_size = pick_size(desc, box_w, 0.85, [11, 10.5, 10])
                p = d_box.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                p.line_spacing = 1.2
                set_run(p.add_run(), d_size, theme["muted"], font=theme["font"])
                p.runs[0].text = desc


def _layout_columns(slide: Any, theme: dict[str, Any], st: dict[str, Any],
                    item: dict[str, Any], index: int) -> None:
    heading = str(item.get("heading") or item.get("title") or "")
    _chrome(slide, theme, st, index, heading)
    cols = [str(p) for p in (item.get("points") or []) if str(p).strip()]
    if not cols:
        cols = ["第一栏", "第二栏", "第三栏"]
    n = min(max(len(cols), 2), 4)
    gap = 0.35
    col_w = (CONTENT_W - gap * (n - 1)) / n
    col_h = CONTENT_H - 0.2
    for i in range(n):
        x = MARGIN + i * (col_w + gap)
        _card_panel(slide, theme, st, x, CONTENT_TOP, col_w, col_h)
        text = cols[i] if i < len(cols) else ""
        pad_top = 0.34 if st.get("card_toprule") else 0.3
        box = add_textbox(slide, x + 0.28, CONTENT_TOP + pad_top, col_w - 0.56,
                          col_h - 0.6)
        size = pick_size(text, col_w - 0.56, col_h - 0.6, [13, 12, 11, 10.5])
        tf = box.text_frame
        p = tf.paragraphs[0]
        p.line_spacing = 1.3
        set_run(p.add_run(), size, theme["text"], font=theme["font"])
        p.runs[0].text = text


def _layout_compare(slide: Any, theme: dict[str, Any], st: dict[str, Any],
                    item: dict[str, Any], index: int) -> None:
    left = item.get("left") if isinstance(item.get("left"), dict) else None
    right = item.get("right") if isinstance(item.get("right"), dict) else None
    if not left or not right:
        _layout_columns(slide, theme, st, item, index)
        return
    heading = str(item.get("heading") or item.get("title") or "")
    _chrome(slide, theme, st, index, heading)

    panel_w = (CONTENT_W - 1.0) / 2
    panel_h = CONTENT_H - 0.15
    badge_square = bool(st.get("card_dot_square")) or float(st.get("card_radius") or 0.07) == 0
    tag_colors = (theme["accent"], blend(theme["text"], theme["accent"], 0.3))
    for spec_obj, tag_color, px in ((left, tag_colors[0], MARGIN),
                                    (right, tag_colors[1],
                                     MARGIN + panel_w + 1.0)):
        _card_panel(slide, theme, st, px, CONTENT_TOP, panel_w, panel_h)
        tag_text = str(spec_obj.get("title") or "方案")
        tag = add_rect(slide, px + 0.3, CONTENT_TOP + 0.3, 1.7, 0.44,
                       tag_color, rounded=not badge_square,
                       radius=0.0 if badge_square else 0.5)
        if st.get("card_shadow"):
            tag.line.color.rgb = rgb(theme["text"])
            tag.line.width = Pt(2.5)
        tp = tag.text_frame
        tp.word_wrap = False
        p = tp.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p.add_run(), 12, "FFFFFF", bold=True, font=theme["font"])
        p.runs[0].text = tag_text[:6]
        points = [str(pt) for pt in (spec_obj.get("points") or [])][:5]
        box = add_textbox(slide, px + 0.3, CONTENT_TOP + 0.95, panel_w - 0.6,
                          panel_h - 1.3)
        _bullets(box, theme, points, [13.5, 12.5, 11.5], panel_h - 1.3)

    badge_x = 6.667 - 0.36
    badge_y = CONTENT_TOP + panel_h / 2 - 0.36
    if badge_square:
        badge_fill = theme["accent"] if st.get("cover") == "blocks" else theme["text"]
        badge = add_rect(slide, badge_x, badge_y, 0.72, 0.72, badge_fill)
        if st.get("card_shadow"):
            badge.line.color.rgb = rgb(theme["text"])
            badge.line.width = Pt(2.5)
    else:
        badge = add_oval(slide, badge_x, badge_y, 0.72, 0.72, theme["text"])
    p = badge.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    set_run(p.add_run(), 14, "FFFFFF", bold=True, font=theme["font"])
    p.runs[0].text = "VS"


def _layout_kpi(slide: Any, theme: dict[str, Any], st: dict[str, Any],
                item: dict[str, Any], index: int) -> None:
    heading = str(item.get("heading") or item.get("title") or "")
    _chrome(slide, theme, st, index, heading)
    points = [str(p) for p in (item.get("points") or []) if str(p).strip()][:4]
    if not points:
        points = ["--|待补充"]
    n = len(points)
    gap = 0.3
    card_w = (CONTENT_W - gap * (n - 1)) / n
    card_h = 2.6
    y = CONTENT_TOP + 0.6
    for i, raw in enumerate(points):
        parts = [part.strip() for part in raw.split("|")]
        value = parts[0] or "--"
        label = parts[1] if len(parts) > 1 else ""
        delta = parts[2] if len(parts) > 2 else ""
        x = MARGIN + i * (card_w + gap)
        _card_panel(slide, theme, st, x, y, card_w, card_h)
        v_size = pick_size(value, card_w - 0.5, 0.85, [34, 30, 26, 22])
        v_box = add_textbox(slide, x + 0.25, y + 0.4, card_w - 0.5, 0.9)
        set_run(v_box.text_frame.paragraphs[0].add_run(), v_size,
                theme["accent"], bold=True, font=theme["font"])
        v_box.text_frame.paragraphs[0].runs[0].text = value
        if label:
            l_box = add_textbox(slide, x + 0.25, y + 1.38, card_w - 0.5, 0.45)
            l_size = pick_size(label, card_w - 0.5, 0.45, [14, 13, 12])
            set_run(l_box.text_frame.paragraphs[0].add_run(), l_size,
                    theme["text"], bold=True, font=theme["font"])
            l_box.text_frame.paragraphs[0].runs[0].text = label
        if delta:
            d_box = add_textbox(slide, x + 0.25, y + 1.86, card_w - 0.5, 0.4)
            set_run(d_box.text_frame.paragraphs[0].add_run(), 11.5,
                    theme["muted"], font=theme["font"])
            d_box.text_frame.paragraphs[0].runs[0].text = delta


def _layout_quote(slide: Any, theme: dict[str, Any], st: dict[str, Any],
                  item: dict[str, Any], index: int) -> None:
    dark_quote = st.get("cover") == "darktech"
    if dark_quote:
        _cover_bg(slide, theme, st, index)
        _grid_lines(slide, theme["cover_bg"], 0.05)
        quote_color = theme["cover_text"]
        author_color = theme["meta"]
        mark_color = theme["accent"]
    else:
        add_rect(slide, 0, 0, 13.333, 7.5, theme["content_bg"])
        quote_color = theme["text"]
        author_color = theme["muted"]
        mark_color = blend(theme["content_bg"], theme["accent"], 0.45)
        if st.get("cover") == "editorial":
            add_rect(slide, 0.7, 0.35, 11.93, 0.045, theme["accent2"])
            add_rect(slide, 0.7, 0.425, 11.93, 0.014, theme["accent2"])
            add_rect(slide, 0.7, 6.905, 11.93, 0.014, theme["accent2"])
            add_rect(slide, 0.7, 6.95, 11.93, 0.045, theme["accent2"])

    mark = add_textbox(slide, 1.1, 0.75, 2.5, 1.6)
    set_run(mark.text_frame.paragraphs[0].add_run(), 88, mark_color,
            bold=True, font="Georgia")
    mark.text_frame.paragraphs[0].runs[0].text = "\u201C"

    quote = str(item.get("heading") or item.get("body") or "").strip()
    q_font = _heading_font(theme, st) if st.get("quote_serif") else theme["font"]
    size = pick_size(quote, 10.2, 2.3, [24, 21, 18, 16])
    box = add_textbox(slide, 1.55, 2.55, 10.2, 2.4)
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.line_spacing = 1.45
    set_run(p.add_run(), size, quote_color, bold=True, font=q_font)
    p.runs[0].text = quote or ""

    author = str(item.get("subtitle") or "").strip()
    if author:
        a_box = add_textbox(slide, 1.55, 5.25, 10.2, 0.5)
        p = a_box.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        a_font = _mono_font() if st.get("kicker_mono") else theme["font"]
        set_run(p.add_run(), 13, author_color, font=a_font)
        p.runs[0].text = f"—— {author}"


def _layout_table(slide: Any, theme: dict[str, Any], st: dict[str, Any],
                  item: dict[str, Any], index: int) -> None:
    heading = str(item.get("heading") or item.get("title") or "")
    _chrome(slide, theme, st, index, heading)
    spec = item.get("table") if isinstance(item.get("table"), dict) else None
    if not spec:
        raise ValueError("table 版式需要 table: {columns, rows} 数据")
    caption = str((item.get("image") or {}).get("caption") or "").strip()
    box_h = CONTENT_H - (0.45 if caption else 0.1)
    tables.add_table(slide, theme, spec, MARGIN, CONTENT_TOP, CONTENT_W, box_h)
    if caption:
        cap = add_textbox(slide, MARGIN, 6.55, CONTENT_W, 0.32)
        set_run(cap.text_frame.paragraphs[0].add_run(), 10.5, theme["muted"],
                font=theme["font"])
        cap.text_frame.paragraphs[0].runs[0].text = caption


def _layout_chart(slide: Any, theme: dict[str, Any], st: dict[str, Any],
                  item: dict[str, Any], index: int) -> None:
    heading = str(item.get("heading") or item.get("title") or "")
    _chrome(slide, theme, st, index, heading)
    spec = item.get("chart") if isinstance(item.get("chart"), dict) else None
    if not spec:
        raise ValueError(
            "chart 版式需要 chart: {type, categories, series} 数据；"
            "复杂图表请用 generate_chart 生成 SVG 后以 image.path 引用"
        )
    points = [str(p) for p in (item.get("points") or []) if str(p).strip()]
    caption = str((item.get("image") or {}).get("caption") or "").strip()

    if points:
        chart_w = CONTENT_W * 0.6
        charts.add_native_chart(slide, theme, spec, MARGIN, CONTENT_TOP,
                                chart_w, CONTENT_H - 0.25)
        box = add_textbox(slide, MARGIN + chart_w + 0.45, CONTENT_TOP + 0.05,
                          CONTENT_W - chart_w - 0.45, CONTENT_H - 0.4)
        _bullets(box, theme, points, [14.5, 13, 12, 11], CONTENT_H - 0.4)
    else:
        charts.add_native_chart(slide, theme, spec, 1.15, CONTENT_TOP,
                                CONTENT_W - 0.9, CONTENT_H - 0.25)
    if caption:
        cap = add_textbox(slide, MARGIN, 6.62, CONTENT_W, 0.32)
        set_run(cap.text_frame.paragraphs[0].add_run(), 10.5, theme["muted"],
                font=theme["font"])
        cap.text_frame.paragraphs[0].runs[0].text = caption


def _layout_end(slide: Any, theme: dict[str, Any], st: dict[str, Any],
                item: dict[str, Any], index: int) -> None:
    _cover_bg(slide, theme, st, index)
    if st.get("cover") == "editorial":
        add_rect(slide, 0.7, 0.35, 11.93, 0.045, theme["accent2"])
        add_rect(slide, 0.7, 0.425, 11.93, 0.014, theme["accent2"])
        add_rect(slide, 0.7, 6.905, 11.93, 0.014, theme["accent2"])
        add_rect(slide, 0.7, 6.95, 11.93, 0.045, theme["accent2"])
    if st.get("cover") == "darktech":
        _corner_brackets(slide, blend(theme["cover_bg"], theme["accent"], 0.5))
    if st.get("cover") == "blocks":
        yellow = add_rect(slide, 10.2, 5.35, 1.9, 1.35, "FFD400")
        yellow.line.color.rgb = rgb(theme["text"])
        yellow.line.width = Pt(3)

    heading = str(item.get("heading") or item.get("title") or "谢谢观看")
    size = pick_size(heading, 10, 1.2, [42, 38, 32])
    box = add_textbox(slide, 1.6, 2.7, 10.1, 1.3)
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    set_run(p.add_run(), size, theme["cover_text"], bold=True,
            font=_heading_font(theme, st))
    p.runs[0].text = heading

    sub = str(item.get("subtitle") or "").strip()
    if sub:
        s_box = add_textbox(slide, 1.6, 4.05, 10.1, 0.55)
        p = s_box.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p.add_run(), 15, theme["meta"], font=theme["font"])
        p.runs[0].text = sub

    contact = str(item.get("body") or "").strip()
    if contact:
        c_box = add_textbox(slide, 1.6, 4.85, 10.1, 0.5)
        p = c_box.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p.add_run(), 12.5,
                blend(theme["cover_bg"], theme["cover_text"], 0.5),
                font=theme["font"])
        p.runs[0].text = contact


# ---------------------------------------------------------------- 分发

_POINT_TITLE_KEYS = ("title", "heading", "name", "label", "stage", "round",
                     "step", "period", "time", "date", "year", "metric")
_POINT_DESC_KEYS = ("desc", "description", "content", "text", "detail",
                    "summary", "note", "caption", "value")


def _normalize_point(p: Any) -> str:
    """points 条目健壮化：LLM 常违反 schema 传 dict，按字段语义拼接，
    避免把 Python dict 原文印到页面上。"""
    if not isinstance(p, dict):
        return str(p).strip()
    keys = set(p.keys())
    # KPI 形态：{value, label, delta} → "数值|标签|说明"
    if "value" in keys:
        parts = [str(p.get(k) or "").strip()
                 for k in ("value", "label", "delta")]
        joined = "|".join(x for x in parts if x)
        if joined:
            return joined
    title = next((str(p[k]).strip() for k in _POINT_TITLE_KEYS
                  if k in keys and str(p[k]).strip()), "")
    desc = next((str(p[k]).strip() for k in _POINT_DESC_KEYS
                 if k in keys and str(p[k]).strip()), "")
    used = ({k for k in _POINT_TITLE_KEYS if k in keys}
            | {k for k in _POINT_DESC_KEYS if k in keys})
    extras = " · ".join(
        f"{k}: {v}" for k, v in p.items()
        if k not in used and str(v).strip()
    )
    desc_full = " ".join(x for x in (desc, extras) if x)
    if title and desc_full:
        return f"{title}|{desc_full}"
    return title or desc_full or str(p)


def _normalize_points(points: Any) -> list[str]:
    if not isinstance(points, list):
        return points if isinstance(points, list) else []
    return [_normalize_point(p) for p in points if _normalize_point(p)]


def _bullets_overloaded(item: dict[str, Any]) -> bool:
    """要点 ≥7 条（或正文 ≥7 行）的 bullets 页视为超载。"""
    points = item.get("points")
    if isinstance(points, list):
        if len([p for p in points if str(p).strip()]) >= 7:
            return True
    body = str(item.get("body") or item.get("content") or item.get("text") or "")
    if body:
        lines = [line for line in body.splitlines() if line.strip()]
        if len(lines) >= 7:
            return True
    return False


def guess_layout(item: Any) -> str:
    """未显式指定 layout 时的保守推断。"""
    heading = str(item.get("heading") or item.get("title") or "") if isinstance(item, dict) else ""
    lowered = heading.lower()
    if "目录" in heading or "contents" in lowered:
        return "toc"
    if any(word in heading for word in ("时间线", "路线", "里程碑", "阶段")) \
            or any(word in lowered for word in ("timeline", "roadmap")):
        return "timeline"
    points = item.get("points") if isinstance(item, dict) else None
    count = len(points) if isinstance(points, list) else 0
    if count >= 5:
        return "cards"
    if 0 < count <= 3:
        return "columns"
    return "bullets"


def add_slide(prs: Any, item: Any, theme_id: str, index: int,
              resolve_media: Callable[[str], Any] | None = None) -> str:
    """渲染一页内容版式，返回实际使用的 layout 名。"""
    theme = get_ppt_theme(theme_id)
    st = get_style(theme)
    item = item if isinstance(item, dict) else {"body": str(item)}
    if isinstance(item.get("points"), list):
        item = dict(item)
        item["points"] = _normalize_points(item["points"])
    layout = str(item.get("layout") or "").strip().lower()
    if layout not in LAYOUTS:
        layout = guess_layout(item)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if layout == "bullets" and _bullets_overloaded(item):
        # 超密列表自动升级为卡片版式（数据不变，观感兜底）
        layout = "cards"
    if layout == "cover":
        _layout_cover(slide, theme, st, item, index)
    elif layout == "toc":
        _layout_toc(slide, theme, st, item, index)
    elif layout == "section":
        _layout_section(slide, theme, st, item, index)
    elif layout == "two_col":
        _layout_two_col(slide, theme, st, item, index, resolve_media)
    elif layout == "cards":
        _layout_cards(slide, theme, st, item, index)
    elif layout == "timeline":
        _layout_timeline(slide, theme, st, item, index)
    elif layout == "columns":
        _layout_columns(slide, theme, st, item, index)
    elif layout == "compare":
        _layout_compare(slide, theme, st, item, index)
    elif layout == "kpi":
        _layout_kpi(slide, theme, st, item, index)
    elif layout == "quote":
        _layout_quote(slide, theme, st, item, index)
    elif layout == "table":
        _layout_table(slide, theme, st, item, index)
    elif layout == "chart":
        _layout_chart(slide, theme, st, item, index)
    elif layout == "end":
        _layout_end(slide, theme, st, item, index)
    else:
        _layout_bullets(slide, theme, st, item, index)
    return layout
