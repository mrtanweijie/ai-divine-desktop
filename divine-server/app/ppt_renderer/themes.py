"""PPT 主题注册表与低层绘制助手。

主题 token：颜色 / 字体 / 图表调色板。所有版式渲染统一从这里取 token，
保证封面、正文、图表、表格在同一主题下色系一致。
"""

from __future__ import annotations

from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

PPT_THEMES: dict[str, dict[str, Any]] = {
    "universal": {
        "cover_bg": "10131A", "content_bg": "FFFFFF", "accent": "2563EB",
        "accent_light": "EFF6FF", "text": "1A1D24", "muted": "6B7280",
        "cover_text": "FFFFFF", "meta": "9AA3B2", "font": "PingFang SC",
        "chart_palette": ["2563EB", "0EA5E9", "7C3AED", "10B981", "F59E0B"],
    },
    "notion": {
        "cover_bg": "FFFFFF", "content_bg": "FFFFFF", "accent": "37352F",
        "accent_light": "F7F6F3", "text": "37352F", "muted": "9B9B97",
        "cover_text": "37352F", "meta": "B9B9B5", "font": "PingFang SC",
        "chart_palette": ["37352F", "9B9B97", "D9730D", "2F6FEB", "5B8A72"],
    },
    "linear": {
        "cover_bg": "202129", "content_bg": "FFFFFF", "accent": "5E6AD2",
        "accent_light": "F0F1FF", "text": "202129", "muted": "5C5F6E",
        "cover_text": "FFFFFF", "meta": "8F92A3", "font": "PingFang SC",
        "chart_palette": ["5E6AD2", "8B5CF6", "0EA5E9", "22C55E", "F59E0B"],
    },
    "apple": {
        "cover_bg": "F5F5F7", "content_bg": "FFFFFF", "accent": "0071E3",
        "accent_light": "E8F2FF", "text": "1D1D1F", "muted": "6E6E73",
        "cover_text": "1D1D1F", "meta": "AEAEB2", "font": "PingFang SC",
        "chart_palette": ["0071E3", "30D158", "FF9F0A", "BF5AF2", "64D2FF"],
    },
    "modern": {
        "cover_bg": "10131A", "content_bg": "FFFFFF", "accent": "2563EB",
        "accent_light": "EFF6FF", "text": "1A1D24", "muted": "6B7280",
        "cover_text": "FFFFFF", "meta": "9AA3B2", "font": "PingFang SC",
        "chart_palette": ["2563EB", "0EA5E9", "7C3AED", "10B981", "F59E0B"],
    },
    "academic": {
        "cover_bg": "F7FAFC", "content_bg": "FFFFFF", "accent": "742A2A",
        "accent_light": "FEF5F5", "text": "1A202C", "muted": "718096",
        "cover_text": "1A202C", "meta": "A0AEC0", "font": "Noto Serif SC",
        "chart_palette": ["742A2A", "2C5282", "276749", "975A16", "6B46C1"],
    },
    "travel": {
        "cover_bg": "FFF7ED", "content_bg": "FFFDF7", "accent": "EA580C",
        "accent_light": "FFF3E6", "text": "1C1917", "muted": "A8A29E",
        "cover_text": "9A3412", "meta": "FB923C", "font": "PingFang SC",
        "chart_palette": ["EA580C", "D97706", "059669", "0EA5E9", "DC2626"],
    },
    "roadmap": {
        "cover_bg": "1E1B4B", "content_bg": "FFFFFF", "accent": "7C3AED",
        "accent_light": "F3E8FF", "text": "1E1B4B", "muted": "A78BFA",
        "cover_text": "FFFFFF", "meta": "C4B5FD", "font": "PingFang SC",
        "chart_palette": ["7C3AED", "4F46E5", "A855F7", "0EA5E9", "EC4899"],
    },
    "minimal": {
        "cover_bg": "FFFFFF", "content_bg": "FFFFFF", "accent": "222222",
        "accent_light": "F5F5F5", "text": "222222", "muted": "999999",
        "cover_text": "111111", "meta": "BBBBBB", "font": "PingFang SC",
        "chart_palette": ["222222", "666666", "3B82F6", "10B981", "999999"],
    },
}

DEFAULT_THEME = "universal"


def get_ppt_theme(theme_id: str | None) -> dict[str, Any]:
    return PPT_THEMES.get((theme_id or "").strip(), PPT_THEMES[DEFAULT_THEME])


def blend(c1: str, c2: str, t: float) -> str:
    """把两个 hex 颜色按 t（0=c1，1=c2）混合，返回 hex。"""
    a = [int(c1[i:i + 2], 16) for i in (0, 2, 4)]
    b = [int(c2[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(a[i] + (b[i] - a[i]) * t):02X}" for i in range(3))


def is_light(hex_color: str) -> bool:
    r, g, b = (int(hex_color[i:i + 2], 16) for i in (0, 2, 4))
    return (0.299 * r + 0.587 * g + 0.114 * b) / 255 > 0.65


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def solid(shape: Any, color: str, line: str | None = None) -> Any:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def set_alpha(shape: Any, opacity_pct: float) -> Any:
    """给纯色填充叠加透明度（opacity_pct：0-100，100 为不透明）。"""
    sp_pr = shape.fill._xPr
    solid_fill = sp_pr.find(qn("a:solidFill"))
    if solid_fill is None:
        return shape
    srgb = solid_fill.find(qn("a:srgbClr"))
    if srgb is None:
        return shape
    for old in srgb.findall(qn("a:alpha")):
        srgb.remove(old)
    alpha = srgb.makeelement(qn("a:alpha"), {"val": str(int(opacity_pct * 1000))})
    srgb.append(alpha)
    return shape


def gradient(shape: Any, c1: str, c2: str, angle: float = 115.0) -> Any:
    """双色线性渐变填充。"""
    fill = shape.fill
    fill.gradient()
    stops = fill.gradient_stops
    stops[0].color.rgb = rgb(c1)
    stops[0].position = 0.0
    stops[1].color.rgb = rgb(c2)
    stops[1].position = 1.0
    try:
        fill.gradient_angle = angle
    except Exception:  # noqa: BLE001 — 部分渐变类型不支持角度，保持默认
        pass
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def set_run(run: Any, size: float, color: str, bold: bool = False,
            font: str | None = None, italic: bool = False) -> Any:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    run.font.name = font or "PingFang SC"
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.find(qn("a:ea"))
    if ea is None:
        ea = r_pr.makeelement(qn("a:ea"), {})
        r_pr.append(ea)
    ea.set("typeface", font or "PingFang SC")
    return run


def add_textbox(slide: Any, x: float, y: float, w: float, h: float) -> Any:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for margin in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, margin, 0)
    return box


def add_rect(slide: Any, x: float, y: float, w: float, h: float,
             color: str, rounded: bool = False, radius: float = 0.08) -> Any:
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    if rounded:
        try:
            shape.adjustments[0] = radius
        except Exception:  # noqa: BLE001 — 圆角系数不可调时保持默认
            pass
    return solid(shape, color)


def add_oval(slide: Any, x: float, y: float, w: float, h: float, color: str) -> Any:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    return solid(shape, color)
