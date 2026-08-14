from __future__ import annotations

from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


WORD_THEMES: dict[str, dict[str, Any]] = {
    "universal": {
        "name": "通用现代",
        "font": "PingFang SC",
        "size": 10.5,
        "line": 1.6,
        "margin": 2.2,
        "accent": "2563EB",
        "title": "111827",
        "heading": "1E40AF",
    },
    "notion": {
        "name": "Notion 极简",
        "font": "PingFang SC",
        "size": 10.5,
        "line": 1.8,
        "margin": 2.5,
        "accent": "37352F",
        "title": "37352F",
        "heading": "37352F",
    },
    "linear": {
        "name": "Linear 紧凑",
        "font": "PingFang SC",
        "size": 9.5,
        "line": 1.4,
        "margin": 1.8,
        "accent": "5E6AD2",
        "title": "202129",
        "heading": "202129",
    },
    "apple": {
        "name": "Apple 舒适",
        "font": "PingFang SC",
        "size": 10.5,
        "line": 1.9,
        "margin": 2.4,
        "accent": "0071E3",
        "title": "1D1D1F",
        "heading": "1D1D1F",
    },
    "modern": {
        "name": "现代科技",
        "font": "PingFang SC",
        "size": 10.5,
        "line": 1.6,
        "margin": 2.2,
        "accent": "2563EB",
        "title": "111827",
        "heading": "1E40AF",
    },
    "academic": {
        "name": "学术典雅",
        "font": "Noto Serif SC",
        "size": 11,
        "line": 2.0,
        "margin": 2.5,
        "accent": "742A2A",
        "title": "1A202C",
        "heading": "742A2A",
    },
    "travel": {
        "name": "旅行暖橙",
        "font": "PingFang SC",
        "size": 10.5,
        "line": 1.7,
        "margin": 2.0,
        "accent": "EA580C",
        "title": "9A3412",
        "heading": "EA580C",
    },
    "roadmap": {
        "name": "知识图谱",
        "font": "PingFang SC",
        "size": 10.5,
        "line": 1.7,
        "margin": 2.0,
        "accent": "7C3AED",
        "title": "4C1D95",
        "heading": "7C3AED",
    },
    "minimal": {
        "name": "极简黑白",
        "font": "PingFang SC",
        "size": 10.5,
        "line": 2.0,
        "margin": 2.6,
        "accent": "222222",
        "title": "111111",
        "heading": "222222",
    },
}


def _set_style_font(style: Any, font_name: str, size: float, color: str) -> None:
    style.font.name = font_name
    style.font.size = Pt(size)
    if isinstance(color, str):
        style.font.color.rgb = RGBColor.from_string(color)
    else:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)
    rfonts.set(qn("w:eastAsia"), font_name)


def apply_word_theme(document: Document, theme_id: str) -> None:
    theme = WORD_THEMES.get(theme_id, WORD_THEMES["universal"])
    font_name = theme["font"]
    accent = RGBColor.from_string(theme["accent"])
    title_color = RGBColor.from_string(theme["title"])
    heading_color = RGBColor.from_string(theme["heading"])

    normal = document.styles["Normal"]
    _set_style_font(normal, font_name, theme["size"], "111827")
    normal.paragraph_format.line_spacing = theme["line"]
    normal.paragraph_format.space_after = Pt(6)

    heading_map = {
        "Heading 1": (18, title_color, 10),
        "Heading 2": (14, heading_color, 8),
        "Heading 3": (12, heading_color, 6),
    }
    for style_name, (size, color, space_before) in heading_map.items():
        style = document.styles[style_name]
        _set_style_font(style, font_name, size, color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(4)

    for section in document.sections:
        section.top_margin = Cm(theme["margin"])
        section.bottom_margin = Cm(theme["margin"] + 0.4)
        section.left_margin = Cm(theme["margin"] - 0.2)
        section.right_margin = Cm(theme["margin"] - 0.2)


def add_title_page(document: Document, title: str, theme_id: str) -> None:
    theme = WORD_THEMES.get(theme_id, WORD_THEMES["universal"])
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(theme["title"])
    paragraph.paragraph_format.space_before = Pt(120)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("AI Divine 生成文档")
    subtitle_run.font.size = Pt(12)
    subtitle_run.font.color.rgb = RGBColor.from_string(theme["accent"])
    subtitle.paragraph_format.space_after = Pt(36)


def shade_table_header(document: Document, table: Any, theme_id: str) -> None:
    if not table.rows:
        return
    theme = WORD_THEMES.get(theme_id, WORD_THEMES["universal"])
    accent = theme["accent"]
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = tc_pr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            tc_pr.append(shd)
        shd.set(qn("w:fill"), accent)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor.from_string("FFFFFF")
                run.font.bold = True
